"""Same-fixture budget pagination: only Q15 overview paragraph styles may change."""
import argparse
import copy
import json
from pathlib import Path
import subprocess
import zipfile
from bs4 import BeautifulSoup
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from lxml import etree
from artifact_utils import sha
from check_word_rhythm_outputs import assert_styles, compare_questions, inspect_preview
from check_identifier_outputs import body_pages
from check_identifier_followup_outputs import check_followup_page_text
from check_narrative_outputs import compact, page_bounds


def xml(node): return etree.tostring(node,method='c14n',exclusive=True)


def body(document):
    nodes=list(document.element.body)
    first=next(i for i,n in enumerate(nodes) if n.tag==qn('w:p') and Paragraph(n,document).text.startswith('1. '))
    return nodes[first:]


def compare_word(before,after,eligible):
    left,right=body(before),body(after); assert len(left)==len(right)
    in_overview=False; changed=0
    for old,new in zip(left,right):
        assert old.tag==new.tag
        if old.tag==qn('w:p'):
            p,q=Paragraph(old,before),Paragraph(new,after)
            if p.style.name=='Heading 3' and p.text.startswith('15. '): in_overview=True
        elif old.tag==qn('w:tbl'): in_overview=False
        if xml(old)==xml(new): continue
        assert eligible and in_overview and old.tag==qn('w:p'), 'Unexpected non-overview XML change'
        p,q=Paragraph(old,before),Paragraph(new,after)
        assert p.text==q.text
        list_item=old.find('.//'+qn('w:numPr')) is not None
        assert q.style.name==('Pilot List Lead' if list_item else 'Pilot Lead')
        assert p.style.name in ['Normal','First Paragraph','Body Text','Compact'], 'Original style must be retained'
        current=copy.deepcopy(new); props=current.find(qn('w:pPr')); style=props.find(qn('w:pStyle'))
        assert style is not None; props.remove(style)
        old_props=old.find(qn('w:pPr'))
        old_style=old_props.find(qn('w:pStyle')) if old_props is not None else None
        if old_style is not None: props.insert(0,copy.deepcopy(old_style))
        if old_props is None and not len(props) and not props.attrib: current.remove(props)
        assert xml(old)==xml(current), 'Only the declared paragraph style may change'
        changed+=1
    assert bool(changed)==eligible, 'Small overview must change; long/many cases must not'
    links=lambda d: sorted(r.target_ref for r in d.part.rels.values() if r.is_external)
    assert links(before)==links(after), 'Link destinations changed'
    return changed


def question_body(pdf,soup):
    first=compact(soup.select_one('.question h3').get_text())
    text=body_pages(pdf); assert text.count(first)==1
    return text.split(first,1)[1]


def overview_pages(pdf,document):
    pages=[compact(p) for p in subprocess.check_output(['pdftotext','-layout',str(pdf),'-'],text=True).split('\f')]
    active=False; values=[]
    for n in body(document):
        if n.tag==qn('w:p'):
            p=Paragraph(n,document)
            if p.style.name=='Heading 3' and p.text.startswith('15. '): active=True
            if active and p.text.strip(): values.append(compact(p.text))
        elif active and n.tag==qn('w:tbl'):
            # Keep paragraph identities; table reading order differs from HTML.
            for node in n.iter(qn('w:p')):
                value=compact(Paragraph(node,document).text)
                if value: values.append(value)
    assert values
    locations=[]
    for value in values:
        hits={i for i,p in enumerate(pages,1) if value in p}
        assert hits, ('Missing Q15 paragraph in preview',value)
        locations.append(hits)
    return {'all_q15_on_one_page':sorted(set.intersection(*locations)),
            'q15_heading_page':sorted(locations[0]),'last_budget_paragraph_pages':sorted(locations[-1])}


def main():
    p=argparse.ArgumentParser(description=__doc__)
    for n in ['build','prior']: p.add_argument('--'+n,type=Path,required=True)
    p.add_argument('--cases',nargs='+',required=True); a=p.parse_args()
    helpers=['artifact_utils.py','check_word_rhythm_outputs.py','check_identifier_outputs.py','check_identifier_followup_outputs.py','check_narrative_outputs.py']
    report={'selected_checks_passed':False,'release_acceptance':False,'rows':[],
        'checker_sha256':sha(Path(__file__)),'helper_sha256':{n:sha(Path(__file__).with_name(n)) for n in helpers},
        'package_sha256':{n:sha(a.build/n) for n in ['english.zip','chinese.zip']},
        'prior_package_sha256':{n:sha(a.prior/n) for n in ['english.zip','chinese.zip']},
        'artifact_sha256':{str(f.relative_to(a.build)):sha(f) for folder in ['renders','word-preview'] for f in sorted((a.build/folder).glob('*')) if f.is_file()},
        'prior_artifact_sha256':{str(f.relative_to(a.prior)):sha(f) for folder in ['renders','word-preview'] for f in sorted((a.prior/folder).glob('*')) if f.is_file()},
        'limits':['Selected synthetic fixtures, not all DMPs','Only short Q15 Word overview keep styles may change','Same-font LibreOffice preview, not Microsoft Word acceptance','Stock Markdown table failure remains a release blocker']}
    target=a.build/'budget-pagination-report.json'
    try:
        for case in a.cases:
            for language in ['english','chinese']:
                name=case+'-'+language; old=a.prior/'renders'/name; new=a.build/'renders'/name
                for fmt in ['html','pdf','docx']:
                    fs=[json.loads(n.with_suffix('.'+fmt+'.fixture.json').read_text()) for n in [old,new]]
                    for key in ['recipe_sha256','events_sha256','km_sha256']: assert fs[0][key]==fs[1][key]
                    assert fs[0]['package_sha256']==report['prior_package_sha256'][language+'.zip']
                    assert fs[1]['package_sha256']==report['package_sha256'][language+'.zip']
                soups=[BeautifulSoup(n.with_suffix('.html').read_text(),'html.parser') for n in [old,new]]
                count=compare_questions(*soups); documents=[Document(n.with_suffix('.docx')) for n in [old,new]]
                assert_styles(documents[1])
                with zipfile.ZipFile(old.with_suffix('.docx')) as z1,zipfile.ZipFile(new.with_suffix('.docx')) as z2:
                    assert z1.read('word/styles.xml')==z2.read('word/styles.xml'), 'Style definitions changed'
                eligible=case=='preservation-complete'
                changes=compare_word(*documents,eligible)
                row={'case':case,'language':language,'question_comparisons':count,'changed_overview_styles':changes,'eligible':eligible}
                for kind,paths in [('pdf',[old.with_suffix('.pdf'),new.with_suffix('.pdf')]),('word_preview',[a.prior/'word-preview'/(name+'.pdf'),a.build/'word-preview'/(name+'.pdf')])]:
                    assert question_body(paths[0],soups[0])==question_body(paths[1],soups[1]), (name,kind,'Question text changed')
                    row['prior_'+kind+'_pages']=page_bounds(paths[0]); row[kind+'_pages']=page_bounds(paths[1])
                    assert row[kind+'_pages']==row['prior_'+kind+'_pages'],(name,kind,'Unexpected page count change')
                    row[kind+'_q13_pages']=check_followup_page_text(subprocess.check_output(['pdftotext','-layout',str(paths[1]),'-'],text=True),soups[1])
                preview=a.build/'word-preview'/(name+'.pdf'); inspect_preview(preview)
                row['prior_q15_pagination']=overview_pages(a.prior/'word-preview'/(name+'.pdf'),documents[0])
                row['q15_pagination']=overview_pages(preview,documents[1])
                if eligible: assert len(row['q15_pagination']['all_q15_on_one_page'])==1, 'Small Q15 must include its budget on one page'
                else: assert not row['q15_pagination']['all_q15_on_one_page'], 'Long/many budget must remain splittable'
                if case=='budget-long':
                    cells='\n'.join(c.text for t in documents[1].tables for r in t.rows for c in r.cells)
                    for i in range(1,61): assert cells.count(f'BUDGET-PARA-{i:02d}:')==1
                    assert 'Budget-2027-12-31.csv' in cells
                if case=='budget-many': assert len(documents[1].tables[-1].rows)==9
                row['passed']=True; report['rows'].append(row)
    except Exception as e:
        report['failure']=str(e); target.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n'); raise
    report['selected_checks_passed']=True; target.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'passed':True,'pairs':len(report['rows']),'question_comparisons':sum(r['question_comparisons'] for r in report['rows'])}))


if __name__=='__main__': main()
