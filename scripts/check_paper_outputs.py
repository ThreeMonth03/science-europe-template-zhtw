"""Reference relocation, exact scalar values and native links; not full acceptance."""
import argparse
from collections import Counter
import copy
import json
import re
from pathlib import Path
import subprocess
import sys
from xml.etree import ElementTree as ET
from bs4 import BeautifulSoup
from docx import Document
from check_context_outputs import check_flow, compare_html as unchanged_html, signature
from check_preservation_outputs import inspect as preservation_inspect
from check_narrative_outputs import compact, sha, pdf_text
from compare_runtime_outputs import markers

LABELS = {'english': 'Related paper:', 'chinese': '相關論文：'}


def expected_papers(replies, ids):
    data = '.'.join(ids[n] for n in ['preservingCUuid', 'producedDataQUuid'])
    result = {}
    for item in replies.get(data, []):
        stage = '.'.join([data, item, ids['producedDataStageQUuid']])
        value = replies.get('.'.join([stage, ids['producedDataStagePublishedAUuid'], ids['producedDataPaperQUuid']]), '')
        if replies.get(stage) == ids['producedDataStagePublishedAUuid'] and value.strip(): result[item] = value
    return result


def check_paper_flow(soup):
    # Preserve the previous context guard, except the separately checked reference.
    other = copy.deepcopy(soup)
    for node in other.select('[data-fact-id="preservation-related-paper"]'):
        assert 'paper-reference' in node.get('class', [])
        assert node.parent in other.select('#q-data-preservation .dataset-section')
        assert 'preservation-summary' in node.find_previous_sibling().get('class', [])
        assert len(node.select('p')) == 1 and len(node.select('.paper-reference-value')) == 1
        assert not node.find_parent(class_='dataset-policy')
        node.decompose()
    return check_flow(other)


def check_values(soup, expected, language):
    check_paper_flow(soup)
    nodes = soup.select('#q-data-preservation .paper-reference')
    assert len(nodes) == len(expected)
    for node in nodes:
        item = node.parent['data-item-id']; raw = expected[item]
        assert node.get('data-status') == 'complete'
        assert node.select_one('.paper-reference-label').get_text() == LABELS[language]
        value = node.select_one('.paper-reference-value')
        assert value.get_text() == raw, ('Reference value altered', item)
        assert compact(node.get_text()) == compact(LABELS[language] + raw), 'Template punctuation added'
        assert not value.select('script, style, img, br, wbr'), 'Scalar reference parsed as markup'
        if value.a:
            assert len(value.select('a')) == 1 and value.a['href'] == raw and value.a.get_text() == raw
            assert raw.startswith(('http://', 'https://'))
    return nodes


def compare_prior(old, new, expected, language):
    current = copy.deepcopy(new)
    nodes = check_values(current, expected, language)
    assert len(old.select('[data-fact-id="preservation-related-paper"]')) == len(nodes)
    for node in nodes:
        item = node.parent['data-item-id']; raw = expected[item]
        old_dataset = old.select_one(f'#q-data-preservation .dataset-section[data-item-id="{item}"]')
        prior = old_dataset.select_one('[data-fact-id="preservation-related-paper"]')
        assert prior and prior.name == 'p' and prior.parent.get('class') == ['preservation-summary', 'dataset-policy']
        punctuation = '.' if language == 'english' else '。'
        assert compact(prior.get_text()) == compact(LABELS[language] + raw + punctuation), 'Unexpected prior reference text'
        dataset = node.parent; node.decompose()
        stage = dataset.select_one('[data-fact-id="preservation-data-stage"]')
        assert stage is not None
        stage.insert_after(copy.deepcopy(prior))
    return unchanged_html(old, current)


def word_stream(document, soup):
    heading = compact(soup.select_one('.question h3').get_text())
    rows = [signature(p) for p in document.paragraphs]
    return rows[next(i for i, p in enumerate(rows) if p[0] == heading):]


def compare_word(old_path, new_path, old_soup, new_soup, expected, language):
    old, new = Document(old_path), Document(new_path)
    rows = word_stream(old, old_soup)
    cursor = 0
    for dataset in old_soup.select('#q-data-preservation .dataset-section'):
        item = dataset['data-item-id']
        if item not in expected: continue
        paper = dataset.select_one('[data-fact-id="preservation-related-paper"]')
        owned = compact(paper.get_text())
        index = next(i for i in range(cursor, len(rows)) if owned in rows[i][0])
        assert rows[index][0].count(owned) == 1
        rows[index] = (rows[index][0].replace(owned, '', 1), *rows[index][1:])
        # Other authored/gap blocks can interrupt the policy. Insert after its
        # final actual paragraph, not immediately after the stage's paragraph.
        policy = dataset.select_one('.preservation-summary')
        tail = compact(policy.find_all(['p', 'li'])[-1].get_text())
        end = next(i for i in range(index, len(rows)) if tail in rows[i][0])
        rows.insert(end + 1, (compact(LABELS[language] + expected[item]), 'Body Text', None, None))
        cursor = end + 2
    assert rows == word_stream(new, new_soup), 'Unexpected Word paragraph/text/style change'
    tables = lambda doc: [[[(c.text, [signature(p) for p in c.paragraphs]) for c in r.cells] for r in t.rows] for t in doc.tables]
    assert tables(old) == tables(new), 'Word tables changed'


def native_links(base, nodes, preview=None):
    word = Document(base.with_suffix('.docx'))
    w = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    r = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
    links = []
    for p in word.paragraphs:
        for link in p._p.findall(w + 'hyperlink'):
            key = link.get(r + 'id')
            if key and word.part.rels[key].is_external:
                links.append((word.part.rels[key].target_ref, ''.join(n.text or '' for n in link.iter(w + 't'))))
    expected = Counter((n.a['href'], n.a.get_text()) for n in nodes if n.a)
    actual = Counter(v for v in links if v[0] in {k[0] for k in expected})
    assert actual == expected, ('Native Word links differ', actual, expected)
    pdfs = [base.with_suffix('.pdf')]
    if preview is not None:
        assert preview.exists(), 'Word preview required for this review'
        pdfs.append(preview)
    for pdf in pdfs: check_pdf_references(pdf, nodes, allow_label_overlap=(pdf == preview))
    return sum(expected.values())


def assert_link_text(fragments, value, count, label, allow_label_overlap=False):
    actual, raw = compact(''.join(fragments)), compact(value)
    if not allow_label_overlap:
        assert actual == raw * count, 'PDF linked text changed'
        return
    # LibreOffice's PDF annotation rectangle can overlap the adjacent CJK label;
    # pdftohtml then attributes e.g. "文：" to the link. Only a suffix of this
    # exact label is allowed, never arbitrary text or changes inside the value.
    label = compact(label)
    prefixes = '|'.join(re.escape(label[i:]) for i in range(len(label)))
    assert re.fullmatch(f'(?:(?:{prefixes})?{re.escape(raw)}){{{count}}}', actual), 'Unexpected PDF label overlap or changed linked value'


def check_pdf_references(pdf, nodes, allow_label_overlap=False):
    expected = Counter((n.a['href'], n.a.get_text()) for n in nodes if n.a)
    values = [compact(n.select_one('.paper-reference-value').get_text()) for n in nodes]
    text = compact(pdf_text(pdf))
    assert all(text.count(value) >= count for value, count in Counter(values).items()), ('Reference missing from PDF/Word preview', pdf)
    xml = ET.fromstring(subprocess.check_output(['pdftohtml', '-xml', '-hidden', '-stdout', str(pdf)], stderr=subprocess.DEVNULL))
    for (href, text), count in expected.items():
        fragments = [''.join(a.itertext()) for a in xml.iter('a') if a.get('href') == href]
        label = next(n.select_one('.paper-reference-label').get_text() for n in nodes if n.a and n.a['href'] == href)
        assert_link_text(fragments, text, count, label, allow_label_overlap)
    # Neither a physical-page overrun nor a body-margin overrun is acceptable.
    bbox = ET.fromstring(subprocess.check_output(['pdftotext', '-bbox', str(pdf), '-']))
    matched = 0
    for page in bbox.findall('.//{*}page'):
        for word in page.findall('.//{*}word'):
            text = compact(word.text or '')
            if len(text) >= 12 and any(text in value for value in values):
                assert float(word.get('xMin')) >= 40 and float(word.get('xMax')) <= float(page.get('width')) - 40, ('Reference crosses body margins', text)
                matched += 1
    assert matched or not values


def main():
    p = argparse.ArgumentParser(description=__doc__)
    for name in ['build', 'prior', 'english']: p.add_argument('--' + name, type=Path, required=True)
    p.add_argument('--cases', nargs='+', required=True)
    a = p.parse_args(); sys.path.insert(0, str(a.english.resolve() / 'scripts'))
    from generate_pilot_fixtures import IDS
    report = {'selected_checks_passed': False, 'release_acceptance': False, 'rows': [],
              'checker_sha256': sha(Path(__file__)),
              'helper_sha256': {n: sha(Path(__file__).with_name(n)) for n in ['check_context_outputs.py', 'check_preservation_outputs.py', 'check_narrative_outputs.py', 'compare_runtime_outputs.py']},
              'package_sha256': {n: sha(a.build / n) for n in ['english.zip', 'chinese.zip']},
              'prior_package_sha256': {n: sha(a.prior / n) for n in ['english.zip', 'chinese.zip']},
              'artifact_sha256': {str(f.relative_to(a.build)): sha(f) for folder in ['renders', 'word-preview'] for f in sorted((a.build / folder).glob('*')) if f.is_file()},
              'prior_artifact_sha256': {str(f.relative_to(a.prior)): sha(f) for folder in ['renders', 'word-preview'] for f in sorted((a.prior / folder).glob('*')) if f.is_file()},
              'limits': ['Synthetic selected cases only', 'Only the related-paper node is relocated and loses its template-owned trailing period', 'Conservative HTTP(S) linking, not general URL validation', 'LibreOffice previews are not Microsoft Word acceptance; PDF link-hitbox overlap allows only an exact adjacent-label suffix', 'Stock Markdown tables remain blocked']}
    target = a.build / 'paper-report.json'
    target.write_text(json.dumps(report, indent=2) + '\n')
    try:
        for case in a.cases:
            pair = []
            for language, locale in [('english', 'en'), ('chinese', 'zh-Hant')]:
                base = a.build / 'renders' / f'{case}-{language}'
                old_base = a.prior / 'renders' / base.name
                for fmt in ['html', 'pdf', 'docx']:
                    fixtures = [json.loads(p.with_suffix('.' + fmt + '.fixture.json').read_text()) for p in [old_base, base]]
                    for key in ['recipe_sha256', 'events_sha256', 'km_sha256']: assert fixtures[0][key] == fixtures[1][key]
                replies = {e['path']: e['value']['value'] for e in json.loads((a.english / 'fixtures/pilot' / locale / (case + '.events.json')).read_text())}
                expected = expected_papers(replies, IDS)
                soup, row = preservation_inspect(a.build, None, case, language)
                nodes = check_values(soup, expected, language)
                old = BeautifulSoup(old_base.with_suffix('.html').read_text(), 'html.parser')
                row['controlled_question_comparisons'] = compare_prior(old, soup, expected, language)
                compare_word(old_base.with_suffix('.docx'), base.with_suffix('.docx'), old, soup, expected, language)
                row['standalone_references'] = len(nodes)
                row['native_http_links'] = native_links(base, nodes, a.build/'word-preview'/(base.name+'.pdf'))
                row['word_preview_reference_text_links_margins_passed'] = True
                report['rows'].append(row); pair.append(soup)
            assert markers(pair[0]) == markers(pair[1])
    except Exception as e:
        report['failure'] = str(e); target.write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n'); raise
    report['selected_checks_passed'] = True
    target.write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps({'passed': True, 'pairs': len(report['rows']), 'comparisons': sum(r['controlled_question_comparisons'] for r in report['rows'])}))


if __name__ == '__main__': main()
