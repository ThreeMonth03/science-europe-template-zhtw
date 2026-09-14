"""Bounded quality prose / empty-section checks on actual bilingual artifacts."""
import argparse
import copy
import json
from pathlib import Path
from bs4 import BeautifulSoup
from docx import Document
from check_narrative_outputs import compact, pdf_text, page_bounds, sha

METHODS = {
    'english': ['measurement calibration', 'repeat sampling or measurement', 'standardized data capture and recording', 'data entry validation', 'data peer review', 'use of controlled vocabularies', 'measurement of samples with known outcomes to monitor consistency'],
    'chinese': ['校正量測結果', '重複取樣或量測', '以標準化方式擷取與記錄資料', '資料輸入檢核', '資料同儕審查', '使用受控詞彙', '量測已知結果的樣本以監測一致性'],
}
GAPS = {
    'english': {
        'q-copyright-ipr': 'This document does not yet describe arrangements for data ownership, intellectual property rights, or applicable legislation. Please review and complete this section.',
        'q-dm-responsible': 'This document does not yet identify a named person or team with a data management responsibility. Please check the contributor names and roles.',
    },
    'chinese': {
        'q-copyright-ipr': '本文件尚未呈現資料所有權、智慧財產權或適用法規的安排，請核對並補充本節內容。',
        'q-dm-responsible': '本文件尚未列出具名且負有資料管理職責的人員或團隊，請核對參與人員的姓名與角色。',
    },
}


def compare_unchanged(before, after, language):
    """Only Q1/Q4 are prose rewrites. Other changes must be exact new fallbacks."""
    count = 0
    assert len(before.select('.question')) == len(after.select('.question')) == 15
    for old in before.select('.question'):
        qid = old['id']; new = after.find(id=qid)
        if qid in ('q-how-data', 'q-quality-control'):
            continue
        if qid in GAPS[language] and not old.select_one('.answer').get_text(strip=True):
            assert compact(new.select_one('.answer').get_text()) == compact(GAPS[language][qid]), (language, qid)
            assert new.select_one('[data-status="missing-output"]')
        else:
            assert compact(old.get_text()) == compact(new.get_text()), (language, qid, 'unplanned change')
        count += 1
    # Q1 instrument quality used to follow the instrument list until the end of
    # its dataset li. Exclude only that old tail / the new shared quality nodes.
    old, new = copy.deepcopy(before.find(id='q-how-data')), copy.deepcopy(after.find(id='q-how-data'))
    old_items = old.select('.answer > ul:first-of-type > li[data-item-id]')
    for item in old_items:
        heading = item.find('h5', recursive=False)
        instruments = item.find('ul', recursive=False)
        if not heading or not instruments: continue
        for sibling in list(instruments.next_siblings): sibling.extract()
    for element in list(new.select('[data-fact-id^="quality-"]')): element.decompose()
    assert compact(old.get_text()) == compact(new.get_text()), (language, 'Q1 outside instrument-quality tail')
    return count + 1


def check_summary(text, language, methods, name):
    expected = (f'Quality control measures for {name}: ' + ', '.join(methods) + '.' if language == 'english'
                else f'{name}：品質管控措施包括' + '、'.join(methods) + '。')
    assert compact(text) == compact(expected), (language, text, expected)


def inspect(build, case, language):
    base = build / 'renders' / f'{case}-{language}'
    soup = BeautifulSoup(base.with_suffix('.html').read_text(), 'html.parser')
    word = Document(base.with_suffix('.docx'))
    texts = {'html': soup.get_text(), 'pdf': pdf_text(base.with_suffix('.pdf')), 'docx': '\n'.join(p.text for p in word.paragraphs)}
    q1, q4 = soup.find(id='q-how-data'), soup.find(id='q-quality-control')
    assert len(soup.select('.question')) == 15
    assert not q4.select('ul:not(.answer-detail ul), ol:not(.answer-detail ol)')
    assert not q4.select('p p, p ul, p div, p table')
    summaries = q4.select('.quality-summary')
    for summary in summaries:
        assert not summary.select('br, ul, p, div')
        assert compact(summary.get_text()) in compact(q1.get_text())
        for fmt, text in texts.items():
            assert compact(summary.get_text()) in compact(text), (case, language, fmt, 'summary lost')
        matches = [p for p in word.paragraphs if compact(p.text) == compact(summary.get_text())]
        assert len(matches) == 2, (case, language, 'Q1/Q4 must each be one Word paragraph')
    for qid in ('q-copyright-ipr', 'q-dm-responsible'):
        assert soup.find(id=qid).select_one('.answer').get_text(strip=True), (case, language, qid)
        if case in ('empty', 'archive-only'):
            assert soup.find(id=qid).select_one('[data-status="missing-output"]')
            for fmt, text in texts.items():
                assert compact(GAPS[language][qid]) in compact(text), (case, language, fmt, qid)
    if case in ('structured', 'storage-sharing', 'storage-sharing-partial', 'narrative-long'):
        name = 'Coastal temperature measurements' if language == 'english' else '沿岸水溫觀測資料'
        assert len(summaries) == 1
        check_summary(summaries[0].get_text(), language, [METHODS[language][0], METHODS[language][3]], name)
    if case == 'quality-rich':
        assert len(summaries) == 2
        check_summary(summaries[0].get_text(), language, METHODS[language], 'Coastal temperature measurements' if language == 'english' else '沿岸水溫觀測資料')
        check_summary(summaries[1].get_text(), language, [METHODS[language][3]], 'Salinity observations' if language == 'english' else '鹽度觀測資料')
        detail = q4.select_one('[data-fact-id="quality-other"]')
        assert len(detail.select('p')) == 22 and len(detail.select('ul > li')) == 2
        sentence = 'Extended quality procedure remains separate.' if language == 'english' else '延伸品質管控程序須保留獨立段落。'
        for fmt, text in texts.items():
            assert text.count(sentence) == 40, (language, fmt, 'Q1 and Q4 each retain 20 authored paragraphs')
            assert 'v1.2' in text and '0.05' in text, (language, fmt, 'authored punctuation changed')
        assert sum(sentence in p.text for p in word.paragraphs) == 40
        assert not detail.find_parent(class_='short-reading-unit')
    if case == 'quality-partial':
        assert not summaries
        for fact in ('quality-control', 'quality-methods', 'quality-other'):
            marker = q4.select_one(f'[data-fact-id="{fact}"][data-status="missing"]')
            assert marker
            for fmt, text in texts.items(): assert compact(marker.get_text()) in compact(text), (language, fmt, fact)
        assert not q4.select('[data-status="explicit-no"]')
    return {'case': case, 'language': language, 'pages': page_bounds(base.with_suffix('.pdf')), 'summary_paragraphs_per_question': len(summaries), 'passed': True}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--build', type=Path, required=True)
    parser.add_argument('--prior', type=Path)
    parser.add_argument('--cases', nargs='+', required=True)
    args = parser.parse_args(); rows = []; comparisons = 0
    for case in args.cases:
        for language in ('english', 'chinese'):
            rows.append(inspect(args.build, case, language))
            if args.prior:
                old = args.prior / 'renders' / f'{case}-{language}.html'
                if old.exists():
                    before, after = [BeautifulSoup(p.read_text(), 'html.parser') for p in (old, args.build / 'renders' / old.name)]
                    # Fixture identities must be the same, not merely case names.
                    for key in ('events_sha256', 'recipe_sha256', 'km_sha256'):
                        a = json.loads(old.with_suffix('.html.fixture.json').read_text())
                        b = json.loads((args.build / 'renders' / old.with_suffix('.html.fixture.json').name).read_text())
                        assert key in a and a[key] == b[key], (case, language, key)
                    comparisons += compare_unchanged(before, after, language)
    report = {'selected_checks_passed': True, 'release_acceptance': False, 'checker_sha256': sha(Path(__file__)),
              'rows': rows, 'controlled_question_comparisons': comparisons,
              'package_sha256': {name: sha(args.build / name) for name in ('english.zip', 'chinese.zip')},
              'render_sha256': {p.name: sha(p) for p in sorted((args.build / 'renders').iterdir()) if p.is_file()},
              'limits': ['Selected synthetic cases only; not full Science Europe answer coverage', 'Q8/Q14 output-empty fallbacks do not establish topic completeness', 'Word paragraphs checked, Microsoft Word pagination not tested', 'Quality text remains in Q1 and Q4; duplication deliberately not hidden']}
    (args.build / 'quality-report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps({'selected_checks_passed': True, 'pairs': len(rows), 'comparisons': comparisons}))


if __name__ == '__main__': main()
