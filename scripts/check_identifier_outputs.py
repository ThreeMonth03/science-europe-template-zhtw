"""0.3.13 -> 0.3.14: only Q13 owned headings and policy paragraphs may join."""
import argparse
import copy
import json
import subprocess
import sys
from pathlib import Path
from bs4 import BeautifulSoup
from docx import Document
from check_answer_state_outputs import canonical
from check_context_outputs import fact_scope, joined_paragraphs, signature
from check_narrative_outputs import compact, sha, page_bounds
from check_paper_outputs import expected_papers, check_values, native_links
from check_preservation_outputs import inspect as preservation_inspect
from check_word_rhythm_outputs import inspect_preview, assert_styles
from check_repository_outputs import scoped_pages
from compare_runtime_outputs import markers

Q13 = '#q-persistent-identifier'


def check_units(soup):
    groups = []
    for distro in soup.select(Q13 + ' .distribution-section'):
        heading = distro.select_one('.identifier-heading')
        assert heading is not None and heading.parent is distro
        children = heading.find_all(recursive=False)
        assert 1 <= len(children) <= 2 and all(p.name == 'p' and p.get('class') == ['answer-lead'] for p in children)
        assert not heading.select('.answer-detail, .data-gap, ul, table')
        if compact(heading.get_text()): groups.append(compact(heading.get_text()))
        policy = distro.select_one('.identifier-arrangement')
        if policy:
            assert policy.parent is distro and policy.get('class') == ['identifier-arrangement','dataset-policy']
            children = policy.find_all(recursive=False)
            assert 1 <= len(children) <= 3 and all(p.name == 'p' for p in children)
            assert children[0].get('data-fact-id') == 'persistent-identifier' and children[0].get('data-status') == 'complete'
            assert not policy.select('.answer-detail, .data-gap, ul, table')
            groups.append(compact(policy.get_text()))
    return groups


def compare_html(before, after):
    check_units(after)
    current = copy.deepcopy(after)
    # Restore ONLY the two declared container edits, then compare full Q13 DOM.
    for heading in current.select(Q13 + ' .identifier-heading'): heading.unwrap()
    for policy in current.select(Q13 + ' .identifier-arrangement'):
        first = policy.find('p', recursive=False)
        policy.insert_before(first.extract())
    assert markers(before) == markers(current) and fact_scope(before) == fact_scope(current)
    assert len(before.select('.question')) == len(current.select('.question')) == 15
    for old in before.select('.question'):
        new = current.find(id=old['id'])
        assert new and canonical(old) == canonical(new), (old['id'],'Unapproved question DOM/text change')
        assert [n.decode_contents() for n in old.select('.answer-detail')] == [n.decode_contents() for n in new.select('.answer-detail')]
        if old['id'] != Q13[1:]: assert str(old) == str(new), (old['id'],'Unrelated HTML changed')
    return 15


def split_word(document):
    rows = [signature(p) for p in document.paragraphs]
    start = next(i for i,p in enumerate(rows) if p[0].startswith('1.') and p[1]=='Heading 3')
    q13 = next(i for i,p in enumerate(rows) if p[0].startswith('13.') and p[1]=='Heading 3')
    q14 = next(i for i,p in enumerate(rows) if p[0].startswith('14.') and p[1]=='Heading 3')
    return rows[start:q13], rows[q13:q14], rows[q14:]


def compare_word(old, new, soup):
    before, after = map(split_word, [old,new])
    assert before[0] == after[0] and before[2] == after[2], 'Word changed outside Q13'
    expected, reductions = joined_paragraphs(before[1], check_units(soup))
    assert expected == after[1], 'Unexpected Q13 Word text, style or paragraph edit'
    cells = lambda doc: [[[(c.text,[signature(p) for p in c.paragraphs]) for c in r.cells] for r in t.rows] for t in doc.tables]
    assert cells(old) == cells(new), 'Native Word tables changed'
    for group in check_units(soup):
        para = next(p for p in new.paragraphs if compact(p.text)==group)
        if para.style.name == 'Pilot Label':
            assert compact(''.join(r.text for r in para.runs if r.bold)) == group, 'Q13 heading emphasis lost'
    return reductions


def body_pages(pdf):
    # Q13 text changes its line/page position, not its wording. Remove ONLY the
    # exact known numeric footer line; do not normalize authored punctuation.
    import re
    text = subprocess.check_output(['pdftotext','-layout',str(pdf),'-'],text=True)
    pages = text.split('\f')
    if not pages[-1].strip(): pages.pop()
    result = []
    for number,page in enumerate(pages,1):
        lines = page.splitlines()
        while lines and not lines[-1].strip(): lines.pop()
        if lines and re.fullmatch(r'\s*\d+/\d+\s*',lines[-1]):
            assert lines[-1].strip()==f'{number}/{len(pages)}', 'Unexpected numeric footer'
            lines.pop()
        result.append(compact('\n'.join(lines)))
    value=''.join(result)
    # Provenance version changes on the cover. Compare from the first numbered
    # question; its exact prefix is found independently in both languages.
    return value


def check_heading_pages(pdf, soup):
    text = subprocess.check_output(['pdftotext','-layout',str(pdf),'-'],text=True)
    start = compact(soup.select_one(Q13+' h3').get_text())[:16]
    end = compact(soup.select_one('#q-dm-responsible h3').get_text())[:16]
    pages = scoped_pages(text,start,end)
    result = []
    for heading in soup.select(Q13+' .identifier-heading'):
        label = compact(heading.get_text())
        if not label: continue
        following = heading.find_next_sibling()
        paragraph = following if following.name=='p' else following.find('p',recursive=False)
        assert paragraph is not None
        sentence = compact(paragraph.get_text())
        matched = [number for number,value in pages.items() if label in value and sentence in value]
        assert len(matched)==1, (pdf,label,'Heading detached from its first answer sentence')
        result.append({'distribution':heading.parent['data-item-id'],'page':matched[0]})
    return result


def main():
    p=argparse.ArgumentParser(description=__doc__)
    for n in ['build','prior','english']: p.add_argument('--'+n,type=Path,required=True)
    p.add_argument('--cases',nargs='+',required=True); a=p.parse_args()
    sys.path.insert(0,str(a.english.resolve()/'scripts'))
    from generate_pilot_fixtures import IDS
    report={'selected_checks_passed':False,'release_acceptance':False,'rows':[],
            'checker_sha256':sha(Path(__file__)),
            'helper_sha256':{n:sha(Path(__file__).with_name(n)) for n in ['check_answer_state_outputs.py','check_context_outputs.py','check_narrative_outputs.py','check_paper_outputs.py','check_preservation_outputs.py','check_word_rhythm_outputs.py','check_repository_outputs.py','compare_runtime_outputs.py']},
            'package_sha256':{n:sha(a.build/n) for n in ['english.zip','chinese.zip']},
            'prior_package_sha256':{n:sha(a.prior/n) for n in ['english.zip','chinese.zip']},
            'artifact_sha256':{str(f.relative_to(a.build)):sha(f) for folder in ['renders','word-preview'] for f in sorted((a.build/folder).glob('*')) if f.is_file()},
            'prior_artifact_sha256':{str(f.relative_to(a.prior)):sha(f) for folder in ['renders','word-preview'] for f in sorted((a.prior/folder).glob('*')) if f.is_file()},
            'limits':['Only selected synthetic cases, not full DMP coverage','Only declared Q13 wrappers and adjacent owned paragraphs may change','No smaller fonts or changed table/cell content','LibreOffice is not Microsoft Word acceptance','Stock Markdown tables still fail']}
    target=a.build/'identifier-report.json'; target.write_text(json.dumps(report,indent=2)+'\n')
    try:
        for case in a.cases:
            pair=[]
            for language,locale in [('english','en'),('chinese','zh-Hant')]:
                name=f'{case}-{language}'; base=a.build/'renders'/name; prior=a.prior/'renders'/name
                for fmt in ['html','pdf','docx']:
                    rows=[json.loads(f.with_suffix('.'+fmt+'.fixture.json').read_text()) for f in [prior,base]]
                    for key in ['recipe_sha256','events_sha256','km_sha256']: assert rows[0][key]==rows[1][key]
                    assert rows[0]['package_sha256']==report['prior_package_sha256'][language+'.zip']
                    assert rows[1]['package_sha256']==report['package_sha256'][language+'.zip']
                soup,row=preservation_inspect(a.build,None,case,language)
                old=BeautifulSoup(prior.with_suffix('.html').read_text(),'html.parser')
                row['controlled_question_comparisons']=compare_html(old,soup)
                old_word,new_word=[Document(f.with_suffix('.docx')) for f in [prior,base]]
                assert_styles(new_word)
                row['owned_paragraphs_joined_away']=compare_word(old_word,new_word,soup)
                replies={e['path']:e['value']['value'] for e in json.loads((a.english/'fixtures/pilot'/locale/(case+'.events.json')).read_text())}
                nodes=check_values(soup,expected_papers(replies,IDS),language)
                preview=a.build/'word-preview'/(name+'.pdf'); old_preview=a.prior/'word-preview'/(name+'.pdf')
                row['preserved_paper_links']=native_links(base,nodes,preview)
                start=compact(soup.select_one('.question h3').get_text())
                for label,paths in [('pdf',[prior.with_suffix('.pdf'),base.with_suffix('.pdf')]),('word_preview',[old_preview,preview])]:
                    texts=[body_pages(f) for f in paths]
                    assert all(t.count(start)==1 for t in texts)
                    assert texts[0].split(start,1)[1]==texts[1].split(start,1)[1], (name,label,'Body text changed')
                    row['prior_'+label+'_pages']=page_bounds(paths[0])
                    row[label+'_pages']=page_bounds(paths[1])
                    row[label+'_heading_with_answer_pages']=check_heading_pages(paths[1],soup)
                inspect_preview(preview)
                report['rows'].append(row); pair.append(soup)
            assert markers(pair[0])==markers(pair[1])
    except Exception as e:
        report['failure']=str(e); target.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n'); raise
    report['selected_checks_passed']=True
    target.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'passed':True,'pairs':len(report['rows']),'comparisons':sum(r['controlled_question_comparisons'] for r in report['rows'])}))


if __name__=='__main__': main()
