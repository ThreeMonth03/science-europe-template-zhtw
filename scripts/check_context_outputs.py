"""0.3.11 -> 0.3.12: same facts/text; only adjacent Q11 owned paragraphs may join."""
import argparse
from collections import Counter
import json
from pathlib import Path
from bs4 import BeautifulSoup
from docx import Document
from check_narrative_outputs import compact, sha
from check_preservation_outputs import inspect as preservation_inspect
from check_sharing_outputs import direct_runs
from compare_runtime_outputs import markers


def fact_scope(soup):
    return [(n.get('data-fact-id'), [p['data-item-id'] for p in n.parents if p.get('data-item-id')])
            for n in soup.select('[data-fact-id]')]


def compare_html(old, new):
    assert markers(old) == markers(new) and fact_scope(old) == fact_scope(new), 'Fact/state/scope changed'
    assert len(old.select('.question')) == len(new.select('.question')) == 15
    for a in old.select('.question'):
        b = new.find(id=a['id'])
        assert b and compact(a.get_text()) == compact(b.get_text()), (a['id'], 'Text changed')
        assert [n.decode_contents() for n in a.select('.answer-detail')] == [n.decode_contents() for n in b.select('.answer-detail')], 'Authored blocks changed'
        assert [(n.get('href'), n.get_text()) for n in a.select('a')] == [(n.get('href'), n.get_text()) for n in b.select('a')], 'Links changed'
        if a['id'] != 'q-data-preservation':
            assert str(a) == str(b), (a['id'], 'Unrelated HTML changed')
    return 15


def check_flow(soup):
    q = soup.find(id='q-data-preservation')
    groups = []
    for dataset in q.select('.dataset-section'):
        policy = dataset.select_one('.preservation-summary.dataset-policy')
        assert policy is not None and dataset.select('.dataset-policy') == [policy]
        for node in dataset.select('[data-fact-id="preservation-data-stage"], [data-fact-id="preservation-related-paper"], [data-fact-id="preservation-dataset-description"]'):
            assert node.parent is policy, 'Context is separated from the owned policy'
        assert not policy.select('.repository-destinations, .repository-contact, .post-project-archive, p.data-gap:not(.reading-gap > p)')
        groups.extend(compact(''.join(run)) for run in direct_runs(policy))
    return groups


def joined_paragraphs(before, targets):
    """Reconstruct only whole adjacent old paragraphs matching an owned DOM run.

    Signatures contain text and style/keep properties. Unchanged author paragraphs,
    gaps, lists and every other question must still match the complete Word stream.
    """
    result = list(before)
    cursor = 0
    reductions = 0
    for target in targets:
        found = None
        for start in range(cursor, len(result)):
            combined = ''
            for end in range(start, len(result)):
                combined += result[end][0]
                if combined == target:
                    found = (start, end + 1)
                    break
                if not target.startswith(combined): break
            if found: break
        assert found, ('Owned run not present as consecutive prior Word paragraphs', target)
        start, end = found
        assert all(p[1:] == result[start][1:] for p in result[start:end]), 'Cannot merge different paragraph styles'
        reductions += end - start - 1
        result[start:end] = [(target, *result[start][1:])]
        cursor = start + 1
    return result, reductions


def signature(p):
    f = p.paragraph_format
    return (compact(p.text), p.style.name, f.keep_with_next, f.keep_together)


def compare_word(old, new, soup):
    a, b = Document(old), Document(new)
    # Exclude package provenance/cover, whose version deliberately changes.
    heading = compact(soup.select_one('.question h3').get_text())
    streams = []
    for doc in (a, b):
        rows = [signature(p) for p in doc.paragraphs]
        start = next(i for i, p in enumerate(rows) if p[0] == heading)
        streams.append(rows[start:])
    expected, reductions = joined_paragraphs(streams[0], check_flow(soup))
    assert expected == streams[1], 'Unexpected Word paragraph/text/style change'
    assert [[[(c.text, [signature(p) for p in c.paragraphs]) for c in r.cells] for r in t.rows] for t in a.tables] == [[[(c.text, [signature(p) for p in c.paragraphs]) for c in r.cells] for r in t.rows] for t in b.tables], 'Word table changed'
    return reductions


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--build', type=Path, required=True)
    p.add_argument('--prior', type=Path, required=True)
    p.add_argument('--cases', nargs='+', required=True)
    a = p.parse_args()
    report = {'selected_checks_passed': False, 'release_acceptance': False, 'rows': [],
              'checker_sha256': sha(Path(__file__)),
              'helper_sha256': {n: sha(Path(__file__).with_name(n)) for n in ['check_narrative_outputs.py', 'check_preservation_outputs.py', 'check_sharing_outputs.py', 'compare_runtime_outputs.py']},
              'package_sha256': {n: sha(a.build / n) for n in ['english.zip', 'chinese.zip']},
              'prior_package_sha256': {n: sha(a.prior / n) for n in ['english.zip', 'chinese.zip']},
              'artifact_sha256': {str(f.relative_to(a.build)): sha(f) for folder in ['renders', 'word-preview'] for f in sorted((a.build / folder).glob('*')) if f.is_file()},
              'prior_artifact_sha256': {str(f.relative_to(a.prior)): sha(f) for case in a.cases for language in ['english', 'chinese'] for f in sorted((a.prior / 'renders').glob(f'{case}-{language}.*')) if f.is_file()},
              'limits': ['Only selected synthetic cases; no full SE or whole-document visual acceptance', 'Word paragraph signatures are checked; full OOXML binary equality is not expected', 'LibreOffice preview is not Microsoft Word pagination', 'Stock Markdown table failure is not fixed by this structural change']}
    target = a.build / 'context-report.json'
    target.write_text(json.dumps(report, indent=2) + '\n')
    try:
        for case in a.cases:
            pair = []
            for language in ['english', 'chinese']:
                base = f'{case}-{language}'
                for fmt in ['html', 'pdf', 'docx']:
                    fixtures = [json.loads((root / 'renders' / f'{base}.{fmt}.fixture.json').read_text()) for root in [a.prior, a.build]]
                    for key in ['recipe_sha256', 'events_sha256', 'km_sha256']:
                        assert fixtures[0][key] == fixtures[1][key], (base, fmt, key)
                soup, row = preservation_inspect(a.build, None, case, language)
                old = BeautifulSoup((a.prior / 'renders' / (base + '.html')).read_text(), 'html.parser')
                row['controlled_question_comparisons'] = compare_html(old, soup)
                groups = check_flow(soup)
                paragraphs = Counter(compact(p.text) for p in Document(a.build / 'renders' / (base + '.docx')).paragraphs)
                assert all(paragraphs[t] >= n for t, n in Counter(groups).items()), 'Owned run fragmented in Word'
                row['owned_paragraphs_joined_away'] = compare_word(a.prior / 'renders' / (base + '.docx'), a.build / 'renders' / (base + '.docx'), soup)
                if case.startswith('preservation-'): assert row['owned_paragraphs_joined_away'] == 2
                else: assert row['owned_paragraphs_joined_away'] == 0
                report['rows'].append(row); pair.append(soup)
            assert markers(pair[0]) == markers(pair[1]), 'Bilingual state mismatch'
    except Exception as e:
        report['failure'] = str(e); target.write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n'); raise
    report['selected_checks_passed'] = True
    target.write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps({'passed': True, 'pairs': len(report['rows']), 'comparisons': sum(r['controlled_question_comparisons'] for r in report['rows'])}))


if __name__ == '__main__': main()
