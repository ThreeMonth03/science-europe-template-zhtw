"""Native 0.3.14 -> 0.3.15: exact child states, grouped notices, retained answers."""
import argparse
from collections import Counter
import copy
import json
from pathlib import Path
import subprocess
import sys
from bs4 import BeautifulSoup
from docx import Document
from identifier_followup_contract import FIELDS, check_followups
from check_identifier_outputs import check_units, split_word, body_pages
from check_repository_outputs import scoped_pages
from check_answer_state_outputs import canonical
from check_context_outputs import fact_scope, signature
from check_narrative_outputs import compact, sha, page_bounds
from check_preservation_outputs import inspect as preservation_inspect, assert_no_punctuation_only_lines
from check_paper_outputs import expected_papers, check_values, native_links
from check_word_rhythm_outputs import assert_styles, inspect_preview
from compare_runtime_outputs import markers


def compare_html(before,after,replies,ids,language):
    check_followups(after,replies,ids,language)
    current=copy.deepcopy(after)
    for wrapper in current.select('#q-persistent-identifier .identifier-followups'): wrapper.decompose()
    for unit in current.select('#q-persistent-identifier .identifier-followup-unit'): unit.unwrap()
    check_units(current)
    for n in current.select('#q-persistent-identifier [data-fact-id="identifier-assigner"], #q-persistent-identifier [data-fact-id="identifier-resolution"]'):
        assert n.name=='p' and n.get('data-status') in ['complete','explicit-no']
        for key in ['data-fact-id','data-status','data-requirement-id']: del n[key]
    assert markers(before)==markers(current) and fact_scope(before)==fact_scope(current)
    assert len(before.select('.question'))==len(current.select('.question'))==15
    for old in before.select('.question'):
        new=current.find(id=old['id'])
        assert new and canonical(old)==canonical(new),(old['id'],'Unapproved HTML/text change')
        assert [n.decode_contents() for n in old.select('.answer-detail')]==[n.decode_contents() for n in new.select('.answer-detail')]
        if old['id']!='q-persistent-identifier': assert str(old)==str(new),(old['id'],'Unrelated question changed')
    return 15


def compare_word(old,new,soup):
    before,after=map(split_word,[old,new]); expected=list(before[1]); cursor=0; inserted=0
    assert before[0]==after[0] and before[2]==after[2], 'Word changed outside Q13'
    for distro in soup.select('#q-persistent-identifier .distribution-section'):
        policy=distro.select_one('.identifier-arrangement')
        if policy is None: continue
        text=compact(policy.get_text())
        index=next(i for i in range(cursor,len(expected)) if expected[i][0]==text)
        gap=distro.select_one('.identifier-followups')
        notices=[(compact(p.get_text()),'Body Text',None,None) for p in gap.find_all('p',recursive=False)] if gap else []
        if notices:
            expected[index]=(expected[index][0],'Pilot Lead',None,None)
            notices=[(text,'Pilot Lead' if i<len(notices)-1 else style,keep,next_) for i,(text,style,keep,next_) in enumerate(notices)]
        expected[index+1:index+1]=notices; inserted+=len(notices); cursor=index+1+len(notices)
    assert expected==after[1], 'Only warning paragraphs and their bounded keep chain may change in Word'
    cells=lambda d:[[[(c.text,[signature(p) for p in c.paragraphs]) for c in r.cells] for r in t.rows] for t in d.tables]
    assert cells(old)==cells(new), 'Native Word table content/style changed'
    return inserted


def restore_body_text(value,soup,notices):
    first=compact(soup.select_one('.question h3').get_text())
    q13=compact(soup.select_one('#q-persistent-identifier h3').get_text())
    q14=compact(soup.select_one('#q-dm-responsible h3').get_text())
    assert value.count(first)==value.count(q13)==value.count(q14)==1
    body=value.split(first,1)[1]; prefix,tail=body.split(q13,1); middle,suffix=tail.split(q14,1)
    cursor=0
    for notice in notices:
        needle=compact(notice.get_text()); index=middle.find(needle,cursor)
        assert index>=0, 'New warning not present in Q13 PDF body'
        middle=middle[:index]+middle[index+len(needle):]; cursor=index
    return prefix+q13+middle+q14+suffix


def check_followup_page_text(text,soup):
    start=compact(soup.select_one('#q-persistent-identifier h3').get_text())[:16]
    end=compact(soup.select_one('#q-dm-responsible h3').get_text())[:16]
    pages=scoped_pages(text,start,end); result=[]
    for distro in soup.select('#q-persistent-identifier .distribution-section'):
        heading=distro.select_one('.identifier-heading')
        unit=distro.select_one('.identifier-followup-unit')
        if unit is None: continue
        needle=compact(heading.get_text())+compact(unit.get_text())
        matched=[number for number,value in pages.items() if needle in value]
        assert len(matched)==1, (distro['data-item-id'],'Heading, policy and child notices must share a page')
        result.append({'distribution':distro['data-item-id'],'page':matched[0],
                       'warning_paragraphs':len(unit.select('.identifier-followups p'))})
    return result


def main():
    p=argparse.ArgumentParser(description=__doc__)
    for n in ['build','prior','english']: p.add_argument('--'+n,type=Path,required=True)
    p.add_argument('--cases',nargs='+',required=True); a=p.parse_args()
    sys.path.insert(0,str(a.english.resolve()/'scripts'))
    from generate_pilot_fixtures import IDS
    helpers=['identifier_followup_contract.py','check_identifier_outputs.py','check_repository_outputs.py','check_answer_state_outputs.py','check_context_outputs.py','check_narrative_outputs.py','check_preservation_outputs.py','check_paper_outputs.py','check_word_rhythm_outputs.py','compare_runtime_outputs.py']
    report={'selected_checks_passed':False,'release_acceptance':False,'rows':[],
            'checker_sha256':sha(Path(__file__)),'helper_sha256':{n:sha(Path(__file__).with_name(n)) for n in helpers},
            'package_sha256':{n:sha(a.build/n) for n in ['english.zip','chinese.zip']},
            'prior_package_sha256':{n:sha(a.prior/n) for n in ['english.zip','chinese.zip']},
            'artifact_sha256':{str(f.relative_to(a.build)):sha(f) for folder in ['renders','word-preview'] for f in sorted((a.build/folder).glob('*')) if f.is_file()},
            'prior_artifact_sha256':{str(f.relative_to(a.prior)):sha(f) for folder in ['renders','word-preview'] for f in sorted((a.prior/folder).glob('*')) if f.is_file()},
            'limits':['Selected synthetic Common KM 2.7.0 cases, not all SE requirements','Unknown UUID paths are checked offline, not asserted to be valid native fixtures','Only Q13 identifier child notices and markers may change','LibreOffice is not Microsoft Word acceptance','Word tail page and stock Markdown table limitations remain']}
    target=a.build/'identifier-followup-report.json'; target.write_text(json.dumps(report,indent=2)+'\n')
    try:
        for case in a.cases:
            pair=[]
            for language,locale in [('english','en'),('chinese','zh-Hant')]:
                name=f'{case}-{language}'; base=a.build/'renders'/name; prior=a.prior/'renders'/name
                for fmt in ['html','pdf','docx']:
                    fixtures=[json.loads(f.with_suffix('.'+fmt+'.fixture.json').read_text()) for f in [prior,base]]
                    for key in ['recipe_sha256','events_sha256','km_sha256']: assert fixtures[0][key]==fixtures[1][key]
                    assert fixtures[0]['package_sha256']==report['prior_package_sha256'][language+'.zip']
                    assert fixtures[1]['package_sha256']==report['package_sha256'][language+'.zip']
                events=a.english/'fixtures/pilot'/locale/(case+'.events.json')
                assert sha(events)==fixtures[1]['events_sha256']
                replies={e['path']:e['value']['value'] for e in json.loads(events.read_text())}
                soup,row=preservation_inspect(a.build,None,case,language); notices=check_followups(soup,replies,IDS,language)
                old=BeautifulSoup(prior.with_suffix('.html').read_text(),'html.parser')
                row['controlled_question_comparisons']=compare_html(old,soup,replies,IDS,language)
                old_word,new_word=[Document(f.with_suffix('.docx')) for f in [prior,base]]
                assert_styles(new_word); row['warning_paragraphs_added']=compare_word(old_word,new_word,soup)
                row['child_states']=dict(Counter(n['data-status'] for n in soup.select('#q-persistent-identifier [data-fact-id="identifier-assigner"], #q-persistent-identifier [data-fact-id="identifier-resolution"]')))
                nodes=check_values(soup,expected_papers(replies,IDS),language)
                preview=a.build/'word-preview'/(name+'.pdf'); old_preview=a.prior/'word-preview'/(name+'.pdf')
                row['preserved_paper_links']=native_links(base,nodes,preview)
                for label,paths in [('pdf',[prior.with_suffix('.pdf'),base.with_suffix('.pdf')]),('word_preview',[old_preview,preview])]:
                    assert restore_body_text(body_pages(paths[0]),old,[])==restore_body_text(body_pages(paths[1]),soup,notices),(name,label,'Body changed beyond exact new warnings')
                    row['prior_'+label+'_pages']=page_bounds(paths[0]); row[label+'_pages']=page_bounds(paths[1])
                    row[label+'_heading_with_answer_pages']=check_followup_page_text(subprocess.check_output(['pdftotext','-layout',str(paths[1]),'-'],text=True),soup)
                    assert_no_punctuation_only_lines(subprocess.check_output(['pdftotext','-bbox-layout',str(paths[1]),'-']),compact(soup.select_one('#q-persistent-identifier h3').get_text())[:16],compact(soup.select_one('#q-dm-responsible h3').get_text())[:16])
                inspect_preview(preview)
                report['rows'].append(row); pair.append(soup)
            assert markers(pair[0])==markers(pair[1])
    except Exception as e:
        report['failure']=str(e); target.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n'); raise
    report['selected_checks_passed']=True; target.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'passed':True,'pairs':len(report['rows']),'comparisons':sum(r['controlled_question_comparisons'] for r in report['rows'])}))


if __name__=='__main__': main()
