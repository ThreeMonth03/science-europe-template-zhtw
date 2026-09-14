"""Q10 references and full Q11 answers in native outputs; compare only allowed edits."""
import argparse
import copy
import json
import re
import subprocess
import sys
import zipfile
from collections import Counter
from pathlib import Path
from bs4 import BeautifulSoup, NavigableString
from docx import Document
from lxml import etree
from check_repository_outputs import inspect as inspect_repository, keep_value, scoped_pages
from check_answer_state_outputs import canonical
from check_narrative_outputs import compact, pdf_text, sha
from check_word_rhythm_outputs import word_body
from compare_runtime_outputs import markers

ROOT = Path(__file__).resolve().parents[1]
LEAD = {'english': 'Other repository contact arrangements:', 'chinese': '儲存庫聯繫的其他安排：'}
MISSING = {'english': 'The other repository contact arrangements have not been described.', 'chinese': '尚待補充：儲存庫聯繫的其他安排。'}
OLD = {'english': ['We have made other arrangements: ', 'We have made other arrangements instead of contacting the repository directly: '],
       'chinese': ['我們另有安排：', '我們未直接聯絡資料儲存庫，而是另作安排：']}


def reference_text(language, position, multiple):
    if language == 'english':
        return (f'For repository contact arrangements, see Question 11, distribution {position} of this dataset.' if multiple
                else 'For repository contact arrangements, see Question 11 under this dataset.')
    return (f'儲存庫聯繫安排請見第 11 題本資料集的管道 {position}。' if multiple else '儲存庫聯繫安排請見第 11 題本資料集的說明。')


def check_references(soup, language):
    refs = soup.select('#q-share-restrictions .repository-contact-reference a')
    targets = soup.select('#q-data-preservation .repository-contact')
    assert len(refs) == len(targets)
    assert len({n['id'] for n in targets}) == len(targets), 'Duplicate target ID'
    for ref in refs:
        dest = soup.find_all(id=ref['href'][1:]); assert len(dest) == 1, 'Broken or ambiguous reference'
        target = dest[0]; assert target in targets
        source_data = ref.find_parent(class_='dataset-section'); data = target.find_parent(class_='dataset-section')
        source_dist = ref.find_parent(class_='distribution-section'); dist = target.find_parent(class_='repository-distribution')
        assert source_data['data-item-id'] == data['data-item-id'] and source_dist['data-item-id'] == dist['data-item-id'], 'Reference points to another answer'
        datasets = soup.select('#q-share-restrictions .dataset-section'); distributions = source_data.select('.distribution-section')
        n, m = datasets.index(source_data) + 1, distributions.index(source_dist) + 1
        assert target['id'] == f'repository-contact-{n}-{m}'
        assert compact(ref.get_text()) == compact(reference_text(language, m, len(distributions) > 1))
        fact = target.select_one('[data-fact-id="repository-contact-arrangements"]'); assert fact is not None
        heading = dist.select_one('.repository-contact-heading')
        assert heading is not None and 'answer-lead' in heading['class']
        if fact['data-status'] == 'complete':
            assert target.select_one('.answer-lead > p').get_text() == LEAD[language]
            assert fact.get('class') == ['answer-detail'] and fact.get_text(strip=True)
        else:
            assert fact['data-status'] == 'missing' and fact.get_text() == MISSING[language]
            assert not target.select('.answer-detail')
    return refs


def compare_prior(before, after, language):
    """Reconstruct the exact old duplication from the retained new body, not a blanket Q10/Q11 exclusion."""
    old, new = copy.deepcopy(before), copy.deepcopy(after)
    for ref in list(new.select('.repository-contact-reference')):
        target = new.find(id=ref.a['href'][1:])
        detail = target.select_one('.answer-detail'); assert detail is not None, 'Historical comparison only covers supplied Other answers'
        body = [copy.deepcopy(n) for n in detail.contents]
        unit = ref.find_parent(class_='distribution-reading-unit')
        dataset = ref.find_parent(class_='dataset-section')['data-item-id']; distro = ref.find_parent(class_='distribution-section')['data-item-id']
        prior_unit = old.select_one(f'#q-share-restrictions .dataset-section[data-item-id="{dataset}"] .distribution-section[data-item-id="{distro}"] .distribution-reading-unit')
        # Removing long repeated prose can newly enable the existing bounded Q10 keep hint.
        assert set(unit.get('class', [])) <= {'distribution-reading-unit', 'short-reading-unit'}
        unit['class'] = prior_unit['class']
        ref.replace_with(NavigableString(OLD[language][0]), *[copy.deepcopy(n) for n in body])
        target.replace_with(NavigableString(OLD[language][1]), *body)
    for heading in new.select('.repository-contact-heading'): heading.unwrap()
    old.smooth(); new.smooth()
    assert len(old.select('.question')) == len(new.select('.question')) == 15
    assert markers(old) == markers(new), 'Unexpected fact change'
    for a, b in zip(old.select('.question'), new.select('.question')):
        assert a['id'] == b['id'] and canonical(a) == canonical(b), (a['id'], 'Unexpected question or authored-content change')
    return 15


def compare_word_outside(before, after):
    previous, current = [word_body(d) for d in [before, after]]
    assert previous[1] == current[1], 'Existing table cells changed'
    def outside(rows):
        start = next(i for i, (style, text) in enumerate(rows) if style == 'Heading 3' and text.startswith('10.'))
        end = next(i for i, (style, text) in enumerate(rows) if i > start and style == 'Heading 3' and text.startswith('12.'))
        return rows[:start], rows[end:]
    assert outside(previous[0]) == outside(current[0]), 'Word prose/styles outside Q10/Q11 changed'


def native_bookmarks(docx, refs):
    with zipfile.ZipFile(docx) as z: root = etree.fromstring(z.read('word/document.xml'))
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}; w = '{' + ns['w'] + '}'
    for ref in refs:
        name = ref['href'][1:]
        starts = root.xpath('//w:bookmarkStart[@w:name=$name]', namespaces=ns, name=name)
        assert len(starts) == 1, (name, 'Word target missing/duplicated')
        links = root.xpath('//w:hyperlink[@w:anchor=$name]', namespaces=ns, name=name)
        assert len(links) == 1 and compact(''.join(links[0].itertext())) == compact(ref.get_text()), (name, 'Word internal link lost')
        assert len(root.xpath('//w:bookmarkEnd[@w:id=$id]', namespaces=ns, id=starts[0].get(w + 'id'))) == 1


def pdf_references(file, refs):
    if not refs: return {}
    destinations = subprocess.check_output(['pdfinfo', '-dests', str(file)], text=True)
    entries = re.findall(r'^\s*(\d+)\s+.*?"([^"]+)"\s*$', destinations, re.M)
    xml = subprocess.check_output(['pdftohtml', '-xml', '-i', '-stdout', str(file)])
    root = etree.fromstring(xml, parser=etree.XMLParser(resolve_entities=False, no_network=True))
    anchors = root.findall('.//a'); result = {}
    for ref in refs:
        name = ref['href'][1:]; pages = [int(p) for p, n in entries if n == name]
        assert len(pages) == 1, (name, 'PDF destination missing/duplicated')
        label = ''.join(''.join(n.itertext()) for n in anchors if n.get('href', '').endswith('#' + str(pages[0])))
        assert compact(ref.get_text()) in compact(label), (name, 'PDF reference does not link to target page')
        result[name] = pages[0]
    return result


def inspect(build, english, case, language, ids):
    soup, row = inspect_repository(build, english, case, language, ids, contact_owner='q11')
    refs = check_references(soup, language); row['contact_references'] = len(refs)
    expected_count = {'repository-long': 1, 'contact-mixed': 4}.get(case, 0)
    assert len(refs) == expected_count
    base = build / 'renders' / f'{case}-{language}'
    word = Document(base.with_suffix('.docx')); native_bookmarks(base.with_suffix('.docx'), refs)
    row['pdf_reference_target_pages'] = pdf_references(base.with_suffix('.pdf'), refs)
    for ref in refs:
        target = soup.find(id=ref['href'][1:]); heading = target.find_parent(class_='repository-distribution').select_one('.repository-contact-heading')
        matches = [p for p in word.paragraphs if compact(p.text) == compact(heading.get_text())]
        assert matches and all(keep_value(p, 'keep_with_next') for p in matches), 'Contact heading orphanable in Word'
    q10 = soup.find(id='q-share-restrictions'); q11 = soup.find(id='q-data-preservation')
    if case == 'repository-long':
        assert 'Repo-review-2027-' not in q10.get_text()
        for i in range(1, 61): assert q11.get_text().count(f'Repo-review-2027-{i:03d}.csv') == 1
    if case == 'contact-mixed':
        targets = q11.select('.repository-contact')
        assert [t.select_one('[data-fact-id]')['data-status'] for t in targets] == ['complete', 'missing', 'complete', 'complete']
        detail = targets[0].select_one('.answer-detail')
        patched = json.loads((build / 'manifest.json').read_text())['status'] == 'runtime-experiment'
        assert len(detail.find_all('p', recursive=False)) == (2 if patched else 3) and len(detail.select('ul > li')) == 2
        assert bool(detail.table) == patched, 'Preserve stock table failure; require a real table on patched worker'
        assert q11.get_text().count('Same-contact-2027.csv') == 2 and 'Same-contact-2027.csv' not in q10.get_text()
        for token in ['Contact-2027.csv', 'Contact-list-2027.csv', 'Contact-table-2027.csv']:
            assert q11.get_text().count(token) == 1 and token not in q10.get_text()
        if patched:
            assert any('Contact-table-2027.csv' in c.text for t in word.tables for r in t.rows for c in r.cells)
    for file in [base.with_suffix('.pdf'), build / 'word-preview' / (base.name + '.pdf')]:
        text = compact(pdf_text(file))
        pages = scoped_pages(pdf_text(file), compact(q11.h3.get_text())[:16], compact(soup.find(id='q-access-data').h3.get_text())[:16])
        for ref in refs:
            target = soup.find(id=ref['href'][1:]); heading = target.find_parent(class_='repository-distribution').select_one('.repository-contact-heading')
            opening = target.select_one('.answer-detail > p') or target.select_one('.data-gap')
            assert opening is not None
            lead = target.select_one('.answer-lead')
            snippets = [compact(heading.get_text()), compact(opening.get_text())] + ([compact(lead.get_text())] if lead else [])
            assert any(all(s in content for s in snippets) for content in pages.values()), 'Contact heading separated from opening'
        if case == 'repository-long':
            for i in range(1, 61): assert text.count(f'Repo-review-2027-{i:03d}.csv') == 1
        if case == 'contact-mixed':
            assert text.count('Same-contact-2027.csv') == 2
            for token in ['Contact-2027.csv', 'Contact-list-2027.csv', 'Contact-table-2027.csv']: assert text.count(token) == 1
        for ref in refs: assert compact(ref.get_text()) in text, 'Reference text lost in PDF'
    return soup, row


def main():
    p = argparse.ArgumentParser(description=__doc__)
    for n in ['build', 'english', 'prior']: p.add_argument('--' + n, type=Path, required=True)
    p.add_argument('--cases', nargs='+', required=True); p.add_argument('--new-cases', nargs='*', default=[])
    a = p.parse_args(); sys.path.insert(0, str(a.english.resolve() / 'scripts'))
    from generate_pilot_fixtures import IDS
    report = {'selected_checks_passed': False, 'release_acceptance': False, 'rows': [], 'checker_sha256': sha(Path(__file__)),
              'package_sha256': {n: sha(a.build / n) for n in ['english.zip', 'chinese.zip']},
              'artifact_sha256': {str(f.relative_to(a.build)): sha(f) for folder in ['renders', 'word-preview'] for f in sorted((a.build / folder).glob('*')) if f.is_file()},
              'helper_sha256': {n: sha(ROOT / 'scripts' / n) for n in ['check_repository_outputs.py', 'check_answer_state_outputs.py', 'check_narrative_outputs.py', 'check_word_rhythm_outputs.py', 'compare_runtime_outputs.py']},
              'prior_package_sha256': {n: sha(a.prior / n) for n in ['english.zip', 'chinese.zip']}, 'prior_artifact_sha256': {},
              'new_cases_without_prior': a.new_cases, 'limits': ['Selected synthetic contact fields, not full SE coverage', 'LibreOffice is not target Microsoft Word acceptance']}
    target = a.build / 'contact-report.json'
    try:
        for case in a.cases:
            pair = []
            for lang in ['english', 'chinese']:
                soup, row = inspect(a.build, a.english, case, lang, IDS); pair.append(soup)
                row['controlled_question_comparisons'] = 0
                if case not in a.new_cases:
                    old = a.prior / 'renders' / f'{case}-{lang}.html'; new = a.build / 'renders' / old.name
                    x, y = [json.loads(f.with_suffix('.html.fixture.json').read_text()) for f in [old, new]]
                    for key in ['recipe_sha256', 'events_sha256', 'km_sha256']: assert x[key] == y[key]
                    row['controlled_question_comparisons'] = compare_prior(BeautifulSoup(old.read_text(), 'html.parser'), soup, lang)
                    compare_word_outside(Document(old.with_suffix('.docx')), Document(new.with_suffix('.docx')))
                    row['outside_q10_q11_word_unchanged'] = True
                    for f in [old, old.with_suffix('.html.fixture.json'), old.with_suffix('.docx'), old.with_suffix('.docx.fixture.json')]: report['prior_artifact_sha256'][str(f.relative_to(a.prior))] = sha(f)
                report['rows'].append(row)
            assert markers(pair[0]) == markers(pair[1])
    except Exception as e:
        report['failure'] = str(e); target.write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n'); raise
    report['selected_checks_passed'] = True; target.write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps({'passed': True, 'pairs': len(report['rows']), 'comparisons': sum(r['controlled_question_comparisons'] for r in report['rows'])}))


if __name__ == '__main__': main()
