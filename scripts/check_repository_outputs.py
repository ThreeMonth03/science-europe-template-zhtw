"""Native Q11 identity, bounded pagination, authored paragraphs and 0.3.9 comparison."""
import argparse
import copy
import json
import re
import subprocess
import sys
from collections import Counter
from pathlib import Path
from bs4 import BeautifulSoup
from docx import Document
from check_answer_state_outputs import inspect as inspect_states, canonical
from check_narrative_outputs import compact, sha
from compare_runtime_outputs import markers
from check_word_rhythm_outputs import word_body

OLD_LEAD = {'english': 'The distributions will be stored in:', 'chinese': '資料的發布版本將存放於：'}
LEAD = {'english': 'Preservation destinations by distribution:', 'chinese': '各資料提供管道的保存位置：'}


def compare_prior(before, after, language):
    old, new = copy.deepcopy(before), copy.deepcopy(after)
    assert len(old.select('.question')) == len(new.select('.question')) == 15
    for unit in new.select('#q-data-preservation .repository-destinations'):
        p = unit.select_one('.answer-lead > p'); assert p.get_text() == LEAD[language]
        p.string = OLD_LEAD[language]
        rows = unit.select('.repository-distribution')
        for i, row in enumerate(rows, 1):
            labels = row.select('.repository-label')
            assert len(labels) == (1 if len(rows) > 1 else 0)
            for n in labels:
                assert n.get_text() == (f'Distribution {i}:' if language == 'english' else f'資料提供管道 {i}：')
                n.decompose()
        unit.unwrap()
    old.smooth(); new.smooth()
    assert markers(old) == markers(new), 'Unexpected fact change'
    for a, b in zip(old.select('.question'), new.select('.question')):
        assert a['id'] == b['id'] and canonical(a) == canonical(b), (a['id'], 'Unexpected question change')
        assert [str(n) for n in a.select('.answer-detail')] == [str(n) for n in b.select('.answer-detail')]
    return 15


def scoped_pages(text, start, end):
    """Return numbered page fragments only within Q11, excluding Q10 labels."""
    pages = [compact(p) for p in text.split('\f')]
    assert sum(p.count(start) for p in pages) == sum(p.count(end) for p in pages) == 1
    active = False; result = {}
    for number, p in enumerate(pages, 1):
        if start in p: active = True; p = p.split(start, 1)[1]
        if active:
            result[number] = p.split(end, 1)[0]
            if end in p: break
    return result


def keep_value(paragraph, property_name):
    value = getattr(paragraph.paragraph_format, property_name)
    style = paragraph.style
    while value is None and style is not None:
        value = getattr(style.paragraph_format, property_name); style = style.base_style
    return bool(value)


def compare_word_outside_q11(before, after):
    previous, current = [word_body(d) for d in [before, after]]
    assert previous[1] == current[1], 'Native Word table cells changed'
    def outside(rows):
        first = next(i for i, (style, text) in enumerate(rows) if style == 'Heading 3' and text.startswith('11.'))
        last = next(i for i, (style, text) in enumerate(rows) if i > first and style == 'Heading 3' and text.startswith('12.'))
        return rows[:first], rows[last:]
    assert outside(previous[0]) == outside(current[0]), 'Native Word prose or styles outside Q11 changed'


def inspect(build, english, case, language, ids):
    soup, row = inspect_states(build, english, case, language, ids)
    q = soup.find(id='q-data-preservation'); base = build / 'renders' / f'{case}-{language}'
    document = Document(base.with_suffix('.docx')); paragraphs = document.paragraphs
    units = q.select('.repository-destinations')
    for unit in units:
        assert unit.select_one('.answer-lead > p').get_text() == LEAD[language]
        dataset = unit.find_parent(class_='dataset-section')['data-item-id']
        rows = unit.select('.repository-distribution')
        for other_id in ['q-share-restrictions', 'q-persistent-identifier']:
            other_q = soup.find(id=other_id)
            assert other_q is not None
            other = other_q.select_one(f'.dataset-section[data-item-id="{dataset}"]')
            assert [n['data-item-id'] for n in rows] == [n['data-item-id'] for n in other.select('.distribution-section')]
        for i, n in enumerate(rows, 1):
            labels = n.select('.repository-label')
            assert [l.get_text() for l in labels] == ([f'Distribution {i}:' if language == 'english' else f'資料提供管道 {i}：'] if len(rows) > 1 else [])
            if 'short-repository-list' in unit.get('class', []):
                matches = [p for p in paragraphs if compact(p.text) == compact(n.get_text())]
                assert len(matches) == 1, 'Short repository row fragmented or duplicated in Word'
                p = matches[0]
                assert p.style.name == ('Pilot Repository Lead' if i < len(rows) else 'Pilot Repository Item')
                assert keep_value(p, 'keep_together') and keep_value(p, 'keep_with_next') == (i < len(rows))
    if case == 'repository-gap':
        n = q.select('.repository-distribution')[1]
        assert n.select_one('[data-fact-id="repository-destination"][data-status="missing"]')
        assert not n.select('[data-fact-id="repository-long-term-support"]')
    start = compact(q.h3.get_text())[:16]; end = compact(soup.find(id='q-access-data').h3.get_text())[:16]
    row['short_list_pages'] = {}
    for name, file in [('pdf', base.with_suffix('.pdf')), ('word', build / 'word-preview' / (base.name + '.pdf'))]:
        pages = scoped_pages(subprocess.check_output(['pdftotext', '-layout', str(file), '-'], text=True), start, end)
        spans = []
        for unit in units:
            if 'short-repository-list' not in unit.get('class', []): continue
            texts = [compact(unit.select_one('.answer-lead').get_text())] + [compact(n.get_text()) for n in unit.select('.repository-distribution')]
            candidates = [p for p, text in pages.items() if all(t in text for t in texts)]
            assert len(candidates) == 1, (case, language, name, 'Short repository list split across pages')
            spans.append(candidates[0])
        row['short_list_pages'][name] = spans
        if case == 'repository-long':
            assert not q.select('.short-repository-list')
            token_pages = {p for p, text in pages.items() if 'Repo-review-2027-' in text}
            assert len(token_pages) >= 2, (name, 'Long Q11 answer did not cross pages')
            for i in range(1, 61): assert sum(t.count(f'Repo-review-2027-{i:03d}.csv') for t in pages.values()) == 1
            row[name + '_long_answer_pages'] = sorted(token_pages)
    if case == 'repository-long':
        authored = q.select('.repository-distribution')[0].find_all('p')
        assert len(authored) == 60
        available = Counter(compact(p.text) for p in paragraphs)
        for n in authored:
            assert available[compact(n.get_text())] == 2, 'Authored paragraph lost/merged (same answer is used in Q10 and Q11)'
        matches = [p for p in paragraphs if re.search(r'Repo-review-2027-\d{3}\.csv', p.text)]
        assert sum(not keep_value(p, 'keep_with_next') for p in matches) >= 118, 'Long authored answer locked into a keep chain'
        row['q11_authored_paragraphs'] = 60
    return soup, row


def main():
    p = argparse.ArgumentParser(description=__doc__)
    for n in ['build', 'english', 'prior']: p.add_argument('--' + n, type=Path, required=True)
    p.add_argument('--cases', nargs='+', required=True); p.add_argument('--new-cases', nargs='*', default=[])
    a = p.parse_args(); sys.path.insert(0, str(a.english.resolve() / 'scripts'))
    from generate_pilot_fixtures import IDS
    assert set(a.new_cases) <= set(a.cases)
    report = {'selected_checks_passed': False, 'release_acceptance': False, 'rows': [],
              'checker_sha256': sha(Path(__file__)), 'package_sha256': {n: sha(a.build / n) for n in ['english.zip', 'chinese.zip']},
              'helper_sha256': {n: sha(Path(__file__).with_name(n)) for n in ['check_answer_state_outputs.py', 'check_narrative_outputs.py', 'check_word_rhythm_outputs.py', 'compare_runtime_outputs.py']},
              'artifact_sha256': {str(f.relative_to(a.build)): sha(f) for folder in ['renders', 'word-preview'] for f in sorted((a.build / folder).glob('*')) if f.is_file()},
              'prior_package_sha256': {n: sha(a.prior / n) for n in ['english.zip', 'chinese.zip']}, 'prior_artifact_sha256': {},
              'new_cases_without_prior': a.new_cases, 'limits': ['Selected synthetic Q11 paths, not full SE coverage', 'LibreOffice is not Microsoft Word acceptance', 'Short-list geometry is not an overall beauty score']}
    target = a.build / 'repository-report.json'; target.write_text(json.dumps(report, indent=2) + '\n')
    try:
        for case in a.cases:
            pair = []
            for language in ['english', 'chinese']:
                soup, row = inspect(a.build, a.english, case, language, IDS); pair.append(soup)
                row['controlled_question_comparisons'] = 0
                if case not in a.new_cases:
                    old = a.prior / 'renders' / f'{case}-{language}.html'; new = a.build / 'renders' / old.name
                    x, y = [json.loads(f.with_suffix('.html.fixture.json').read_text()) for f in [old, new]]
                    for key in ['recipe_sha256', 'events_sha256', 'km_sha256']: assert x[key] == y[key]
                    for f in [old, old.with_suffix('.html.fixture.json')]: report['prior_artifact_sha256'][str(f.relative_to(a.prior))] = sha(f)
                    row['controlled_question_comparisons'] = compare_prior(BeautifulSoup(old.read_text(), 'html.parser'), soup, language)
                    compare_word_outside_q11(Document(old.with_suffix('.docx')), Document(new.with_suffix('.docx')))
                    for f in [old.with_suffix('.docx'), old.with_suffix('.docx.fixture.json')]: report['prior_artifact_sha256'][str(f.relative_to(a.prior))] = sha(f)
                    row['outside_q11_word_paragraphs_and_cells_unchanged'] = True
                report['rows'].append(row)
            assert markers(pair[0]) == markers(pair[1])
    except Exception as e:
        report['failure'] = str(e); target.write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n'); raise
    report['selected_checks_passed'] = True; target.write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps({'passed': True, 'pairs': len(report['rows']), 'comparisons': sum(r['controlled_question_comparisons'] for r in report['rows'])}))


if __name__ == '__main__': main()
