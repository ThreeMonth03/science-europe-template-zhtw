"""Native 0.3.16 -> 0.3.17: preserve all answers and scope long-budget layout."""
import argparse
import copy
import json
from pathlib import Path
import subprocess
import zipfile
from bs4 import BeautifulSoup
from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from artifact_utils import sha
from check_budget_outputs import body,xml,compare_word,overview_pages,question_body
from check_word_rhythm_outputs import assert_styles,compare_questions,inspect_preview
from check_narrative_outputs import compact,page_bounds


def paragraphs(node): return [xml(p) for p in node.iter(qn('w:p'))]


def empty_table_separator(node):
    assert node.tag==qn('w:p') and not node.attrib and len(node)==0 and node.text in (None,''), 'Only an empty Pandoc table separator is allowed'


def compare_long_table(old,long,tail):
    old_rows=old.findall(qn('w:tr')); rows=long.findall(qn('w:tr')); last=tail.findall(qn('w:tr'))
    assert len(old_rows)==3 and len(last)==2 and len(rows)>=14
    assert xml(old_rows[0])==xml(rows[0])==xml(last[0])
    assert xml(old_rows[2])==xml(last[1]), 'Short trailing resource changed'
    assert xml(old.find(qn('w:tblPr')))==xml(tail.find(qn('w:tblPr')))
    assert xml(old.find(qn('w:tblGrid')))==xml(long.find(qn('w:tblGrid')))==xml(tail.find(qn('w:tblGrid')))
    props=copy.deepcopy(long.find(qn('w:tblPr'))); style=props.find(qn('w:tblStyle'))
    assert style.get(qn('w:val'))=='PilotLongBudget'; style.set(qn('w:val'),'Table')
    assert xml(props)==xml(old.find(qn('w:tblPr'))), 'Only the table style may change'
    original=old_rows[1].findall(qn('w:tc')); identity=rows[1].findall(qn('w:tc'))
    assert len(original)==len(identity)==3
    assert paragraphs(original[0])[:1]==paragraphs(identity[0]), 'Resource title changed'
    for index in [1,2]: assert paragraphs(original[index])==paragraphs(identity[index]), 'Amount/funding changed'
    assert all(row.find(qn('w:trPr')+'/'+qn('w:tblHeader')) is not None for row in rows[:2])
    preserved=[]
    for row in rows[2:]:
        assert row.find(qn('w:trPr')+'/'+qn('w:tblHeader')) is None
        cells=row.findall(qn('w:tc')); assert len(cells)==1
        span=cells[0].find(qn('w:tcPr')+'/'+qn('w:gridSpan'))
        assert span is not None and span.get(qn('w:val'))=='3'
        assert 1<=len(paragraphs(cells[0]))<=8
        preserved.extend(paragraphs(cells[0]))
    assert paragraphs(original[0])[1:]==preserved, 'Purpose paragraph/run/list content changed'
    return len(preserved)


def compare_documents(old,new,long_case):
    if not long_case: return compare_word(old,new,False)
    before,after=body(old),body(new)
    assert len(after)==len(before)+2
    i=max(i for i,n in enumerate(before) if n.tag==qn('w:tbl'))
    assert all(xml(a)==xml(b) for a,b in zip(before[:i],after[:i]))
    assert [xml(n) for n in before[i+1:]]==[xml(n) for n in after[i+3:]], 'Non-budget body changed'
    assert after[i].tag==after[i+2].tag==qn('w:tbl')
    empty_table_separator(after[i+1])
    count=compare_long_table(before[i],after[i],after[i+2])
    links=lambda d:sorted(r.target_ref for r in d.part.rels.values() if r.is_external)
    assert links(old)==links(new), 'External relationship targets changed'
    return count


def styles_only_added(old,new):
    with zipfile.ZipFile(old) as a,zipfile.ZipFile(new) as b:
        left,right=[etree.fromstring(z.read('word/styles.xml')) for z in [a,b]]
    added=[n for n in right if n.get(qn('w:styleId'))=='PilotLongBudget']; assert len(added)==1
    style=added[0]; assert style.get(qn('w:type'))=='table'
    assert style.find(qn('w:basedOn')).get(qn('w:val'))=='Table'
    assert style.find('.//'+qn('w:insideH')).get(qn('w:val'))=='nil'
    base=next(n for n in left if n.get(qn('w:styleId'))=='Table')
    assert [xml(n) for n in style.findall(qn('w:tblStylePr'))]==[xml(n) for n in base.findall(qn('w:tblStylePr'))]
    right.remove(style); assert xml(left)==xml(right), 'Existing styles changed'


def page_texts(path):
    pages=subprocess.check_output(['pdftotext','-layout',str(path),'-'],text=True).split('\f')
    if not pages[-1].strip(): pages.pop()
    result=[]
    for i,page in enumerate(pages,1):
        lines=[l for l in page.splitlines() if l.strip()]
        if lines and compact(lines[-1])==f'{i}/{len(pages)}': lines.pop()
        result.append(compact('\n'.join(lines)))
    return result


def check_long_pages(pages,old_document,soup):
    table=old_document.tables[-1]; title=compact(table.cell(1,0).paragraphs[0].text)
    amount=compact(table.cell(1,1).text); funding=compact(table.cell(1,2).text)
    headings=[compact(c.text) for c in table.rows[0].cells]
    purposes=[compact(p.text) for p in table.cell(1,0).paragraphs if 'BUDGET-PARA-' in p.text]
    assert len(purposes)==60
    locations=[]
    for text in purposes:
        hits=[i for i,p in enumerate(pages,1) if text in p]
        assert len(hits)==1, ('Missing, split or duplicated purpose',text)
        page=pages[hits[0]-1]
        assert all(value in page for value in [title,amount,funding]+headings), 'Continuation lost resource identity'
        locations.extend(hits)
    budget=compact(soup.select_one('#q-required-resources h4').get_text())
    first=compact(table.cell(1,0).paragraphs[1].text)
    hits=[i for i,p in enumerate(pages,1) if budget in p and all(v in p for v in [title,amount,funding,first,purposes[0]])]
    assert len(hits)==1, 'Budget heading must start with the resource and its purpose'
    return {'purpose_pages':sorted(set(locations)),'budget_start_page':hits[0],'complete_purpose_paragraphs':60}


def main():
    p=argparse.ArgumentParser(description=__doc__)
    for n in ['build','prior']: p.add_argument('--'+n,type=Path,required=True)
    p.add_argument('--cases',nargs='+',required=True); a=p.parse_args()
    helpers=['artifact_utils.py','check_budget_outputs.py','check_word_rhythm_outputs.py','check_narrative_outputs.py','check_identifier_outputs.py','check_identifier_followup_outputs.py']
    report={'selected_checks_passed':False,'release_acceptance':False,'rows':[],
        'checker_sha256':sha(Path(__file__)),'helper_sha256':{n:sha(Path(__file__).with_name(n)) for n in helpers},
        'package_sha256':{n:sha(a.build/n) for n in ['english.zip','chinese.zip']},
        'prior_package_sha256':{n:sha(a.prior/n) for n in ['english.zip','chinese.zip']},
        'artifact_sha256':{str(f.relative_to(a.build)):sha(f) for folder in ['renders','word-preview'] for f in sorted((a.build/folder).glob('*')) if f.is_file()},
        'prior_artifact_sha256':{str(f.relative_to(a.prior)):sha(f) for folder in ['renders','word-preview'] for f in sorted((a.prior/folder).glob('*')) if f.is_file()},
        'limits':['Selected synthetic inputs, not whole DMP acceptance','LibreOffice, not Microsoft Word','Native PDF layout unchanged; long PDF remains outside this fix','Stock Markdown failure remains blocked']}
    target=a.build/'long-budget-report.json'
    try:
        for case in a.cases:
            for language in ['english','chinese']:
                name=case+'-'+language; old=a.prior/'renders'/name; new=a.build/'renders'/name
                for fmt in ['html','pdf','docx']:
                    fs=[json.loads(n.with_suffix('.'+fmt+'.fixture.json').read_text()) for n in [old,new]]
                    for key in ['recipe_sha256','events_sha256','km_sha256']: assert fs[0][key]==fs[1][key]
                    assert fs[0]['package_sha256']==report['prior_package_sha256'][language+'.zip']
                    assert fs[1]['package_sha256']==report['package_sha256'][language+'.zip']
                soups=[BeautifulSoup(n.with_suffix('.html').read_text(),'html.parser') for n in [old,new]]
                count=compare_questions(*soups); docs=[Document(n.with_suffix('.docx')) for n in [old,new]]
                assert_styles(docs[1]); styles_only_added(old.with_suffix('.docx'),new.with_suffix('.docx'))
                long_case=case=='budget-long'; preserved=compare_documents(*docs,long_case)
                assert question_body(old.with_suffix('.pdf'),soups[0])==question_body(new.with_suffix('.pdf'),soups[1])
                prior_preview=a.prior/'word-preview'/(name+'.pdf'); preview=a.build/'word-preview'/(name+'.pdf')
                pages_before,pages=[page_texts(p) for p in [prior_preview,preview]]
                q1=compact(soups[0].select_one('.question h3').get_text()); q15=compact(soups[0].select_one('#q-required-resources h3').get_text())
                bodies=[''.join(p).split(q1,1)[1].split(q15,1)[0] for p in [pages_before,pages]]
                assert bodies[0]==bodies[1], 'Q1-Q14 Word text changed'
                row={'case':case,'language':language,'question_comparisons':count,'purpose_xml_paragraphs_preserved':preserved,
                    'prior_pdf_pages':page_bounds(old.with_suffix('.pdf')),'pdf_pages':page_bounds(new.with_suffix('.pdf')),
                    'prior_word_pages':len(pages_before),'word_pages':inspect_preview(preview)}
                assert row['prior_pdf_pages']==row['pdf_pages']
                if long_case: row.update(check_long_pages(pages,docs[0],soups[0]))
                else:
                    assert pages_before==pages, 'Control preview text/pagination changed'
                    if case=='preservation-complete': assert len(overview_pages(preview,docs[1])['all_q15_on_one_page'])==1
                row['passed']=True; report['rows'].append(row)
    except Exception as e:
        report['failure']=str(e); target.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n'); raise
    report['selected_checks_passed']=True; target.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'passed':True,'pairs':len(report['rows']),'question_comparisons':sum(r['question_comparisons'] for r in report['rows'])}))


if __name__=='__main__': main()
