"""Bounded narrative/pagination checks on actual outputs; not whole-DMP acceptance."""
import argparse
import hashlib
import json
import re
import subprocess
from pathlib import Path
from xml.etree import ElementTree as ET

from bs4 import BeautifulSoup
from docx import Document

CASES = ('structured', 'storage-sharing', 'storage-sharing-partial', 'archive-only', 'narrative-long')


def compact(value):
    return re.sub(r'\s+', '', value)


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def pdf_text(path):
    text = subprocess.check_output(['pdftotext', str(path), '-'], text=True)
    return re.sub(r'(?m)^\s*\d+\s*/\s*\d+\s*$', '', text)


def page_bounds(path):
    root = ET.fromstring(subprocess.check_output(['pdftotext', '-bbox', str(path), '-']))
    pages = root.findall('.//{*}page')
    for page in pages:
        width, height = float(page.get('width')), float(page.get('height'))
        for word in page.findall('.//{*}word'):
            assert float(word.get('xMin')) >= -0.5 and float(word.get('yMin')) >= -0.5, (path, word.text)
            assert float(word.get('xMax')) <= width + 0.5 and float(word.get('yMax')) <= height + 0.5, (path, word.text)
    return len(pages)


def compare_questions(before, after, language):
    count = 0
    for old in before.select('.question'):
        qid = old['id']; new = after.find(id=qid)
        if qid == 'q-access-security':
            # Remove ONLY the known obsolete archival paragraph and new reference.
            prefix = ('We will be archiving data for long-term preservation already during our project.' if language == 'english'
                      else '我們會在專案進行期間即開始典藏資料，以供長期保存。')
            for paragraph in list(old.select('p')):
                if compact(paragraph.get_text()).startswith(compact(prefix)): paragraph.decompose()
            for reference in list(new.select('.answer-reference')): reference.decompose()
        a, b = compact(old.get_text()), compact(new.get_text())
        if qid == 'q-share-restrictions':
            original, replacement = (('Available under some restrictions, which we will follow in our project:', 'Access to this distribution is subject to restrictions.') if language == 'english'
                else ('可在部分限制下取得，我們將在專案中遵循這些限制：', '此管道的資料取用須遵守以下限制。'))
            a = a.replace(compact(original), compact(replacement))
        if qid == 'q-access-security' and language == 'chinese':
            for original, replacement in [
                ('專案成員不會隨身攜帶資料（例如存放在筆記型電腦、USB 隨身碟或其他外接媒體）。', '計畫成員不會隨身攜帶資料（例如存放於筆記型電腦、USB 隨身碟或其他外接媒體）。'),
                ('所有專案網路服務皆可透過安全的 HTTPS（https://...）取用。', '本計畫的所有網路服務皆提供 HTTPS 安全連線。'),
                ('專案成員已接受專案一般風險與特定風險的說明。', '計畫成員已接受一般性風險及本計畫特有風險的說明。')]:
                a = a.replace(compact(original), compact(replacement))
        assert a == b, (language, qid, 'unplanned text change')
        count += 1
    assert count == 15
    return count


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--build', type=Path, required=True)
    parser.add_argument('--prior', type=Path)
    args = parser.parse_args(); rows = []; comparisons = 0
    for case in CASES:
        for language in ('english', 'chinese'):
            base = args.build / 'renders' / f'{case}-{language}'
            soup = BeautifulSoup(base.with_suffix('.html').read_text(), 'html.parser')
            word = Document(base.with_suffix('.docx'))
            word_text = '\n'.join([p.text for p in word.paragraphs] + [c.text for t in word.tables for r in t.rows for c in r.cells])
            texts = {'html': soup.get_text(), 'pdf': pdf_text(base.with_suffix('.pdf')), 'docx': word_text}
            q5, q6 = soup.find(id='q-store-backup'), soup.find(id='q-access-security')
            for marker in q5.select('[data-fact-id]'):
                for fmt, text in texts.items():
                    assert compact(marker.get_text()) in compact(text), (case, language, fmt, marker['data-fact-id'])
            assert 'will be infrequently backed up' not in q6.get_text()
            if case != 'structured':
                assert q6.select_one('.answer-reference a[href="#q-store-backup"]')
                assert len(q5.select('.archive-location-policy')) == 1
                needle = ('Data will be archived in cold storage during the project.' if language == 'english'
                          else '研究期間將以冷儲存方式封存資料。')
                medium = 'The archive will be stored on tape.' if language == 'english' else '封存資料將儲存於磁帶。'
                assert any(needle in p.text and medium in p.text for p in word.paragraphs), (case, language, 'archive paragraph not joined')
            if case == 'archive-only':
                assert q6.select_one('[data-status="missing-output"]')
                no_frequent = ('Frequent backups are not required because the archived data changes infrequently.' if language == 'english'
                               else '由於封存資料變動不頻繁，因此不需要頻繁備份。')
                for fmt, text in texts.items(): assert compact(no_frequent) in compact(text), (case, language, fmt)
            elif case != 'structured':
                assert not q6.select_one('[data-status="missing-output"]')
                needle = ('There is a shared workspace used during the project for working with data.' if language == 'english'
                          else '研究期間將使用共享工作空間處理資料。')
                specialists = 'Dedicated specialists will operate the shared workspace.' if language == 'english' else '共享工作空間將由專責人員維運。'
                assert any(needle in p.text and specialists in p.text for p in word.paragraphs)
            if case != 'archive-only':
                lead = 'The distributions will be stored in:' if language == 'english' else '資料的發布版本將存放於：'
                paragraphs = [p for p in word.paragraphs if p.text.strip() == lead]
                assert paragraphs and all(p.style.name == 'Pilot Lead' and p.style.paragraph_format.keep_with_next for p in paragraphs)
                assert soup.select_one('#q-share-restrictions .distribution-reading-unit.short-reading-unit')
                access = 'Open access: this distribution will be shared with anyone.' if language == 'english' else '此管道提供的資料將公開供任何人取用。'
                paragraphs = [p for p in word.paragraphs if p.text.strip() == access]
                assert paragraphs and all(p.style.name == 'Pilot Lead' for p in paragraphs), 'Short distribution chain not applied in Word'
            if case == 'narrative-long':
                needle = 'Extended access condition remains in the document.' if language == 'english' else '延伸取用條件仍須完整保留於文件中。'
                for fmt, text in texts.items(): assert compact(text).count(compact(needle)) == 80, (case, language, fmt)
                assert sum(needle in p.text for p in word.paragraphs) == 80, 'Free paragraphs must not be joined'
                detail = next(d for d in soup.select('#q-share-restrictions .answer-detail') if needle in d.get_text())
                assert not detail.find_parent(class_='short-reading-unit')
                assert len(detail.find_all('p', recursive=False)) == 81
                assert detail.find('ul') is not None
            for q in soup.select('.question'):
                assert not q.select('p p, p div, p ul, p ol, p table'), (case, language, q['id'])
            if args.prior and case in CASES[:3]:
                old = BeautifulSoup((args.prior / 'renders' / base.with_suffix('.html').name).read_text(), 'html.parser')
                comparisons += compare_questions(old, soup, language)
            rows.append({'case': case, 'language': language, 'selected_checks_passed': True,
                         'pdf_pages': page_bounds(base.with_suffix('.pdf'))})
    report = {'selected_checks_passed': True, 'checks': rows, 'prior_question_comparisons': comparisons,
              'release_acceptance': False, 'checker_sha256': sha(Path(__file__)),
              'limits': ['Declared synthetic cases only', 'Word styles and LibreOffice previews are not Microsoft Word testing',
                         'Page-boundary text boxes do not prove absence of overlap', 'No full Science Europe content acceptance'],
              'sha256': {str(p.relative_to(args.build)): sha(p) for p in sorted((args.build / 'renders').iterdir()) if p.is_file()}}
    (args.build / 'narrative-report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps({'case_language_pairs': len(rows), 'prior_question_comparisons': comparisons, 'selected_checks_passed': True}))


if __name__ == '__main__': main()
