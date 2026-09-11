"""Selected storage/sharing facts in actual bilingual HTML/PDF/DOCX; not full SE acceptance."""
import argparse
import hashlib
import json
import re
import subprocess
from pathlib import Path
from bs4 import BeautifulSoup
from docx import Document


def compact(value):
    return re.sub(r'\s+', '', value)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--build', type=Path, required=True)
    parser.add_argument('--prior', type=Path)
    args = parser.parse_args(); checks = []; snapshots = {}
    for case in ('structured', 'storage-sharing', 'storage-sharing-partial'):
        for language in ('english', 'chinese'):
            base = args.build / 'renders' / f'{case}-{language}'
            html = BeautifulSoup(base.with_suffix('.html').read_text(), 'html.parser')
            q5 = html.find(id='q-store-backup'); q10 = html.find(id='q-share-restrictions')
            for qid in ('q-share-restrictions', 'q-data-preservation', 'q-access-data', 'q-persistent-identifier'):
                root = html.find(id=qid)
                assert root.select_one('.dataset-section'), (case, language, qid)
                assert not root.select('ul ul, ol ol, ul ol, ol ul'), (case, language, qid, 'nested lists in these fixtures')
                assert not root.select('p p, p div, p table, p ul, p ol'), (case, language, qid, 'invalid blocks')
            needles = ['2048', 'station_YYYYMMDD.csv', 'CHANGELOG.md', 'CoastView 1.0 (synthetic test tool)', '2027-12-31']
            assert len(q10.select('.distribution-section')) == (1 if case == 'structured' else 2)
            if case != 'structured':
                needles += ['2027-06-01', 'https://example.org/coast/review-terms']
                needles += (['The archive will be stored on tape.', 'Dedicated specialists will operate the shared workspace.', 'Do not redistribute the preliminary files.'] if language == 'english'
                            else ['封存資料將儲存於磁帶。', '共享工作空間將由專責人員維運。', '不得轉散布初步檔案。'])
                assert q5.select_one('[data-fact-id="archive-medium"][data-status="complete"]')
                if case.endswith('partial'):
                    assert q10.select_one('[data-fact-id="distribution-access"][data-status="missing"]')
                    assert q10.select_one('[data-fact-id="distribution-repository"][data-status="missing"]')
                    assert not q5.select('[data-fact-id="archive-remote-location"], [data-fact-id="archive-frequent-backup-need"]')
                else:
                    assert q5.select_one('[data-fact-id="archive-frequent-backup-need"][data-status="complete"]')
                    assert q5.select_one('[data-fact-id="storage-location"][data-status="partial"]')
                    assert q5.select_one('[data-fact-id="backup-frequency"][data-status="partial"]')
            pdf = subprocess.check_output(['pdftotext', str(base.with_suffix('.pdf')), '-'], text=True)
            pdf = re.sub(r'(?m)^\s*\d+\s*/\s*\d+\s*$', '', pdf)
            word = Document(base.with_suffix('.docx'))
            word_text = '\n'.join([p.text for p in word.paragraphs] + [cell.text for table in word.tables for row in table.rows for cell in row.cells])
            for fmt, text in [('html', html.get_text()), ('pdf', pdf), ('docx', word_text)]:
                for needle in needles: assert compact(needle) in compact(text), (case, language, fmt, needle)
            markers = [(q['id'], [(n.get('data-item-id'), n.get('data-fact-id'), n.get('data-status')) for n in q.select('[data-item-id], [data-fact-id], [data-status]')]) for q in html.select('.question')]
            snapshots[(case, language)] = markers
            checks.append({'case': case, 'language': language, 'selected_facts_in_all_formats': True, 'fixture_max_list_depth': 1})
        assert snapshots[(case, 'english')] == snapshots[(case, 'chinese')], (case, 'bilingual state mismatch')
    unchanged = []
    if args.prior:
        for language in ('english', 'chinese'):
            a, b = [BeautifulSoup((root / 'renders' / f'structured-{language}.html').read_text(), 'html.parser') for root in (args.prior, args.build)]
            for number in [1, 2, 3, 4, 6, 7, 8, 9, 14, 15]:
                old = next(q for q in a.select('.question') if q.h3 and re.match(rf'^{number}\.\s', q.h3.get_text(strip=True)))
                assert compact(old.get_text()) == compact(b.find(id=old['id']).get_text()), (language, number, 'unrelated question changed')
                unchanged.append({'language': language, 'question': number, 'text_unchanged': True})
    report = {'checks': checks, 'selected_checks_passed': True, 'unchanged_question_checks': unchanged,
              'release_acceptance': False, 'checker_sha256': hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
              'limits': ['Only declared synthetic fixtures and selected facts', 'Questionnaire gaps remain; not full Science Europe acceptance', 'Microsoft Word has not been tested'],
              'sha256': {str(p.relative_to(args.build)): hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted((args.build / 'renders').iterdir()) if p.is_file()}}
    (args.build / 'storage-sharing-report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps({'case_language_pairs': len(checks), 'unchanged_question_checks': len(unchanged), 'selected_checks_passed': True}))


if __name__ == '__main__': main()
