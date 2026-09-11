"""Trace selected facts across actual files, and inventory supplied reply strings.

String matches are evidence of text presence, not semantic coverage. Options,
numbers, list controls and unmatched strings remain explicitly unreviewed.
"""

import argparse
import hashlib
import json
import re
import subprocess
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document


def compact(value):
    # Only normalise whitespace and Markdown bullet prefixes, never letter case.
    value = re.sub(r'(?m)^\s*[-*+]\s+', '', value)
    return re.sub(r'\s+', '', value)


def file_texts(build, case, language):
    base = build / 'renders' / f'{case}-{language}'
    soup = BeautifulSoup(base.with_suffix('.html').read_text(), 'html.parser')
    word = Document(base.with_suffix('.docx'))
    paragraphs = [p.text for p in word.paragraphs]
    paragraphs += [c.text for table in word.tables for row in table.rows for c in row.cells]
    pdf = subprocess.check_output(['pdftotext', '-layout', str(base.with_suffix('.pdf')), '-'], text=True)
    pdf = re.sub(r'(?m)^\s*\d+\s*/\s*\d+\s*$', '', pdf)
    return soup, {'html': soup.get_text(' ', strip=True), 'pdf': pdf, 'docx': '\n'.join(paragraphs)}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--build', type=Path, required=True)
    parser.add_argument('--english', type=Path, required=True)
    parser.add_argument('--upstream', type=Path, required=True)
    args = parser.parse_args()
    checks, ledger, structural = [], [], []
    for case in ('representative', 'retention-partial'):
        for language, locale in (('english', 'en'), ('chinese', 'zh-Hant')):
            soup, texts = file_texts(args.build, case, language)
            q3 = soup.find(id='q-docs-metadata').get_text(' ', strip=True)
            q12 = soup.find(id='q-access-data').get_text(' ', strip=True)
            instructions = 'including instructions how to get access to the data' if language == 'english' else '並包含如何取用資料的說明'
            checks.append({'case': case, 'language': language, 'check': 'capacity-in-q3', 'passed': '2048' in q3})
            checks.append({'case': case, 'language': language, 'check': 'instructions-yes-in-q3', 'passed': instructions in q3})
            needles = ['2048', 'Dublin Core', 'CSV (UTF-8)', instructions]
            if case == 'representative':
                needles += ['station_YYYYMMDD.csv', 'CHANGELOG.md', 'CoastView 1.0 (synthetic test tool)', 'https://example.org/coastview/1.0']
            else:
                pending = 'Software is required, but the tools have not yet been listed.' if language == 'english' else '使用此資料集需要特定軟體，但所需工具尚未列出。'
                needles += [pending]
                checks.append({'case': case, 'language': language, 'check': 'required-software-list-pending',
                    'passed': soup.select_one('#q-access-data [data-fact-id="required-software-list"][data-status="missing"]') is not None
                              and pending in q12 and 'There are no tools needed' not in q12 and '不需要工具' not in q12})
            for fmt, text in texts.items():
                for needle in needles:
                    checks.append({'case': case, 'language': language, 'format': fmt, 'fact': needle,
                                   'passed': compact(needle) in compact(text)})
            for question in soup.select('.question'):
                invalid = question.select('p p, p ul, p ol, p div, p table')
                if invalid:
                    structural.append({'case': case, 'language': language, 'question': question['id'],
                                       'nested_block_occurrences': len(invalid)})
            events = json.loads((args.english / f'fixtures/pilot/{locale}/{case}.events.json').read_text())
            for event in events:
                reply = event['value']
                value = reply['value']
                entry = {'case': case, 'language': language, 'path': event['path'], 'type': reply['type']}
                if reply['type'] == 'StringReply' and isinstance(value, str) and len(value) >= 6:
                    needle = compact(value)
                    locations = [q['id'] for q in soup.select('.question') if needle in compact(q.get_text(' ', strip=True))]
                    found = needle in compact(texts['html'])
                    entry.update(status='text-present' if found else 'unmatched-needs-review',
                                 value=value, locations=locations or (['frontmatter-or-other'] if found else []))
                else:
                    entry.update(status='requires-semantic-review')
                ledger.append(entry)
    upstream_partial = BeautifulSoup((args.upstream / 'renders/upstream-partial-english.html').read_text(), 'html.parser')
    upstream_full = BeautifulSoup((args.upstream / 'renders/upstream-representative-english.html').read_text(), 'html.parser')
    before = {
        'baseline': 'dsw:science-europe:1.30.1',
        'storage_amount_missing_in_partial': '2048' not in upstream_partial.find(id='q-docs-metadata').get_text(),
        'required_software_misreported_as_none': 'There are no tools needed' in upstream_partial.find(id='q-access-data').get_text(),
        'file_case_corrupted': 'changelog.md' in upstream_full.find(id='q-docs-metadata').get_text(),
        'instructions_yes_detail_missing': 'including instructions how' not in upstream_full.find(id='q-docs-metadata').get_text(),
    }
    report = {
        'selected_regressions_passed': all(c['passed'] for c in checks),
        'release_acceptance': False,
        'upstream_counterexamples': before,
        'checks': checks,
        'remaining_structure_findings': structural,
        'reply_inventory': ledger,
        'inventory_counts': {state: sum(r['status'] == state for r in ledger) for state in ('text-present', 'unmatched-needs-review', 'requires-semantic-review')},
        'limits': ['Not full answer coverage or Science Europe compliance', 'Text presence does not prove correct placement, meaning, or visual layout',
                   'All option/list/numeric semantics still require explicit mapping review', 'No actual Microsoft Word test', 'Markdown pipe tables remain unsupported in unchanged worker'],
        'sha256': {str(p.relative_to(args.build)): hashlib.sha256(p.read_bytes()).hexdigest()
                   for p in sorted((args.build / 'renders').glob('*')) if p.suffix in ('.html', '.pdf', '.docx')},
    }
    output = args.build / 'answer-retention-report.json'
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps({'output': str(output), 'selected_regressions_passed': report['selected_regressions_passed'],
                     'remaining_structure_findings': len(structural), 'inventory_counts': report['inventory_counts']}))
    if not report['selected_regressions_passed']:
        raise SystemExit(2)


if __name__ == '__main__':
    main()
