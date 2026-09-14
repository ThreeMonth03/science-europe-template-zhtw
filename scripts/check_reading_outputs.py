"""Check actual Q2/Q3 artifacts and bounded table pagination, not release readiness."""
import argparse
import copy
import json
import subprocess
from pathlib import Path
from xml.etree import ElementTree as ET
from bs4 import BeautifulSoup
from docx import Document
from check_narrative_outputs import compact, pdf_text, page_bounds, sha


def compare_unchanged(before, after):
    assert len(before.select('.question')) == len(after.select('.question')) == 15
    for old in before.select('.question'):
        qid = old['id']; new = after.find(id=qid)
        if qid == 'q-what-data':
            # Only remove the identified instrument section, not all of Q2.
            old, new = copy.deepcopy(old), copy.deepcopy(new)
            for node in (old, new):
                headings = [h for h in node.select('h4') if h.get_text(strip=True) in ('Instrument datasets', '儀器資料集')]
                assert len(headings) <= 1
                if headings: headings[0].parent.decompose()
        assert compact(old.get_text()) == compact(new.get_text()), (qid, 'unexpected text change')
    return 15


def locations(path, labels, start, end):
    # A normal author answer elsewhere may repeat a table-cell label. Keep the
    # real page numbers, but search only the question containing provenance.
    pages = [compact(page) for page in subprocess.check_output(['pdftotext', str(path), '-'], text=True).split('\f')]
    text = '\f'.join(pages)
    start, end = compact(start), compact(end)
    assert text.count(start) == text.count(end) == 1, (path, 'Ambiguous question boundaries')
    first, last = text.index(start), text.index(end)
    assert first < last, (path, 'Reversed question boundaries')
    found, offset = {label: [] for label in labels}, 0
    for index, page in enumerate(pages, 1):
        scoped = page[max(0, first - offset):max(0, last - offset)]
        for label in labels:
            if compact(label) in scoped: found[label].append(index)
        offset += len(page) + 1
    return found


def authored_lines(path, language, expected):
    first = '保留 v1.2 與 station_YYYYMMDD.csv。' if language == 'chinese' else 'Keep v1.2 and station_YYYYMMDD.csv.'
    second = '這是使用者另外撰寫的段落！' if language == 'chinese' else 'This is a separate authored paragraph!'
    root = ET.fromstring(subprocess.check_output(['pdftotext', '-bbox-layout', str(path), '-']))
    found = {first: [], second: []}
    for index, page in enumerate(root.findall('.//{*}page'), 1):
        for line in page.findall('.//{*}line'):
            text = ''.join(w.text or '' for w in line.findall('{*}word'))
            assert not (compact(first) in compact(text) and compact(second) in compact(text)), (path, 'authored paragraphs share a line')
            for phrase in found:
                if compact(phrase) in compact(text): found[phrase].append((index, float(line.get('yMin')), float(line.get('yMax'))))
    assert len(found[first]) == len(found[second]) == expected, (path, found)
    for a, b in zip(found[first], found[second]):
        assert b[0] > a[0] or (b[0] == a[0] and b[1] > a[2]), (path, 'authored paragraphs overlap', a, b)


def inspect(build, case, language):
    base = build / 'renders' / f'{case}-{language}'
    soup = BeautifulSoup(base.with_suffix('.html').read_text(), 'html.parser')
    word = Document(base.with_suffix('.docx'))
    paragraphs = [p.text for p in word.paragraphs]
    word_text = '\n'.join(paragraphs + [c.text for t in word.tables for r in t.rows for c in r.cells])
    texts = {'pdf': pdf_text(base.with_suffix('.pdf')), 'docx': word_text}
    if case == 'table-long':
        for label in (f'ROW-{i:02d}' for i in range(1, 49)):
            for fmt, text in {'html': soup.get_text(), **texts}.items():
                assert text.count(label) == 1, (language, fmt, 'long table row lost or duplicated', label)
    for qid in ('q-what-data', 'q-docs-metadata'):
        q = soup.find(id=qid)
        assert not q.select('p p, p div, p ul, p table')
        for node in q.select('p, li'):
            for fmt, text in texts.items():
                assert compact(node.get_text()) in compact(text), (case, language, fmt, qid, node.get_text())
    for summary in soup.select('.collection-summary'):
        assert not summary.select('ul, p, br')
        assert any(compact(p) == compact(summary.get_text()) for p in paragraphs), (case, language, 'Q2 summary fragmented in Word')
    joined = 0
    for policy in soup.select('#q-docs-metadata .dataset-policy'):
        runs, pending = [], []
        for block in policy.find_all(recursive=False):
            if block.name == 'p': pending.append(block.get_text())
            elif pending: runs.append(pending); pending = []
        if pending: runs.append(pending)
        for run in runs:
            if len(run) > 1:
                assert any(compact(''.join(run)) in compact(p) for p in paragraphs), (case, language, 'Q3 fixed sentences fragmented')
                joined += 1
    if case in ('reading-rich', 'reading-partial'):
        authored_lines(base.with_suffix('.pdf'), language, 4 if case == 'reading-rich' else 3)
        preview = build / 'word-preview' / (base.name + '.pdf')
        if preview.exists(): authored_lines(preview, language, 4 if case == 'reading-rich' else 3)
        q2 = soup.find(id='q-what-data')
        for fact in ('metadata-access-explanation', 'file-naming', 'object-naming', 'external-ownership'):
            detail = soup.select_one(f'#q-docs-metadata [data-fact-id="{fact}"], #q-what-data [data-fact-id="{fact}"]')
            assert detail
            if case == 'reading-partial' and fact == 'external-ownership':
                assert detail['data-status'] == 'missing'; continue
            detail = copy.deepcopy(detail)
            for lead in detail.select('.answer-lead'): lead.decompose()
            assert len(detail.select('p')) == 2 and len(detail.select('ul > li')) == 2, (case, language, fact)
            for p in detail.select('p'):
                assert any(compact(p.get_text()) == compact(v) for v in paragraphs), (case, language, fact, 'authored paragraph joined')
            assert 'v1.2' in detail.get_text() and 'station_YYYYMMDD.csv' in detail.get_text()
        if case == 'reading-partial':
            assert len(q2.select('[data-fact-id="equipment-documentation"][data-status="missing"]')) == 2
            assert len(q2.select('[data-fact-id="data-collector"][data-status="missing"]')) == 1
            assert not q2.select('[data-status="explicit-no"]')
    table = soup.select_one('[data-fact-id="provenance"] table')
    pagination = {}
    if table:
        hint = table.find_parent(class_='short-table-unit')
        labels = ([f'ROW-{i:02d}' for i in range(1, 49)] if case == 'table-long' else
                  ['Processing log', 'Source checksum'] if language == 'english' else ['處理紀錄', '來源校驗碼'])
        if case == 'table-long':
            assert hint is None and len(table.select('tbody tr')) == 48
        else: assert hint is not None, (case, language, 'short table not tagged')
        for fmt, text in texts.items():
            for label in labels: assert label in text, (case, language, fmt, label)
        for kind, pdf in [('pdf', base.with_suffix('.pdf')), ('word-preview', build / 'word-preview' / (base.name + '.pdf'))]:
            if not pdf.exists(): continue
            found = locations(pdf, labels, soup.find(id='q-how-data').h3.get_text(), soup.find(id='q-what-data').h3.get_text())
            assert all(len(pages) == 1 for pages in found.values()), (case, language, kind, found)
            page_set = {pages[0] for pages in found.values()}
            assert (len(page_set) > 1) if case == 'table-long' else (len(page_set) == 1), (case, language, kind, found)
            pagination[kind] = sorted(page_set)
        if case != 'table-long':
            target = next(t for t in word.tables if labels[0] in '\n'.join(c.text for r in t.rows for c in r.cells))
            assert all(p.style.name == 'Pilot Table Lead' for r in target.rows[:-1] for c in r.cells for p in c.paragraphs)
            assert all(p.style.name != 'Pilot Table Lead' for c in target.rows[-1].cells for p in c.paragraphs)
    row = {'case': case, 'language': language, 'joined_q3_runs': joined, 'pdf_pages': page_bounds(base.with_suffix('.pdf')),
           'native_provenance_table': table is not None, 'table_pages': pagination, 'passed': True}
    preview = build / 'word-preview' / (base.name + '.pdf')
    if preview.exists(): row['word_preview_pages'] = page_bounds(preview)
    return row


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--build', type=Path, required=True)
    parser.add_argument('--prior', type=Path)
    parser.add_argument('--cases', nargs='+', required=True)
    args = parser.parse_args(); rows = []; comparisons = 0
    for case in args.cases:
        for language in ('english', 'chinese'):
            rows.append(inspect(args.build, case, language))
            if args.prior:
                name = f'{case}-{language}.html'
                old, new = args.prior / 'renders' / name, args.build / 'renders' / name
                if old.exists():
                    for fmt in ('html', 'pdf', 'docx'):
                        a, b = [json.loads(p.with_suffix('.' + fmt + '.fixture.json').read_text()) for p in (old, new)]
                        for key in ('recipe_sha256', 'events_sha256', 'km_sha256'): assert a[key] == b[key], (case, language, key)
                    comparisons += compare_unchanged(*[BeautifulSoup(p.read_text(), 'html.parser') for p in (old, new)])
    report = {'selected_checks_passed': True, 'release_acceptance': False, 'checker_sha256': sha(Path(__file__)), 'rows': rows,
              'controlled_question_comparisons': comparisons,
              'package_sha256': {name: sha(args.build / name) for name in ('english.zip', 'chinese.zip')},
              'artifact_sha256': {str(p.relative_to(args.build)): sha(p) for folder in ('renders', 'word-preview') for p in sorted((args.build / folder).glob('*')) if p.is_file()},
              'limits': ['Synthetic selected branches only; no full SE topic coverage', 'Comparison excludes only Q2 instrument section; its new summaries checked separately', 'Stock worker cannot render Markdown tables', 'LibreOffice preview is not Microsoft Word pagination', 'Oversized single cells are excluded from keep hints, not a claim that every renderer splits them well']}
    (args.build / 'reading-report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps({'passed': True, 'pairs': len(rows), 'comparisons': comparisons}))


if __name__ == '__main__': main()
