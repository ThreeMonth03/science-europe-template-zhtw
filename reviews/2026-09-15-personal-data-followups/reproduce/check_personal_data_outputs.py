"""Native bilingual Q7/Q9 content, missing-state, block and pagination checks."""
import argparse
import json
from pathlib import Path
import subprocess
import sys
import zipfile
from bs4 import BeautifulSoup
from docx import Document
from artifact_utils import sha
from check_missing_info_outputs import inspect
from check_narrative_outputs import compact
from check_budget_spacing_outputs import pdf_raw_page_texts
from check_word_rhythm_outputs import inspect_preview
from check_budget_outputs import compare_word
from compare_runtime_outputs import markers
from probe_personal_data_translation import EXPECTED

ROOT = Path(__file__).resolve().parents[1]


def section_word(document, number):
    paragraphs = document.paragraphs
    start = next(i for i, p in enumerate(paragraphs) if p.style.name == 'Heading 3' and p.text.startswith(str(number)+'. '))
    end = next(i for i, p in enumerate(paragraphs[start+1:], start+1) if p.style.name == 'Heading 3')
    return paragraphs[start+1:end]


def check_row(build, name, language, english):
    row, soup = inspect(build, name, language)
    q7, q9 = soup.select_one('#q-personal-data'), soup.select_one('#q-ethical-issues')
    assert q7 and q9
    facts = {n['data-fact-id'] for n in q7.select('[data-fact-id][data-status="missing"]')}
    assert facts == EXPECTED[name], (name, language, facts)
    base = build/'renders'/(name+'-'+language)
    word = Document(base.with_suffix('.docx'))
    word_text = compact(''.join(n.text or '' for n in word.element.body.xpath('.//w:t')))
    word_q7 = compact(''.join(p.text for p in section_word(word, 7)))
    assert compact(q7.select_one('.answer').get_text()) == word_q7, (name, language, 'Word Q7 text/ordering changed')
    for n in soup.select('#dmp-content .data-gap, #dmp-content .answer-detail p'):
        assert compact(n.get_text()) in word_text, (name, language, 'Word lost text', n.get_text())
    locale = 'en' if language == 'english' else 'zh-Hant'
    recipe = english/'fixtures/pilot'/locale/(name+'.json')
    events = recipe.parent/json.loads(recipe.read_text())['events_file']
    meta = [json.loads(base.with_suffix('.'+fmt+'.fixture.json').read_text()) for fmt in ['html', 'pdf', 'docx']]
    for item in meta:
        assert item['recipe_sha256'] == sha(recipe) and item['events_sha256'] == sha(events)
        assert item['package_sha256'] == sha(build/(language+'.zip'))
    for key in ['km_sha256', 'events_sha256', 'recipe_sha256', 'package_sha256']:
        assert len({m[key] for m in meta}) == 1
    pages = pdf_raw_page_texts(base.with_suffix('.pdf')); all_text = ''.join(pages)
    q7_text = all_text.split(compact(q7.h3.get_text()), 1)[1].split(compact(soup.select_one('#q-copyright-ipr h3').get_text()), 1)[0]
    cursor = 0
    for leaf in q7.select('.answer h4, .answer p, .answer li'):
        value = compact(leaf.get_text()); position = q7_text.find(value, cursor)
        assert position >= cursor, (name, language, 'PDF Q7 block omitted or reordered', value)
        cursor = position+len(value)
    for bad in ['We explored General Data Protection Regulation', 'more important than the privacy',
                '我們已檢視歐盟《個人資料保護規則》', '該公共利益高於資料主體的隱私利益']:
        assert compact(bad) not in compact(q9.get_text()) and compact(bad) not in all_text
    if name.startswith('personal-followups') or name.startswith('personal-transfer'):
        assert not any(x in q9.get_text() for x in ['not subject to ethical legislation', '不受倫理法規'])
    details = q7.select('.answer-detail[data-status="complete"]')
    for detail in details:
        assert [n.name for n in detail.find_all(recursive=False)] == ['p', 'p', 'ul']
        assert len(detail.select('ul > li')) == 2
        assert detail.find_previous_sibling().get('class') == ['answer-lead']
    if name == 'personal-transfer-complete':
        link = q7.select_one('[data-fact-id="personal-data-transfer-measures"] a')
        expected_link = 'https://example.org/transfer?a=1&b=2'
        assert link and link['href'] == expected_link
        assert expected_link in [r.target_ref for r in word.part.rels.values() if r.is_external]
    preview = build/'word-preview'/(name+'-'+language+'.pdf')
    row['word_pages'] = inspect_preview(preview)
    preview_pages = [compact(p) for p in subprocess.check_output(['pdftotext', '-raw', str(preview), '-'], text=True).split('\f') if p.strip()]
    # Retain every Q7 paragraph / list item in the Word preview too. Locate
    # blocks independently because its generated footer can precede page text.
    for leaf in q7.select('.answer p, .answer li'):
        assert compact(leaf.get_text()) in ''.join(preview_pages), (name, language, 'Word preview Q7 text missing')
    if name == 'empty':
        nodes = soup.select('#q-required-resources h3, #q-required-resources .data-gap')
        assert len(nodes) == 5
        assert any(all(compact(n.get_text()) in page for n in nodes) for page in preview_pages)
    row.update(q7_missing_facts=sorted(facts), q7_original_detail_blocks=len(details),
               q7_word_text_and_pdf_order_retained=True, q9_unsupported_claims_absent=True)
    for fmt in ['docx', 'docx.fixture.json']:
        path = base.with_suffix('.'+fmt); row['artifact_sha256'][str(path.relative_to(build))] = sha(path)
    row['artifact_sha256'][str(preview.relative_to(build))] = sha(preview)
    return row, soup


def compare_control(build, prior, name, language, soup):
    stem = name+'-'+language
    old = prior/'renders'/stem; new = build/'renders'/stem
    a, b = [json.loads(p.with_suffix('.html.fixture.json').read_text()) for p in [old, new]]
    for key in ['recipe_sha256', 'events_sha256', 'km_sha256']: assert a[key] == b[key]
    old_soup = BeautifulSoup(old.with_suffix('.html').read_text(), 'html.parser')
    for left, right in zip(old_soup.select('.question'), soup.select('.question')):
        assert str(left) == str(right), (name, language, 'Control HTML changed', left['id'])
    compare_word(Document(old.with_suffix('.docx')), Document(new.with_suffix('.docx')), False)
    with zipfile.ZipFile(old.with_suffix('.docx')) as left, zipfile.ZipFile(new.with_suffix('.docx')) as right:
        assert left.read('word/styles.xml') == right.read('word/styles.xml')
    def locations(pdf, source):
        pages = pdf_raw_page_texts(pdf)
        return [[i for i, text in enumerate(pages, 1) if compact(q.h3.get_text()) in text] for q in source.select('.question')]
    assert locations(old.with_suffix('.pdf'), old_soup) == locations(new.with_suffix('.pdf'), soup)
    return {'questions': 15, 'word_body_styles_links_unchanged': True, 'pdf_question_pages_unchanged': True,
            'prior_artifact_sha256': {str(old.with_suffix('.'+f)): sha(old.with_suffix('.'+f)) for f in ['html', 'pdf', 'docx', 'html.fixture.json']}}


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--build', type=Path, required=True); p.add_argument('--english', type=Path, required=True)
    p.add_argument('--prior', type=Path, required=True); p.add_argument('--control-prior', type=Path, required=True)
    p.add_argument('--cases', nargs='+', default=list(EXPECTED)); p.add_argument('--output', type=Path)
    a = p.parse_args()
    report = {'selected_checks_passed': False, 'release_acceptance': False, 'version': '0.3.21', 'rows': [],
              'checker_sha256': sha(Path(__file__)), 'package_sha256': {n: sha(a.build/n) for n in ['english.zip', 'chinese.zip']},
              'limits': ['Selected synthetic Q7/Q9 branches, not all reachable questionnaire fields or legal compliance',
                         'Native PDF and LibreOffice previews, not Microsoft Word visual acceptance',
                         'Short budget gap wrapping and overall blank space are unchanged']}
    target = a.output or a.build/'personal-data-report.json'; assert not target.exists()
    try:
        for name in a.cases:
            soups = []
            for language in ['english', 'chinese']:
                row, soup = check_row(a.build, name, language, a.english.resolve()); soups.append(soup)
                if name in ['empty', 'negative', 'preservation-complete']:
                    row['unchanged_control'] = compare_control(a.build, a.control_prior if name == 'preservation-complete' else a.prior,
                                                               name, language, soup)
                report['rows'].append(row)
                print(json.dumps({k: row[k] for k in ['case', 'language', 'pages', 'word_pages', 'errors', 'reading_issues']}, ensure_ascii=False), flush=True)
            assert markers(soups[0]) == markers(soups[1]), (name, 'Bilingual fact/status/ownership drift')
        report['selected_checks_passed'] = len(report['rows']) == len(a.cases)*2 and all(not r['errors'] and not r['reading_issues'] for r in report['rows'])
    finally:
        target.write_text(json.dumps(report, ensure_ascii=False, indent=2)+'\n')
    assert report['selected_checks_passed'], 'See the preserved failure report'


if __name__ == '__main__': main()
