"""Column/row-aware proof for the public marked mixed-budget Word fixtures.

Not a general PDF reading-order repair. Unique, one-line resource titles and
three verified column headings bind every body line to its source resource.
Repeated identity headers are accepted only at the start of a continuation
page and must repeat the complete original cells. No arbitrary text deletion.
"""
from collections import Counter
import hashlib
import json
import subprocess
from bs4 import BeautifulSoup
from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from check_budget_outputs import body
from check_word_short_budget_outputs import paragraph_texts, word_pages
from mixed_row_content import MARKERS, cell_tokens, compact


def line_text(line):
    return compact(''.join(line.itertext()))


def read_rows(soup, bbox, expanded):
    table = soup.select('#q-required-resources .resource-table')
    assert len(table) == 1 and not MARKERS.intersection(soup.get_text())
    source = table[0].tbody.find_all('tr', recursive=False)
    tokens = [[cell_tokens(c) for c in row.find_all('td', recursive=False)] for row in source]
    assert all(len(row) == 3 for row in tokens)
    titles = [row[0][0] for row in tokens]
    assert len(titles) == len(set(titles))
    headings = [compact(c.get_text()) for c in table[0].thead.find_all('th')]
    assert len(headings) == 3
    pages = etree.fromstring(bbox).findall('.//{*}page')
    title_hits = [(n, line) for n, page in enumerate(pages) for line in page.findall('.//{*}line')
                  if line_text(line) == titles[0]]
    assert title_hits
    first_page = title_hits[0][0]
    streams = [[[] for _ in range(3)] for _ in source]
    current = -1; repeats = []; segments = []
    for page_index in range(first_page, len(pages)):
        page = pages[page_index]; number = page_index + 1
        lines = page.findall('.//{*}line')
        header_lines = [l for l in lines if line_text(l) in headings]
        headers = []
        for first in [l for l in header_lines if line_text(l) == headings[0]]:
            group = [[l for l in header_lines if line_text(l) == text and abs(float(l.get('yMin')) - float(first.get('yMin'))) < 4]
                     for text in headings]
            assert all(len(g) == 1 for g in group), 'Ambiguous or incomplete column heading'
            headers.append([g[0] for g in group])
        assert headers and len(header_lines) == len(headers) * 3
        starts = [float(l.get('xMin')) for l in headers[0]]
        assert starts == sorted(starts)
        assert all([float(l.get('xMin')) for l in h] == starts for h in headers)
        top = min(float(l.get('yMin')) for l in header_lines)
        content = []
        for line in lines:
            value = line_text(line)
            if line in header_lines:
                continue
            if value == f'{number}/{len(pages)}':
                assert float(line.get('yMin')) > float(page.get('height')) * .9
                continue
            if float(line.get('yMax')) < top:
                assert page_index == first_page, 'Unexpected text before continuation heading'
                continue
            assert float(line.get('yMin')) >= top
            assert float(line.get('yMax')) < float(page.get('height')) * .93, 'Budget text intrudes into footer'
            content.append(line)
        anchors = sorted([l for l in content if line_text(l) in titles], key=lambda l: float(l.get('yMin')))
        # Title style is 4 pt lower than the other cells in the owned fixture.
        # Windows tolerate only 6 pt and never accept a line crossing a boundary.
        cuts = [(float(l.get('yMin')) - 6, titles.index(line_text(l))) for l in anchors]
        windows = [(-float('inf'), current)] + cuts
        for index, (lower, resource) in enumerate(windows):
            upper = windows[index + 1][0] if index + 1 < len(windows) else float('inf')
            chosen = [l for l in content if lower <= float(l.get('yMin')) < upper]
            if not chosen and index == 0:
                continue
            assert resource >= 0, 'Unbound text before first resource'
            assert all(float(l.get('yMax')) < upper for l in chosen), 'Text crosses resource boundary'
            if index:
                repeat = resource == current
                if repeat:
                    assert resource in expanded and index == 1, 'Unexpected repeated resource title'
                    assert not any(float(l.get('yMin')) < lower for l in content), 'Identity repeat is not page-leading'
                else:
                    assert resource == current + 1, 'Resource order changed'
                current = resource
            else:
                repeat = False
            values = [[] for _ in range(3)]
            for line in sorted(chosen, key=lambda l: (float(l.get('yMin')), float(l.get('xMin')))):
                left, right = float(line.get('xMin')), float(line.get('xMax'))
                assert left >= starts[0] - .1 and right <= float(page.get('width'))
                column = max(c for c in range(3) if left >= starts[c] - .1)
                if resource not in expanded and column < 2:
                    assert right < starts[column + 1], 'Ordinary cell crosses columns'
                text = ''.join(c for c in line_text(line) if c not in MARKERS)
                values[column].append(text)
            joined = [''.join(v) for v in values]
            if repeat:
                assert joined[0].startswith(titles[resource])
                assert joined[1:] == [''.join(v) for v in tokens[resource][1:]], 'Continuation identity changed'
                joined[0] = joined[0][len(titles[resource]):]
                joined[1:] = ['', '']
                repeats.append(dict(resource=resource + 1, page=number))
            for column, text in enumerate(joined):
                streams[resource][column].append((number, text))
            segments.append(dict(resource=resource + 1, page=number, repeated_identity=repeat))
    assert current == len(source) - 1
    results = []
    for index, row in enumerate(tokens):
        split = []; row_pages = set(); first_cell_pages = []
        for column, paragraphs in enumerate(row):
            parts = streams[index][column]
            actual = ''.join(text for _, text in parts)
            assert actual == ''.join(paragraphs), (index + 1, column + 1, 'Cell content/order mismatch', actual[:120])
            locations = [n for n, text in parts for _ in text]
            row_pages.update(locations); cursor = 0
            for paragraph in paragraphs:
                on_pages = sorted(set(locations[cursor:cursor + len(paragraph)]))
                if column == 0: first_cell_pages.append(on_pages)
                if len(on_pages) > 1:
                    split.append(dict(column=column + 1, text=paragraph, pages=on_pages))
                cursor += len(paragraph)
        results.append(dict(identity=source[index]['data-item-id'], index=index + 1,
            cell_paragraphs=sum(len(c) for c in row), pages=sorted(row_pages), split_paragraphs=split,
            last_two_first_cell_paragraph_pages=first_cell_pages[-2:]))
    return dict(rows=results, repeated_identity_headers=repeats, segments=segments,
                original_cell_paragraphs_checked=sum(r['cell_paragraphs'] for r in results))


def inspect(document_path, preview, html):
    document = Document(document_path); soup = BeautifulSoup(html, 'html.parser')
    source = soup.select('#q-required-resources .resource-table tbody > tr')
    titles = [compact(row.find('td').find('p').get_text()) for row in source]
    budget_tables = [t for t in document.tables if any(compact(c.text).startswith(title)
                     for row in t.rows for c in row.cells for title in titles)]
    assert budget_tables
    expanded = set()
    document_cells = []
    for table in budget_tables:
        assert [compact(c.text) for c in table.rows[0].cells] == [compact(c.get_text()) for c in soup.select('#q-required-resources .resource-table th')]
        if table.style.name == 'Pilot Long Budget':
            title = compact(table.rows[1].cells[0].text)
            assert title in titles
            expanded.add(titles.index(title))
            values = [compact(c.text) for c in table.rows[1].cells]
            for row in table.rows[2:]:
                assert len({id(c._tc) for c in row.cells}) == 1
                values[0] += compact(row.cells[0].text)
            document_cells.append(values)
        else:
            assert table.style.name == 'Table'
            document_cells.extend([[compact(c.text) for c in row.cells] for row in table.rows[1:]])
    assert document_cells == [[compact(c.get_text()) for c in row.find_all('td', recursive=False)] for row in source], 'Word/HTML original cell content differs'
    # This fixture ends with the budget; only empty table separators follow.
    nodes = body(document); first = next(i for i, n in enumerate(nodes) if n is budget_tables[0]._tbl)
    assert all(n.tag == qn('w:tbl') or not ''.join(n.itertext()).strip() for n in nodes[first:])
    raw = subprocess.check_output(['pdftotext', '-raw', str(preview), '-'], text=True)
    bbox = subprocess.check_output(['pdftotext', '-bbox-layout', str(preview), '-'])
    full_text = ''.join(word_pages(raw, bbox))
    prefix = []; markers = 0
    pages = etree.fromstring(bbox).findall('.//{*}page')
    for number, page in enumerate(pages, 1):
        footer_words = {w for line in page.findall('.//{*}line')
                        if line_text(line) == f'{number}/{len(pages)}' and float(line.get('yMin')) > float(page.get('height')) * .9
                        for w in line.findall('{*}word')}
        for word in page.findall('.//{*}word'):
            if word in footer_words: continue
            if word.text == '15.': markers += 1
            if not markers: prefix.append((number, dict(word.attrib), word.text))
    assert markers == 1
    counts = Counter(compact(p) for p in paragraph_texts(document, soup) if compact(p))
    budget_counts = Counter(compact(p.text) for t in budget_tables for row in t.rows
                            for cell in {id(c._tc): c for c in row.cells}.values()
                            for p in cell.paragraphs if compact(p.text))
    assert not budget_counts - counts
    outside = counts - budget_counts
    assert all(full_text.count(p) >= n for p, n in outside.items()), 'Non-budget Word paragraph missing'
    result = read_rows(soup, bbox, expanded)
    result.update(non_budget_paragraphs_checked=sum(outside.values()),
                  source_budget_paragraphs=sum(budget_counts.values()),
                  q15_prefix_geometry_sha256=hashlib.sha256(json.dumps(prefix, sort_keys=True).encode()).hexdigest(),
                  all_original_budget_cells_verified=True, microsoft_word_acceptance=False,
                  layout_acceptance=False)
    return result
