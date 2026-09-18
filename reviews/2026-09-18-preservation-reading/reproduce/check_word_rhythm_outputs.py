"""Check Word styles and compare only explicitly matching prior fixtures."""
import argparse
import json
import subprocess
from pathlib import Path
from xml.etree import ElementTree as ET
from bs4 import BeautifulSoup
from docx import Document
from docx.enum.text import WD_LINE_SPACING
from check_narrative_outputs import compact, page_bounds, pdf_text, sha
from compare_runtime_outputs import markers


def assert_styles(document):
    for name in ['Normal', 'Body Text', 'First Paragraph', 'Compact']:
        style = document.styles[name]
        assert style.font.size.pt == 10.5, ('Body font size changed', name)
        p = style.paragraph_format
        assert p.line_spacing == 1.2 and p.line_spacing_rule == WD_LINE_SPACING.MULTIPLE
        assert p.space_after.pt == (2 if name == 'Compact' else 4)
        assert p.widow_control
    for level in range(1, 6):
        p = document.styles[f'Heading {level}'].paragraph_format
        assert p.space_before.pt == (8 if level >= 4 else 12)
        assert p.space_after.pt == (3 if level >= 4 else 4 if level == 3 else 6)
        assert p.keep_with_next and p.keep_together and not p.page_break_before


def compare_questions(before, after):
    assert len(before.select('.question')) == len(after.select('.question')) == 15
    assert markers(before) == markers(after), 'Semantic markers changed in a style-only iteration'
    for old, new in zip(before.select('.question'), after.select('.question')):
        assert old['id'] == new['id'] and str(old) == str(new), (old['id'], 'Question HTML changed')
    return 15


def compare_pdf_text(before, after):
    # Compare the same renderer's whole extracted document. Nested lists add
    # bullets and table extraction interleaves columns, so HTML container text
    # is not a valid contiguous-PDF oracle. Preserve punctuation and case.
    assert compact(before) == compact(after), 'PDF text changed in a style-only iteration'


def assert_line_geometry(xml):
    # Bound the claim to lines inside each text block, not a whole-page visual
    # score. Tables may have separate blocks at the same vertical position.
    for block in ET.fromstring(xml).findall('.//{*}block'):
        lines = block.findall('{*}line')
        for i, first in enumerate(lines):
            for second in lines[i+1:]:
                width = min(float(first.get('xMax')),float(second.get('xMax'))) - max(float(first.get('xMin')),float(second.get('xMin')))
                height = min(float(first.get('yMax')),float(second.get('yMax'))) - max(float(first.get('yMin')),float(second.get('yMin')))
                assert not (width>0.5 and height>0.5), 'Overlapping lines within a Word-preview text block'


def inspect_preview(path):
    assert_line_geometry(subprocess.check_output(['pdftotext','-bbox-layout',str(path),'-']))
    return page_bounds(path)


def word_body(document):
    # Cover metadata changes version. Compare all body paragraphs and table cells
    # from the first numbered question, retaining paragraph/style boundaries.
    rows = []
    active = False
    for p in document.paragraphs:
        if p.text.startswith('1. '): active = True
        if active: rows.append((p.style.name, p.text))
    assert rows, 'No numbered body questions'
    tables = [[[c.text for c in r.cells] for r in t.rows] for t in document.tables]
    return rows, tables


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--build', type=Path, required=True)
    parser.add_argument('--prior', type=Path, required=True)
    parser.add_argument('--cases', nargs='+', required=True)
    parser.add_argument('--uncompared-cases', nargs='*', default=[], help='Additional style checks without a matching historical output')
    args = parser.parse_args()
    assert not set(args.cases) & set(args.uncompared_cases)
    report = {'selected_checks_passed': False, 'release_acceptance': False, 'rows': [],
              'checker_sha256': sha(Path(__file__)),
              'helper_sha256': {n: sha(Path(__file__).with_name(n)) for n in ['check_narrative_outputs.py', 'compare_runtime_outputs.py']},
              'package_sha256': {n: sha(args.build/n) for n in ['english.zip', 'chinese.zip']},
              'prior_package_sha256': {n: sha(args.prior/n) for n in ['english.zip', 'chinese.zip']},
              'artifact_sha256': {str(p.relative_to(args.build)): sha(p) for folder in ['renders','word-preview'] for p in sorted((args.build/folder).glob('*')) if p.is_file()},
              'prior_artifact_sha256': {}, 'uncompared_cases': args.uncompared_cases,
              'limits': ['Historical comparisons only for explicitly matching synthetic fixtures', 'Uncompared cases check styles and geometry, not historical content equality', 'LibreOffice previews are not Microsoft Word acceptance', 'PDF page bounds do not prove no text overlap', 'No content coverage added']}
    target = args.build/'word-rhythm-report.json'
    target.write_text(json.dumps(report, indent=2)+'\n')
    try:
        for case in args.uncompared_cases:
            for language in ['english', 'chinese']:
                stem = f'{case}-{language}'; base = args.build/'renders'/stem
                assert_styles(Document(base.with_suffix('.docx')))
                row = {'case': case, 'language': language, 'passed': True, 'question_comparisons': 0,
                       'pdf_pages': page_bounds(base.with_suffix('.pdf'))}
                preview = args.build/'word-preview'/(stem+'.pdf')
                if preview.exists(): row['word_preview_pages'] = inspect_preview(preview)
                report['rows'].append(row)
        for case in args.cases:
            for language in ['english', 'chinese']:
                stem = f'{case}-{language}'; old = args.prior/'renders'/stem; new = args.build/'renders'/stem
                for suffix in ['.html','.pdf','.docx']:
                    a,b = [json.loads(p.with_suffix(suffix+'.fixture.json').read_text()) for p in [old,new]]
                    for key in ['recipe_sha256','events_sha256','km_sha256']: assert a[key] == b[key], (stem,key)
                    for extra in ['', '.fixture.json']:
                        p = old.with_suffix(suffix+extra); report['prior_artifact_sha256'][str(p.relative_to(args.prior))] = sha(p)
                soups = [BeautifulSoup(p.with_suffix('.html').read_text(),'html.parser') for p in [old,new]]
                count = compare_questions(*soups)
                documents = [Document(p.with_suffix('.docx')) for p in [old,new]]
                assert_styles(documents[1])
                assert word_body(documents[0]) == word_body(documents[1]), (stem,'Native Word paragraph/cell content changed')
                a,b = [pdf_text(p.with_suffix('.pdf')) for p in [old,new]]
                compare_pdf_text(a,b)
                row = {'case': case, 'language': language, 'passed': True, 'question_comparisons': count,
                       'whole_pdf_text_unchanged': True,
                       'pdf_pages': page_bounds(new.with_suffix('.pdf')), 'prior_pdf_pages': page_bounds(old.with_suffix('.pdf'))}
                assert row['pdf_pages'] == row['prior_pdf_pages'], (stem,'Unexpected PDF pagination change')
                preview = args.build/'word-preview'/(stem+'.pdf')
                if preview.exists():
                    row['word_preview_pages'] = inspect_preview(preview)
                    prior_preview = args.prior/'word-preview'/(stem+'.pdf')
                    if prior_preview.exists():
                        row['prior_word_preview_pages'] = page_bounds(prior_preview)
                        report['prior_artifact_sha256'][str(prior_preview.relative_to(args.prior))] = sha(prior_preview)
                report['rows'].append(row)
    except Exception as e:
        report['failure'] = str(e); target.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n'); raise
    report['selected_checks_passed'] = True
    target.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'passed': True, 'pairs': len(report['rows']), 'question_comparisons': sum(r['question_comparisons'] for r in report['rows'])}))


if __name__ == '__main__': main()
