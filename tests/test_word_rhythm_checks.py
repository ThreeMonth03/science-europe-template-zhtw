import sys
import unittest
from pathlib import Path
from bs4 import BeautifulSoup
from docx import Document
sys.path.insert(0, str(Path(__file__).resolve().parents[1]/'scripts'))
from check_word_rhythm_outputs import compare_questions, word_body


class WordRhythmTests(unittest.TestCase):
    def test_style_iteration_rejects_text_structure_and_status_changes(self):
        source = ''.join(f'<div class="question" id="q-{i}"><p data-fact-id="f-{i}" data-status="complete">Answer {i}</p></div>' for i in range(15))
        self.assertEqual(15, compare_questions(*[BeautifulSoup(source,'html.parser') for _ in range(2)]))
        for changed in [source.replace('Answer 2','Changed 2'), source.replace('complete','missing',1), source.replace('<p ', '<div ',1)]:
            with self.assertRaises(AssertionError): compare_questions(*[BeautifulSoup(s,'html.parser') for s in [source,changed]])

    def test_native_word_comparison_ignores_cover_not_body_paragraphs_or_cells(self):
        a,b = Document(),Document()
        for d,version in [(a,'0.3.7'),(b,'0.3.8')]:
            d.add_paragraph(version); d.add_paragraph('1. Question'); d.add_paragraph('One'); d.add_paragraph('Two')
            d.add_table(rows=1, cols=1).cell(0,0).text = 'Cell'
        self.assertEqual(word_body(a), word_body(b))
        b.paragraphs[2].text = 'One Two'; b.paragraphs[3].text = ''
        self.assertNotEqual(word_body(a), word_body(b))
        b.paragraphs[2].text = 'One'; b.paragraphs[3].text = 'Two'; b.tables[0].cell(0,0).text = 'Lost'
        self.assertNotEqual(word_body(a), word_body(b))
