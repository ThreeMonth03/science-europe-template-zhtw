import sys
import unittest
from pathlib import Path
from bs4 import BeautifulSoup
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from check_contact_outputs import check_references, compare_prior, reference_text, LEAD, OLD


class ContactChecksTests(unittest.TestCase):
    def pair(self):
        body = '<p>Full-contact-2027.csv.</p><ul><li>Keep this.</li></ul>'
        old = '<div class="question" id="q-share-restrictions"><div class="dataset-section" data-item-id="d"><div class="distribution-section" data-item-id="r"><div class="distribution-reading-unit"><p>Access.</p><div class="repository-arrangement">Repository. ' + OLD['english'][0] + body + '</div><p>Licence.</p></div></div></div></div>'
        old += '<div class="question" id="q-data-preservation"><div class="dataset-section" data-item-id="d"><ul><li class="repository-distribution" data-item-id="r">Repository. ' + OLD['english'][1] + body + '</li></ul></div></div>'
        old += ''.join(f'<div class="question" id="q-{i}"><p>Untouched {i}.</p></div>' for i in range(13))
        new = old.replace(OLD['english'][0] + body, '<span class="repository-contact-reference"><a href="#repository-contact-1-1">' + reference_text('english', 1, False) + '</a></span>')
        new = new.replace(OLD['english'][1] + body, '<div class="repository-contact" id="repository-contact-1-1"><div class="answer-lead"><p>' + LEAD['english'] + '</p></div><div class="answer-detail" data-fact-id="repository-contact-arrangements" data-status="complete">' + body + '</div></div>')
        new = new.replace('data-item-id="r">Repository. <div class="repository-contact"', 'data-item-id="r"><div class="answer-lead repository-contact-heading">Repository. </div><div class="repository-contact"')
        return BeautifulSoup(old, 'html.parser'), BeautifulSoup(new, 'html.parser')

    def test_reconstruction_proves_authored_content_and_other_facts_unchanged(self):
        old, new = self.pair(); original = str(new)
        self.assertEqual(1, len(check_references(new, 'english')))
        self.assertEqual(15, compare_prior(old, new, 'english'))
        self.assertEqual(original, str(new))

    def test_wrong_target_label_and_item_are_rejected(self):
        for kind in ['target', 'label', 'item', 'missing-detail']:
            _, new = self.pair()
            if kind == 'target': new.select_one('a')['href'] = '#wrong'
            elif kind == 'label': new.select_one('a').string = 'Other dataset'
            elif kind == 'item': new.select_one('.repository-distribution')['data-item-id'] = 'wrong'
            else: new.select_one('.answer-detail').decompose()
            with self.assertRaises(AssertionError): check_references(new, 'english')

    def test_no_blanket_exclusion_of_q10_q11_or_free_answers(self):
        for selector in ['.answer-detail p', '.answer-detail li', '#q-2 p', '.distribution-reading-unit > p']:
            old, new = self.pair(); new.select_one(selector).string = 'Changed.'
            with self.assertRaises(AssertionError): compare_prior(old, new, 'english')

    def test_duplicate_target_is_rejected(self):
        _, new = self.pair(); new.append(BeautifulSoup('<div id="repository-contact-1-1"></div>', 'html.parser'))
        with self.assertRaises(AssertionError): check_references(new, 'english')

    def test_contact_fixture_keeps_stock_table_gate(self):
        from run_pilot import TABLE_CASES
        self.assertIn('contact-mixed', TABLE_CASES)

    def test_runtime_table_comparison_preserves_cells_and_other_text(self):
        from compare_runtime_outputs import compare_contact_tables
        before = '<div><p>Other facts.</p><div class="repository-contact" id="repository-contact-1-1"><div class="answer-detail"><p>| Record | Owner |\n| --- | --- |\n| Contact.csv | Team |</p></div></div></div>'
        after = before.replace('<p>| Record | Owner |\n| --- | --- |\n| Contact.csv | Team |</p>', '<table><thead><tr><th>Record</th><th>Owner</th></tr></thead><tbody><tr><td>Contact.csv</td><td>Team</td></tr></tbody></table>')
        a, b = [BeautifulSoup(s, 'html.parser') for s in [before, after]]
        self.assertEqual(1, compare_contact_tables(a, b))
        for bad in [after.replace('Contact.csv', 'Lost.csv'), after.replace('Other facts.', 'Changed.')]:
            with self.assertRaises(AssertionError): compare_contact_tables(a, BeautifulSoup(bad, 'html.parser'))

    def test_block_checks_keep_paragraphs_lists_and_table_cells(self):
        from check_block_content import assert_native_text
        from docx import Document
        node = BeautifulSoup('<li>Heading.<div><p>Authored.csv.</p><ul><li>List.csv</li></ul><table><tr><td>Cell.csv</td></tr></table></div></li>', 'html.parser')
        d = Document()
        for text in ['Heading.', 'Authored.csv.', 'List.csv']: d.add_paragraph(text)
        d.add_table(1, 1).cell(0, 0).text = 'Cell.csv'
        pdf = 'Heading. Authored.csv. ◦ List.csv Cell.csv'
        assert_native_text(node, pdf, d)
        for text in ['Authored.csv.', 'List.csv', 'Cell.csv']:
            with self.assertRaises(AssertionError): assert_native_text(node, pdf.replace(text, ''), d)
        d.tables[0].cell(0, 0).text = 'Lost.csv'
        with self.assertRaises(AssertionError): assert_native_text(node, pdf, d)
