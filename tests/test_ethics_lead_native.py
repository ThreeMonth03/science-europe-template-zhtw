import copy
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import unittest
from unittest.mock import patch as mock_patch
import zipfile
from bs4 import BeautifulSoup
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
ARCHIVE = ROOT / 'reviews/2026-09-21-ethics-lead'
sys.path[:0] = [str(ROOT / 'experiments/ethics-lead'), str(ROOT / 'scripts')]
from lead_recipe import baseline, patch, TARGET
from lead_native import run, visual_gate, word_lead_delta, visible_with_table_header, long_paragraph


def english_root():
    return next(p for p in [ROOT.parent / 'english', ROOT.parent / 'science-europe-template']
                if (p / 'scripts/output_profile_contract.py').is_file())


class EthicsLeadNativeTests(unittest.TestCase):
    def test_sealed_inventory_and_bounded_acceptance(self):
        seal = ARCHIVE / 'checksums.json'
        self.assertEqual(hashlib.sha256(seal.read_bytes()).hexdigest(),
                         '9015bf798a4fb25db040d2abf713ed497f8ffa4184a0896ea422eafc7947d4ef')
        files = json.loads(seal.read_text())
        self.assertEqual(len(files), 267)
        self.assertEqual(set(files), {str(p.relative_to(ARCHIVE)) for p in ARCHIVE.rglob('*') if p.is_file() and p != seal})
        for name, digest in files.items(): self.assertEqual(hashlib.sha256((ARCHIVE / name).read_bytes()).hexdigest(), digest)
        inventory = json.loads((ARCHIVE / 'inventory.json').read_text())
        for key in ['source_integrated', 'translation_tree_modified', 'production_touched', 'full_visual_acceptance',
                    'global_switch_complete', 'release_acceptance', 'microsoft_word_acceptance']: self.assertFalse(inventory[key])
        self.assertTrue(inventory['lead_separation_resolved']); self.assertFalse(inventory['source_integration_allowed'])
        self.assertFalse(inventory['visual_gate_passed'])
        self.assertEqual((inventory['structural_checks'], inventory['native_pairs'], inventory['native_artifacts'],
                          inventory['word_previews']), (544, 12, 72, 24))
        self.assertEqual(len(inventory['known_preexisting_visual_issues']), 1)

    def test_native_content_style_placement_and_long_paragraphs_reproduce(self):
        rows = run(ARCHIVE / 'before-core', ARCHIVE / 'before-long', ARCHIVE / 'after',
                   ARCHIVE / 'fixtures', english_root(), compacted=True)
        proof = json.loads((ARCHIVE / 'provenance/native.json').read_text())
        self.assertEqual(json.loads(json.dumps(rows)), proof['rows'])
        self.assertFalse(proof['passed']); self.assertTrue(proof['content_contract_passed'])
        self.assertEqual(visual_gate(rows), {key: proof[key] for key in ['visual_gate_passed', 'visual_blockers']})
        self.assertEqual(proof['visual_blockers'], [dict(case='ethics-long', language='chinese', profile='submission',
            engine='native_pdf', issue='page-count-increase', before_pages=7, after_pages=8)])
        fixed = next(r for r in rows if (r['case'], r['language'], r['profile']) == ('ethics-missing', 'chinese', 'submission'))
        self.assertFalse(fixed['rendered']['native_pdf']['lead_placement']['before']['together'])
        self.assertTrue(fixed['rendered']['native_pdf']['lead_placement']['after']['together'])
        issues = proof['known_preexisting_visual_issues']; self.assertEqual(len(issues), 1)
        self.assertEqual((issues[0]['case'], issues[0]['language'], issues[0]['profile'], issues[0]['engine']),
                         ('ethics-long', 'chinese', 'review', 'word_preview'))
        self.assertEqual(issues[0]['before'], issues[0]['after'])

    def test_exact_recipe_assets_and_failed_spacing_attempt_are_retained(self):
        members = json.loads((ARCHIVE / 'provenance/package-members.json').read_text())
        for language in ['english', 'chinese']:
            old = json.loads((ARCHIVE / 'before-core' / (language + '.json')).read_text())
            new = json.loads((ARCHIVE / 'after' / (language + '.json')).read_text())
            self.assertEqual(old, baseline(language)); expected, operations = patch(old, language); self.assertEqual(new, expected)
            prototype = json.loads((ARCHIVE / 'after/prototype.json').read_text())
            self.assertEqual(operations, prototype['packages'][language + '.zip']['operations'])
            a, b, c = [members[p][language] for p in ['before-core', 'before-long', 'after']]
            self.assertEqual(a, b); self.assertEqual(a.keys(), c.keys())
            self.assertEqual({n for n in a if a[n] != c[n]}, {'template/template.json'})
            prior = copy.deepcopy(new)
            file = next(f for f in prior['files'] if f['fileName'] == TARGET)
            release = next(op for op in operations[TARGET] if op['kind'] == 'release-final-name-only-item')
            file['content'] = file['content'].replace(release['after'], release['before'])
            self.assertEqual(prior, json.loads((ARCHIVE / 'failed-spacing-attempt' / (language + '.json')).read_text()))
            file['content'] = file['content'].replace('<p class="answer-lead" style="margin-bottom: 0">', '<p class="answer-lead">')
            self.assertEqual(prior, json.loads((ARCHIVE / 'failed-attempt' / (language + '.json')).read_text()))
        failed = json.loads((ARCHIVE / 'failed-attempt/native-2.json').read_text())
        self.assertFalse(failed['passed']); self.assertIn('More pages', failed['failure'])

    def test_scoped_cleanup_and_both_temporary_allowances_restored(self):
        for name in ['before-core', 'before-long', 'after', 'failed-attempt', 'failed-spacing-attempt']:
            cleanup = json.loads((ARCHIVE / name / 'owned-test-template-cleanup.json').read_text())
            self.assertTrue(cleanup['run_completed_successfully']); self.assertEqual(len(cleanup['deleted']), 2)
            self.assertEqual((cleanup['project_references'], cleanup['document_references']), (0, 0))
        for name in ['storage-allowance', 'first-storage-allowance', 'second-storage-allowance']:
            allowance = json.loads((ARCHIVE / 'provenance' / (name + '.json')).read_text())
            self.assertTrue(allowance['restored']); self.assertEqual(allowance['before'], allowance['after'])
            self.assertEqual(allowance['during'], dict(allowance['before'], storage=-1750000000))
        life = json.loads((ARCHIVE / 'provenance/worker-lifecycle.json').read_text())
        self.assertTrue(life['stock_worker_restored']); self.assertEqual(len(life['after']), 4)
        self.assertTrue(all(row['status'] == 'exited' for row in life['after']))

    def test_word_delta_rejects_authored_mutation_and_extra_style_change(self):
        paths = [ARCHIVE / phase / 'renders/ethics-missing-submission-chinese.docx' for phase in ['before-core', 'after']]
        proof = json.loads((ARCHIVE / 'provenance/native.json').read_text())
        row = next(r for r in proof['rows'] if (r['case'], r['profile'], r['language']) == ('ethics-missing', 'submission', 'chinese'))
        class Package:
            def __init__(self, values): self.values = values
            def __enter__(self): return self
            def __exit__(self, *args): pass
            def namelist(self): return list(self.values)
            def read(self, name): return self.values[name]
        values = []
        for path in paths:
            with zipfile.ZipFile(path) as z: values.append({n: z.read(n) for n in z.namelist()})
        for mutation in ['authored', 'style', 'asset']:
            changed = copy.deepcopy(values[1])
            if mutation == 'asset': changed['word/styles.xml'] += b' '
            else:
                tree = etree.fromstring(changed['word/document.xml'])
                if mutation == 'authored':
                    node = next(n for n in tree.iter(qn('w:t')) if n.text and 'AUTHORED-PURPOSE:' in n.text)
                    node.text = node.text.replace('Original.csv.', 'Original.csv!')
                else: next(tree.iter(qn('w:pStyle'))).set(qn('w:val'), 'BodyText')
                changed['word/document.xml'] = etree.tostring(tree)
            with mock_patch('lead_native.zipfile.ZipFile', side_effect=[Package(values[0]), Package(changed)]):
                with self.assertRaises(AssertionError): word_lead_delta(*paths, row['changes'][-1]['owned_text'])

    def test_repeated_header_and_long_paragraph_oracles_fail_closed(self):
        stem = 'ethics-long-review-chinese'
        soup = BeautifulSoup((ARCHIVE / 'after/renders' / (stem + '.html')).read_text(), 'html.parser')
        document = Document(ARCHIVE / 'after/renders' / (stem + '.docx'))
        pdf = ARCHIVE / 'after/word-preview' / (stem + '.pdf')
        for table in document.tables:
            if table.rows[0].cells[0].text == 'Item':
                table.rows[0]._tr.find(qn('w:trPr') + '/' + qn('w:tblHeader')).set(qn('w:val'), 'off')
        with self.assertRaises(AssertionError): visible_with_table_header(pdf, soup, document, True, 'ethics-long', 'review')
        document = Document(ARCHIVE / 'after/renders' / (stem + '.docx'))
        bbox = subprocess.check_output(['pdftotext', '-bbox-layout', str(pdf), '-'])
        with mock_patch('lead_native.subprocess.check_output', return_value=bbox.replace(b'>Value<', b'>VALUE<', 1)):
            with self.assertRaises(AssertionError): visible_with_table_header(pdf, soup, document, True, 'ethics-long', 'review')
        start = 'The purpose of processing the personal data can be described as follows: AUTHORED-PURPOSE: N/A / 0 / Original.csv.'
        for text in [start, start + ' END-OF-LONG-FIRST-PARAGRAPH.', start + '\fEND-OF-LONG-FIRST-PARAGRAPH!' ]:
            with mock_patch('lead_native.subprocess.check_output', return_value=text):
                with self.assertRaises(AssertionError): long_paragraph(Path('synthetic.pdf'), 'english')


if __name__ == '__main__': unittest.main()
