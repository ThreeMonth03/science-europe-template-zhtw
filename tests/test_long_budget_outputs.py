import copy
import sys
import unittest
from pathlib import Path
from bs4 import BeautifulSoup
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from check_long_budget_outputs import compare_long_table,check_long_pages


def fixture(count=12):
    doc=Document(); table=doc.add_table(rows=3,cols=3)
    prop=OxmlElement('w:tblStyle'); prop.set(qn('w:val'),'Table'); table._tbl.tblPr.insert(0,prop)
    for c,t in zip(table.rows[0].cells,['Purpose','Amount','Funding']): c.text=t
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    table.cell(1,0).text='Resource'; table.cell(1,0).add_paragraph('Original purpose.')
    for i in range(count): table.cell(1,0).add_paragraph(f'BUDGET-PARA-{i+1:02d}: Original.')
    table.cell(1,1).text='0 TWD'; table.cell(1,2).text='Institute'
    table.cell(2,0).text='Second resource'; table.cell(2,1).text='5000 TWD'; table.cell(2,2).text='Other funding'
    return doc,table


def transformed(table):
    old=table._tbl; long=copy.deepcopy(old); tail=copy.deepcopy(old)
    long.tblPr.find(qn('w:tblStyle')).set(qn('w:val'),'PilotLongBudget')
    rows=long.findall(qn('w:tr')); long.remove(rows[2]); identity=rows[1]
    identity.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    first=identity.find(qn('w:tc')); paras=first.findall(qn('w:p'))
    for para in paras[1:]:
        first.remove(para); row=OxmlElement('w:tr'); cell=OxmlElement('w:tc'); props=OxmlElement('w:tcPr')
        span=OxmlElement('w:gridSpan'); span.set(qn('w:val'),'3'); props.append(span); cell.append(props)
        cell.append(para); row.append(cell); long.append(row)
    tail.remove(tail.findall(qn('w:tr'))[1]); return old,long,tail


class LongBudgetOutputTests(unittest.TestCase):
    def test_original_paragraphs_and_zero_survive(self):
        _,t=fixture(); self.assertEqual(13,compare_long_table(*transformed(t)))

    def test_amount_run_and_tail_mutations_fail(self):
        for kind in ['amount','purpose','tail']:
            _,t=fixture(); old,long,tail=transformed(t)
            if kind=='amount': node=long.findall(qn('w:tr'))[1].findall(qn('w:tc'))[1]
            elif kind=='purpose': node=long.findall(qn('w:tr'))[2]
            else: node=tail.findall(qn('w:tr'))[1]
            next(node.iter(qn('w:t'))).text='Changed'
            with self.assertRaises(AssertionError): compare_long_table(old,long,tail)

    def test_full_width_and_repeating_headers_are_required(self):
        for kind in ['span','header']:
            _,t=fixture(); old,long,tail=transformed(t)
            node=long.find('.//'+qn('w:gridSpan')) if kind=='span' else long.findall(qn('w:tr'))[1].find(qn('w:trPr')+'/'+qn('w:tblHeader'))
            node.getparent().remove(node)
            with self.assertRaises(AssertionError): compare_long_table(old,long,tail)

    def test_page_oracle_rejects_lost_identity_or_split_purpose(self):
        doc,t=fixture(60); soup=BeautifulSoup('<div id="q-required-resources"><h4>Budget</h4></div>','html.parser')
        header='PurposeAmountFundingResource0TWDInstitute'
        purpose=[p.text.replace(' ','') for p in t.cell(1,0).paragraphs[2:]]
        pages=['Budget'+header+'Originalpurpose.'+''.join(purpose[:30]),header+''.join(purpose[30:])]
        self.assertEqual([1,2],check_long_pages(pages,doc,soup)['purpose_pages'])
        for broken in [[pages[0],pages[1].replace('0TWD','')],[pages[0],pages[1].replace(purpose[30],'')],['Budget',pages[0].replace('Budget',''),pages[1]]]:
            with self.assertRaises(AssertionError): check_long_pages(broken,doc,soup)
