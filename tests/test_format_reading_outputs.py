import sys
import unittest
from pathlib import Path
from bs4 import BeautifulSoup
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
from check_format_reading_outputs import verified_format_markers,word_delta


class FormatReadingOutputTests(unittest.TestCase):
    def test_exact_cjk_hint_is_plain_but_font_override_is_not(self):
        def doc(text):
            d=Document();d.add_heading('1. First',3);d.add_heading('2. Formats',3)
            run=d.add_paragraph().add_run(text);fonts=OxmlElement('w:rFonts');fonts.set(qn('w:hint'),'eastAsia')
            run._element.get_or_add_rPr().append(fonts);d.add_heading('3. Metadata',3);return d
        a,b=doc('原句。'),doc('新句。');pairs=[('原句。','新句。')]
        self.assertEqual(word_delta(a,b,pairs),1)
        b.paragraphs[2].runs[0]._element.rPr.rFonts.set(qn('w:eastAsia'),'Unreviewed Font')
        with self.assertRaises(AssertionError):word_delta(a,b,pairs)

    def test_mixed_markers_must_have_text_geometry_and_exact_type_count(self):
        soup=BeautifulSoup('<ul><li>First<ul><li>Second</li></ul></li></ul>','html.parser')
        box=b'<html><page><line><word xMin="10">\xe2\x80\xa2</word><word>First</word></line><line><word xMin="20">\xe2\x97\xa6</word><word>Second</word></line></page></html>'
        clean,signatures=verified_format_markers(['FirstSecond•◦'],box,soup,0)
        self.assertEqual(clean,['FirstSecond']);self.assertEqual(sum(signatures.values()),2)
        for pages,bbox,html in [(['FirstSecond••'],box,soup),(['First•Second◦'],box,soup),
             (['FirstSecond•◦'],box.replace(b'<word>Second</word>',b''),soup),
             (['FirstSecond•◦'],box,BeautifulSoup('<p>Author • text</p>','html.parser'))]:
            with self.assertRaises(AssertionError):verified_format_markers(pages,bbox,html,0)

    def test_only_approved_word_summary_changes_not_authored_text_or_styles(self):
        def document(value):
            doc=Document();doc.add_heading('1. First',3);doc.add_heading('2. Formats',3)
            p=doc.add_paragraph();p.add_run('Original.csv').bold=True;p.add_run(' '+value)
            doc.add_paragraph('Keep authored!');doc.add_heading('3. Metadata',3)
            return doc
        before=document('We expect 0 files in this format.')
        after=document('We expect 0 files.')
        pairs=[('Original.csv We expect 0 files in this format.','Original.csv We expect 0 files.')]
        self.assertEqual(word_delta(before,after,pairs),1)
        for kind in ['authored','style','quantity','name']:
            changed=document('We expect 0 files.')
            if kind=='authored':changed.paragraphs[3].text='Edited!'
            if kind=='style':changed.paragraphs[2].runs[0].bold=False
            if kind=='quantity':changed.paragraphs[2].runs[1].text=' We expect 1 files.'
            if kind=='name':changed.paragraphs[2].runs[0].text='original.csv'
            with self.assertRaises(AssertionError):word_delta(before,changed,pairs)
