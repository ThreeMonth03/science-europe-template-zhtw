from io import BytesIO
import sys
import unittest
from pathlib import Path
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT/'scripts'))
from check_storage_gap_outputs import word_delta


def document(capacity, tail=None):
    doc = Document()
    doc.styles.add_style('First Paragraph', WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph('1. First', 'Heading 3')
    doc.add_paragraph('Original.csv')
    doc.add_paragraph('3. Metadata', 'Heading 3')
    doc.add_paragraph('Metadata is open.', 'First Paragraph')
    doc.add_paragraph('Storage', 'Heading 4')
    if capacity is not None: doc.add_paragraph(capacity, 'First Paragraph')
    if tail is not None: doc.add_paragraph(tail, 'Body Text')
    doc.add_paragraph('4. Quality', 'Heading 3')
    doc.add_paragraph('CHANGELOG.md')
    return doc


class StorageGapOutputTests(unittest.TestCase):
    def test_insert_prompt_does_not_admit_missing_or_authored_changes(self):
        old = document(None); new = document('Missing capacity (GB).')
        self.assertEqual(word_delta(old, new, ('', 'Missing capacity (GB).'), True), 1)
        for index, value in [(1, 'original.csv'), (5, 'Missing capacity.'), (7, 'changelog.md')]:
            mutation = clone(new); mutation.paragraphs[index].text = value
            with self.assertRaises(AssertionError): word_delta(old, mutation, ('', 'Missing capacity (GB).'), True)

    def test_blank_capacity_splits_owned_policy_and_retains_styles(self):
        old = document('Estimate to GB. File policy.')
        new = document('Missing capacity (GB).', 'File policy.')
        self.assertEqual(word_delta(old, new, ('Estimate to GB.', 'Missing capacity (GB).'), True), 1)
        mutation = clone(new); mutation.paragraphs[6].style = 'Title'
        with self.assertRaises(AssertionError): word_delta(old, mutation, ('Estimate to GB.', 'Missing capacity (GB).'), True)
        mutation = clone(new); mutation.paragraphs[5].runs[0].bold = True
        with self.assertRaises(AssertionError): word_delta(old, mutation, ('Estimate to GB.', 'Missing capacity (GB).'), True)

    def test_complete_zero_retains_quantity_and_retained_character_formatting(self):
        old = document('Estimate to 0 GB. File policy.')
        new = document('Estimate to be 0 GB. File policy.')
        pairs = ('Estimate to 0 GB.', 'Estimate to be 0 GB.')
        self.assertEqual(word_delta(old, new, pairs, False), 1)
        for text in ['Estimate to be 1 GB. File policy.', 'Estimate to be 0 GB. Changed policy.']:
            mutation = clone(new); mutation.paragraphs[5].text = text
            with self.assertRaises(AssertionError): word_delta(old, mutation, pairs, False)
        mutation = clone(new); mutation.paragraphs[5].runs[0].italic = True
        with self.assertRaises(AssertionError): word_delta(old, mutation, pairs, False)


def clone(doc):
    # deepcopy(Document) can detach the cached paragraph body from element.body.
    # Round-trip through actual DOCX so mutations affect the checked XML.
    stream = BytesIO(); doc.save(stream); stream.seek(0)
    return Document(stream)
