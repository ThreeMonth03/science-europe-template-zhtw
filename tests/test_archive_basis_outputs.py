from io import BytesIO
import sys
import unittest
from pathlib import Path
from bs4 import BeautifulSoup
from docx import Document
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
from check_archive_basis_outputs import without_continuation_headers,word_delta


class ArchiveBasisOutputTests(unittest.TestCase):
    def test_only_verified_page_prefix_repetitions_are_excluded(self):
        soup=BeautifulSoup('<div id="q-required-resources"><h4>Budget</h4><table class="resource-table"><thead><tr><th>Resource</th><th>Amount</th><th>Funding</th></tr></thead><tbody><tr><td><p>First</p><div class="answer-detail"><p>BUDGET-PARA-1.</p><p>BUDGET-PARA-2.</p></div></td><td>5000</td><td>Grant</td></tr><tr><td>Second</td><td>0</td><td>None</td></tr></tbody></table></div>','html.parser')
        prefix='ResourceAmountFundingFirst5000Grant'
        pages=['BeforeBudget'+prefix+'BUDGET-PARA-1.',prefix+'BUDGET-PARA-2.ResourceAmountFundingSecond0None']
        cleaned,count=without_continuation_headers(pages,soup)
        self.assertEqual(count,1);self.assertEqual(cleaned,[pages[0],pages[1][len(prefix):]])
        for replacement in ['4999','50000','']:
            bad=pages.copy();bad[1]=bad[1].replace('5000',replacement)
            with self.assertRaises(AssertionError):without_continuation_headers(bad,soup)
        for replacement in ['BUDGET-PARA-3.','BUDGET-PARA-2..','']:
            bad=pages.copy();bad[1]=bad[1].replace('BUDGET-PARA-2.',replacement)
            with self.assertRaises(AssertionError):without_continuation_headers(bad,soup)

    def test_word_control_does_not_admit_answer_punctuation_or_format_changes(self):
        old=Document();old.add_paragraph('1. First question');old.add_paragraph('Keep MyFile.csv.')
        def clone():
            buffer=BytesIO();old.save(buffer);buffer.seek(0);return Document(buffer)
        self.assertFalse(word_delta(old,clone(),[], '')['new_paragraphs'])
        for text in ['Keep MyFile.csv..','Keep myfile.csv.','']:
            new=clone();new.paragraphs[1].text=text
            with self.assertRaises(AssertionError):word_delta(old,new,[], '')
        new=clone();new.paragraphs[1].runs[0].bold=True
        with self.assertRaises(AssertionError):word_delta(old,new,[], '')
