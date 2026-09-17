from io import BytesIO
import sys
import unittest
from pathlib import Path
from bs4 import BeautifulSoup
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT/'scripts'))
from check_metadata_followup_outputs import word_delta, metadata_blocks


def document(new):
    doc = Document(); doc.styles.add_style('First Paragraph', WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph('1. First', 'Heading 3'); doc.add_paragraph('Original.csv')
    doc.add_paragraph('3. Metadata', 'Heading 3')
    doc.add_paragraph('Not public.' if not new else 'Dictionary: no. Not public.', 'First Paragraph')
    p = doc.add_paragraph('Keep ', 'Body Text'); p.add_run('Original.csv').bold = True
    if new: doc.add_paragraph('Missing instructions.', 'Body Text')
    doc.add_paragraph('Storage', 'Heading 4'); doc.add_paragraph('0 GB.', 'First Paragraph')
    doc.add_paragraph('4. Quality', 'Heading 3'); doc.add_paragraph('CHANGELOG.md')
    return doc


def clone(doc):
    stream = BytesIO(); doc.save(stream); stream.seek(0); return Document(stream)


OLD = [('fixed', 'Not public.'), ('authored', 'Keep Original.csv')]
NEW = [('fixed', 'Dictionary: no. Not public.'), ('authored', 'Keep Original.csv'), ('fixed', 'Missing instructions.')]


class MetadataOutputTests(unittest.TestCase):
    def test_html_boundaries_are_not_arbitrarily_flattened(self):
        html = '<div id="q-docs-metadata"><div class="answer"><div class="metadata-policy"><p>Standard.</p><p>Dictionary.</p><div class="reading-gap"><p>Missing A.</p><p>Missing B.</p></div><p>Public.</p><div class="answer-detail" data-fact-id="metadata-access-explanation"><p>Keep <b>Original.csv</b>.</p><ul><li>One.</li><li>Two.</li></ul></div></div></div></div>'
        self.assertEqual(metadata_blocks(BeautifulSoup(html, 'html.parser')),
            [('fixed','Standard. Dictionary.'),('fixed','Missing A.'),('fixed','Missing B.'),('fixed','Public.'),
             ('authored','Keep Original.csv.'),('authored','One.'),('authored','Two.')])

    def test_exact_owned_delta_retains_authored_and_outside_xml(self):
        old, new = document(False), document(True)
        self.assertEqual(word_delta(old, new, OLD, NEW)['authored_blocks'], 1)
        for index, text in [(1,'original.csv'),(3,'Dictionary: yes. Not public.'),(4,'Keep original.csv'),
                            (5,'Missing instructions'),(7,'1 GB.'),(9,'changelog.md')]:
            altered = clone(new); altered.paragraphs[index].text = text
            with self.assertRaises(AssertionError): word_delta(old, altered, OLD, NEW)

    def test_lost_prompt_and_styled_changes_rejected(self):
        old, new = document(False), document(True)
        altered = clone(new); p = altered.paragraphs[5]._p; p.getparent().remove(p)
        with self.assertRaises(AssertionError): word_delta(old, altered, OLD, NEW)
        for index in [3,4,5]:
            altered = clone(new); altered.paragraphs[index].runs[0].italic = True
            with self.assertRaises(AssertionError): word_delta(old, altered, OLD, NEW)
        altered = clone(new); altered.paragraphs[5].style = 'Title'
        with self.assertRaises(AssertionError): word_delta(old, altered, OLD, NEW)
