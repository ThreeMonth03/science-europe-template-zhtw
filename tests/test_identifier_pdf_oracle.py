import sys
from pathlib import Path
import unittest
from bs4 import BeautifulSoup
from docx import Document
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from check_identifier_concise_outputs import verified_list_suffixes,compare_word


class IdentifierPdfOracle(unittest.TestCase):
    def test_word_oracle_allows_only_the_owned_leading_sentence(self):
        def document(changed=False):
            doc=Document();doc.add_paragraph('1. First','Heading 3')
            doc.add_paragraph('Keep before.');doc.add_paragraph('13. Identifiers','Heading 3')
            doc.add_paragraph('Repository assigns.' if changed else 'Assigned. Repository assigns.')
            doc.add_paragraph('14. Responsibilities','Heading 3');doc.add_paragraph('Keep after.')
            return doc
        old,new=document(),document(True);pairs=[('Assigned. Repository assigns.','Repository assigns.')]
        self.assertEqual(compare_word(old,new,pairs),1)
        for change in ['text','style','outside','format']:
            bad=document(True)
            if change=='text':bad.paragraphs[3].text='Someone else assigns.'
            elif change=='style':bad.paragraphs[3].style='Heading 4'
            elif change=='outside':bad.paragraphs[-1].text='Changed after.'
            else:bad.paragraphs[3].runs[0].bold=True
            with self.assertRaises(AssertionError):compare_word(old,bad,pairs)

    def test_earlier_page_markers_are_not_normalized(self):
        soup = BeautifulSoup('<p>Keep.</p>','html.parser')
        bbox = b'<doc><page/><page/></doc>'
        clean, markers = verified_list_suffixes(['Earlier.\u2022\u25e6','Keep.'],bbox,soup,start=1)
        self.assertEqual(clean,['Earlier.\u2022\u25e6','Keep.'])
        self.assertFalse(markers)

    def test_only_verified_suffix_is_removed_and_item_indentation_recorded(self):
        soup = BeautifulSoup('<ul><li>Keep 0 TWD.</li></ul>','html.parser')
        bbox = b'<doc><page><line><word xMin="55.5">\xe2\x80\xa2</word><word>Keep 0 TWD.</word></line></page></doc>'
        clean, markers = verified_list_suffixes(['Keep0TWD.Other.\u2022'],bbox,soup)
        self.assertEqual(clean,['Keep0TWD.Other.'])
        self.assertEqual(markers,{('\u2022Keep0TWD.',55.5):1})

    def test_authored_or_unpositioned_or_non_suffix_bullets_are_rejected(self):
        soup = BeautifulSoup('<p>Keep.</p>','html.parser')
        bbox = b'<doc><page><line><word xMin="55.5">\xe2\x80\xa2</word><word>Keep.</word></line></page></doc>'
        for pages, geometry, html in [(['Keep.\u2022'],bbox,BeautifulSoup('<p>Authored \u2022</p>','html.parser')),
                (['Keep.\u2022\u2022'],bbox,soup), (['Keep.\u2022After.'],bbox,soup),
                (['Keep.\u2022'],bbox.replace(b'<word>Keep.</word>',b''),soup)]:
            with self.assertRaises(AssertionError): verified_list_suffixes(pages,geometry,html)


if __name__ == '__main__': unittest.main()
