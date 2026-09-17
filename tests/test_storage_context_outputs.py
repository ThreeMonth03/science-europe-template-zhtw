from io import BytesIO
import sys
import unittest
from pathlib import Path
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
ROOT=Path(__file__).resolve().parents[1]
# The local checkout and Actions deliberately use different directory names.
english_scripts=next(path for path in [ROOT.parent/'english/scripts',ROOT.parent/'science-europe-template/scripts']
                     if (path/'probe_storage_context.py').is_file())
sys.path[:0]=[str(ROOT/'scripts'),str(english_scripts)]
from check_storage_context_outputs import check_word


def doc(new):
    value=Document()
    for name in ['First Paragraph','Pilot Lead','Pilot List Lead','Compact']:
        if name not in value.styles:value.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    value.add_paragraph('1. First','Heading 3');value.add_paragraph('Original.csv')
    value.add_paragraph('5. Storage?','Heading 3')
    value.add_paragraph('Keep 0 GB.','Pilot Lead' if new else 'First Paragraph')
    value.add_paragraph('Limitations:','Pilot Lead' if new else 'Body Text')
    value.add_paragraph('Location unknown.','Pilot List Lead')
    value.add_paragraph('Schedule unknown.','Compact')
    value.add_paragraph('6. Security?','Heading 3');value.add_paragraph('CHANGELOG.md')
    return value


def clone(d):
    stream=BytesIO();d.save(stream);stream.seek(0);return Document(stream)


class StorageContextOutputTests(unittest.TestCase):
    def test_only_two_fixed_paragraph_styles_change(self):
        old,new=doc(False),doc(True)
        self.assertEqual(check_word(old,new,True),2)
        for i,text in [(1,'original.csv'),(3,'Keep 1 GB.'),(4,'Limitations'),(5,'Location known.'),(8,'changelog.md')]:
            changed=clone(new);changed.paragraphs[i].text=text
            with self.assertRaises(AssertionError):check_word(old,changed,True)
        for index in [2,6,7]:
            changed=clone(new);changed.paragraphs[index].style='Pilot Lead'
            with self.assertRaises(AssertionError):check_word(old,changed,True)
        changed=clone(new);changed.paragraphs[3].runs[0].bold=True
        with self.assertRaises(AssertionError):check_word(old,changed,True)

    def test_fallback_cannot_gain_style_or_lose_content(self):
        old=doc(False);self.assertEqual(check_word(old,clone(old),False),0)
        with self.assertRaises(AssertionError):check_word(old,doc(True),False)
