import sys
from pathlib import Path
import unittest
from bs4 import BeautifulSoup
from docx import Document
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
from check_word_short_budget_outputs import word_pages,paragraph_texts


class WordShortOutputTests(unittest.TestCase):
    def test_only_owned_complete_date_runs_get_extraction_mapping(self):
        doc=Document();doc.add_paragraph('1. First question')
        p=doc.add_paragraph('Date: ');p.add_run('2027\u201112\u201131');p.add_run('. File Selection-2027-12-31.csv.')
        soup=BeautifulSoup('<span class="date-value">2027-12-31</span>','html.parser')
        self.assertEqual(paragraph_texts(doc,soup)[1],'Date: 2027-12-31. File Selection-2027-12-31.csv.')
        with self.assertRaises(AssertionError):paragraph_texts(doc,BeautifulSoup(str(soup)+str(soup),'html.parser'))
        p.add_run('Duplicate-2027\u201112\u201131.csv')
        with self.assertRaises(AssertionError):paragraph_texts(doc,soup)

    def test_cover_without_footer_and_footer_before_body(self):
        raw='Cover\f2 / 2\nKeep 0 TWD.\f'
        bbox=b'<doc><page height="800"/><page height="800"><line yMin="760">2 / 2</line><line yMin="40">Keep 0 TWD.</line></page></doc>'
        self.assertEqual(word_pages(raw,bbox),['Cover','Keep0TWD.'])

    def test_nonbottom_duplicate_or_missing_footer_is_not_ignored(self):
        raw='Cover\f2 / 2\nKeep 0 TWD.\f'
        bbox=b'<doc><page height="800"/><page height="800"><line yMin="760">2 / 2</line></page></doc>'
        for bad in [bbox.replace(b'760',b'200'),bbox.replace(b'</page></doc>',b'<line yMin="760">2 / 2</line></page></doc>'),bbox.replace(b'2 / 2',b'3 / 3')]:
            with self.assertRaises(AssertionError):word_pages(raw,bad)
        with self.assertRaises(AssertionError):word_pages(raw.replace('Keep','2 / 2\nKeep'),bbox)


if __name__=='__main__':unittest.main()
