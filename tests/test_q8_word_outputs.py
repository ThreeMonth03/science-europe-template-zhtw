import copy
import sys
import unittest
from pathlib import Path
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from check_q8_word_outputs import compare_word


def document():
    doc=Document()
    for name in ['Compact','Pilot List Lead']: doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph('1. First question',style='Heading 3');doc.add_paragraph('Retain original.')
    doc.add_paragraph('8. Ownership?',style='Heading 3');doc.add_paragraph('Name.csv',style='Compact')
    doc.add_paragraph('Permission.',style='Compact');doc.add_paragraph('9. Ethics?',style='Heading 3')
    return doc


class Q8WordOutputTests(unittest.TestCase):
    def test_only_owned_label_style_is_allowed(self):
        before,after=document(),document();after.paragraphs[3].style='Pilot List Lead'
        self.assertEqual(1,compare_word(before,after,['Name.csv']))

    def test_other_question_edit_is_rejected(self):
        before,after=document(),document();after.paragraphs[1].style='Pilot List Lead'
        with self.assertRaises(AssertionError):compare_word(before,after,[])

    def test_permission_edit_is_rejected(self):
        before,after=document(),document();after.paragraphs[4].style='Pilot List Lead'
        with self.assertRaises(AssertionError):compare_word(before,after,['Name.csv'])

    def test_text_change_is_rejected(self):
        before,after=document(),document();after.paragraphs[3].text='changed.csv'
        with self.assertRaises(AssertionError):compare_word(before,after,['Name.csv'])

    def test_ignored_style_is_not_success(self):
        with self.assertRaises(AssertionError):compare_word(document(),document(),['Name.csv'])

    def test_empty_control_stays_unchanged(self):
        self.assertEqual(0,compare_word(document(),document(),[]))
