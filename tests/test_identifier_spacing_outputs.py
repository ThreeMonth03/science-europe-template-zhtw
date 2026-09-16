import sys
from pathlib import Path
import unittest
from unittest.mock import patch
from bs4 import BeautifulSoup
from docx import Document

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from check_identifier_spacing_outputs import pairs,word_delta,pdf_delta


def word(value):
    d=Document();d.add_paragraph('1. First','Heading 3')
    d.add_paragraph('Unchanged. A-B / 0.05 / DOI 10.1000/test')
    d.add_paragraph('13. Identifiers','Heading 3');d.add_paragraph(value)
    d.add_paragraph('14. Responsibilities','Heading 3');d.add_paragraph('Original answer.')
    return d


class IdentifierSpacingOutputTests(unittest.TestCase):
    def test_only_one_planned_separator(self):
        old=word('指派。 解析。');new=word('指派。解析。')
        self.assertEqual(1,word_delta(old,new,[('指派。 解析。','指派。解析。',3)]))
        for changed in ['指派解析。','指派。 解析。','指派。解析。 ', '指派。保證解析。']:
            with self.assertRaises(AssertionError):word_delta(old,word(changed),[('指派。 解析。','指派。解析。',3)])

    def test_unrelated_text_style_and_existing_spaces_are_protected(self):
        old=word('指派。 解析。');planned=[('指派。 解析。','指派。解析。',3)]
        for change in ['before','after','style','internal-space']:
            current=word('指派。解析。')
            if change=='before':current.paragraphs[1].text='Changed.'
            if change=='after':current.paragraphs[-1].text='Changed.'
            if change=='style':current.paragraphs[3].runs[0].bold=True
            if change=='internal-space':current.paragraphs[3].text='指 派。解析。'
            with self.assertRaises(AssertionError):word_delta(old,current,planned)

    def test_only_paired_chinese_policy_has_a_removal_plan(self):
        html='<div id="q-persistent-identifier"><div class="identifier-arrangement dataset-policy"><p>資料將取得持續識別碼。</p><p>資料儲存庫將不保證解析。</p></div><div class="answer-detail"><p>保留。 空格。</p></div></div>'
        soup=BeautifulSoup(html,'html.parser')
        self.assertEqual(1,len(pairs(soup,'chinese')))
        self.assertEqual([],pairs(soup,'english'))
        soup.select_one('.identifier-arrangement p').decompose()
        self.assertEqual([],pairs(soup,'chinese'))

    def test_pdf_oracle_preserves_other_spaces_and_punctuation(self):
        soup=BeautifulSoup('<div class="question"><h3>1. First</h3></div><div id="q-persistent-identifier"><h3>13. Identifiers</h3></div><div id="q-dm-responsible"><h3>14. Roles</h3></div>','html.parser')
        left='1. FirstAuthor。 Keep  two spaces.13. Identifiers指派。 解析。14. RolesOriginal.'
        right=left.replace('指派。 解析。','指派。解析。')
        planned=[('指派。 解析。','指派。解析。',3)]
        with patch('check_identifier_spacing_outputs.pdf_reading_text',side_effect=[[left],[right]]):
            pdf_delta(None,None,soup,planned)
        for mutation in [right.replace('Keep  two','Keep two'),right.replace('指派。','指派'),right.replace('Original.','Changed.'),left]:
            with patch('check_identifier_spacing_outputs.pdf_reading_text',side_effect=[[left],[mutation]]):
                with self.assertRaises(AssertionError):pdf_delta(None,None,soup,planned)


if __name__=='__main__':unittest.main()
