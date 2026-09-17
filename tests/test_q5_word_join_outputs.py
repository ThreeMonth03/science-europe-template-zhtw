import copy
from pathlib import Path
import sys
import unittest
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
english = next(p for p in [ROOT.parent / 'english/scripts', ROOT.parent / 'science-europe-template/scripts']
               if (p / 'q5_word_join_contract.py').is_file())
sys.path[:0] = [str(ROOT / 'scripts'), str(ROOT / 'tests'), str(english)]
from check_q5_word_join_outputs import check_word
from test_storage_context_outputs import doc, clone


def joined():
    old = doc(True); new = clone(old)
    policy, intro = new.paragraphs[3:5]
    r = OxmlElement('w:r'); r.append(OxmlElement('w:br')); policy._p.append(r)
    for child in list(intro._p)[1:]: policy._p.append(copy.deepcopy(child))
    intro._p.getparent().remove(intro._p)
    return old, new


class Q5WordJoinOutputTests(unittest.TestCase):
    def test_only_adjacent_owned_q5_leads_join(self):
        old, new = joined()
        self.assertEqual(check_word(old, new, True), 1)
        for index in (1, 3, 4, 5, 7):
            changed = clone(new); changed.paragraphs[index].runs[0].text += '!'
            with self.assertRaises(AssertionError): check_word(old, changed, True)

    def test_wrong_final_keep_heading_or_format_fails(self):
        old, new = joined()
        for index in (2, 5, 6):
            changed = clone(new); changed.paragraphs[index].style = 'Pilot Lead'
            with self.assertRaises(AssertionError): check_word(old, changed, True)
        changed = clone(new); changed.paragraphs[3].runs[0].bold = True
        with self.assertRaises(AssertionError): check_word(old, changed, True)

    def test_fallback_cannot_join_or_be_restylized(self):
        old, new = joined()
        self.assertEqual(check_word(old, clone(old), False), 0)
        with self.assertRaises(AssertionError): check_word(old, new, False)
        with self.assertRaises(AssertionError): check_word(old, old, True)


if __name__ == '__main__': unittest.main()
