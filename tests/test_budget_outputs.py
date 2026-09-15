import io
import sys
import unittest
from pathlib import Path
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from check_budget_outputs import compare_word


class BudgetOutputTests(unittest.TestCase):
    def pair(self):
        old=Document(); old.styles.add_style('Pilot Lead',WD_STYLE_TYPE.PARAGRAPH)
        old.add_heading('1. Data',level=3); old.add_paragraph('Author.csv.',style='Body Text')
        old.add_heading('15. Resources',level=3); old.add_paragraph('Overview.',style='Body Text')
        old.add_paragraph('Charges.',style='Body Text'); old.add_heading('Budget',level=4)
        table=old.add_table(rows=1,cols=3)
        for c,t in zip(table.rows[0].cells,['Purpose.','0 TWD','Funder.']): c.text=t
        buffer=io.BytesIO(); old.save(buffer); buffer.seek(0); new=Document(buffer)
        new.paragraphs[4].style='Pilot Lead'
        return old,new

    def test_only_q15_overview_style_changes(self):
        old,new=self.pair(); self.assertEqual(1,compare_word(old,new,True))
        with self.assertRaises(AssertionError): compare_word(old,new,False)

    def test_preserve_text_and_inline_formatting(self):
        for index in [1,3,4]:
            old,new=self.pair(); new.paragraphs[index].runs[0].bold=True
            with self.assertRaises(AssertionError): compare_word(old,new,True)

    def test_table_content_and_style_are_not_excluded(self):
        for kind in ['amount','style','row']:
            old,new=self.pair()
            if kind=='amount': new.tables[0].cell(0,1).text='5000 TWD'
            elif kind=='style': new.tables[0].cell(0,0).paragraphs[0].style='Pilot Lead'
            else: new.tables[0].add_row()
            with self.assertRaises(AssertionError): compare_word(old,new,True)

    def test_budget_fixture_keeps_stock_table_blockers(self):
        from run_pilot import TABLE_CASES
        self.assertTrue({'budget-long','budget-many'}<=TABLE_CASES)
