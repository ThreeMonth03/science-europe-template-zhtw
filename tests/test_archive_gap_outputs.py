from io import BytesIO
from pathlib import Path
import sys
import unittest
from bs4 import BeautifulSoup
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1]/'scripts'))
from check_archive_gap_outputs import body_geometry, pair_delta, prompt_geometry, word_unchanged


def page(*lines):
    return '<page width="600" height="800">'+''.join(
        f'<line xMin="50" yMin="{y}" xMax="150" yMax="{y+10}"><word xMin="50" yMin="{y}" xMax="150" yMax="{y+10}">{text}</word></line>'
        for y, text in lines)+'</page>'


def bbox(*pages):
    return ('<doc>'+''.join(pages)+'</doc>').encode()


class ArchiveGapOutputTests(unittest.TestCase):
    def test_word_control_rejects_deleted_punctuation_or_format(self):
        original = Document(); original.add_paragraph('Cover date')
        original.add_paragraph('1. Question'); original.add_paragraph('Keep MyFile.csv.')
        def clone():
            stream = BytesIO(); original.save(stream); stream.seek(0); return Document(stream)
        new = clone(); new.paragraphs[0].text = 'Different cover date'; word_unchanged(original, new)
        for value in ['Keep MyFile.csv..', 'Keep myfile.csv.', '']:
            new = clone(); new.paragraphs[2].text = value
            with self.assertRaises(AssertionError): word_unchanged(original, new)
        new = clone(); new.paragraphs[2].runs[0].italic = True
        with self.assertRaises(AssertionError): word_unchanged(original, new)

    def test_only_cover_pages_are_excluded_from_geometry(self):
        soup = BeautifulSoup('<div class="question"><h3>1. Question</h3></div>', 'html.parser')
        body = page((50, '1. Question'), (70, '2026-09-16 answer'))
        old = body_geometry(bbox(page((50, '2026-09-16')), body), soup)
        self.assertEqual(old, body_geometry(bbox(page((50, '2026-09-17')), body), soup))
        self.assertNotEqual(old, body_geometry(bbox(page((50, '2026-09-16')),
                            page((50, '1. Question'), (70, '2026-09-17 answer'))), soup))
        self.assertNotEqual(old, body_geometry(bbox(page((50, '2026-09-16')),
                            page((50, '1. Question'), (71, '2026-09-16 answer'))), soup))
        with self.assertRaises(AssertionError): body_geometry(bbox(body, body), soup)

    def test_prompts_must_be_complete_unique_and_same_page(self):
        self.assertEqual(prompt_geometry(bbox(page((50, 'First'), (65, 'prompt.'))), 'First prompt.')['bottom'], 75)
        for data in [bbox(page((50, 'First')), page((65, 'prompt.'))),
                     bbox(page((50, 'First'), (95, 'prompt.'))),
                     bbox(page((50, 'First prompt..'))),
                     bbox(page((50, 'First prompt.'), (90, 'First prompt.')))]:
            with self.assertRaises(AssertionError): prompt_geometry(data, 'First prompt.')

    def test_pair_delta_rejects_split_unchanged_or_changed_width(self):
        soup = BeautifulSoup('<div class="post-project-archive"><p data-fact-id="a">First.</p><p data-fact-id="b">Last.</p></div>', 'html.parser')
        old = bbox(page((50, 'First.'), (100, 'Last.')))
        new = bbox(page((50, 'First.'), (80, 'Last.')))
        self.assertEqual(pair_delta(old, new, soup, [('a', 'b')])[0]['reduced_span_pt'], 20)
        for data in [old, bbox(page((50, 'First.')), page((80, 'Last.'))), new.replace(b'xMax="150"', b'xMax="151"')]:
            with self.assertRaises(AssertionError): pair_delta(old, data, soup, [('a', 'b')])


if __name__ == '__main__': unittest.main()
