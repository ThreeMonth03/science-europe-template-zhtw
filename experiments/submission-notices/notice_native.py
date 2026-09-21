"""Native output checks: unchanged review; prompt-free marked submission, not full approval."""
import argparse
from collections import Counter
import json
from pathlib import Path
import subprocess
import sys
import zipfile
from bs4 import BeautifulSoup
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
sys.path[:0] = [str(ROOT / 'scripts'), str(Path(__file__).parent)]
from artifact_utils import sha
from check_budget_outputs import body, xml
from check_header_controls import raster_digest, long_content
from check_native_mixed_header import page_bounds
from check_budget_spacing_outputs import line_box_overlaps
from rehearse_profile_pdf import snapshot
from rehearse_profile_pagination import geometry
from check_word_short_budget_outputs import paragraph_texts, verify_preview_paragraphs, word_pages
from notice_probe import compare, compact, owned_gaps


def generated_o_markers(document, bbox):
    """Bind the ambiguous letter 'o' to actual DOCX bullet levels and PDF lines."""
    from lxml import etree
    from docx.text.paragraph import Paragraph
    numbering = document.part.numbering_part.element
    nums = {n.get(qn('w:numId')): n.find(qn('w:abstractNumId')).get(qn('w:val')) for n in numbering.findall(qn('w:num'))}
    levels = {(a.get(qn('w:abstractNumId')), l.get(qn('w:ilvl'))): l
              for a in numbering.findall(qn('w:abstractNum')) for l in a.findall(qn('w:lvl'))}
    expected = Counter()
    for node in document.element.body.iter(qn('w:p')):
        num = node.find(qn('w:pPr') + '/' + qn('w:numPr'))
        if num is None: continue
        level = levels[(nums[num.find(qn('w:numId')).get(qn('w:val'))], num.find(qn('w:ilvl')).get(qn('w:val')))]
        if level.find(qn('w:numFmt')).get(qn('w:val')) == 'bullet' and level.find(qn('w:lvlText')).get(qn('w:val')) == 'o':
            expected[compact(Paragraph(node, document).text)] += 1
    total = sum(expected.values())
    for page in etree.fromstring(bbox).findall('.//{*}page'):
        lines = page.findall('.//{*}line')
        for line in lines:
            words = line.findall('{*}word')
            if not words or words[0].text != 'o': continue
            content_words = words[1:]
            if not content_words:
                # Poppler can put the bullet in its own block. Bind only the
                # immediately adjacent line at the same vertical position.
                adjacent = [v for v in lines if v is not line and
                    abs(float(v.get('yMax')) - float(line.get('yMax'))) < 2 and
                    0 < float(v.get('xMin')) - float(line.get('xMax')) < 30]
                assert len(adjacent) == 1, 'Unbound bullet geometry'
                content_words = adjacent[0].findall('{*}word')
            tail = compact(''.join(w.text or '' for w in content_words)); assert tail
            matches = [text for text, count in expected.items() if count and text.startswith(tail)]
            assert len(matches) == 1, 'Unbound or ambiguous letter-o marker; do not erase authored text'
            expected[matches[0]] -= 1
    assert not +expected, 'A generated o marker was not found on its paragraph line'
    return total


def check_visible(pdf, soup, word=False):
    """Exact character inventory, paragraph counts and known repeated long-row headers.

    Paragraph ordering is checked in HTML; table column/long-row display order
    differs across engines. Do not label this a generic PDF semantic parser.
    """
    if word:
        raw = subprocess.check_output(['pdftotext', '-raw', str(pdf), '-'], text=True)
        bbox = subprocess.check_output(['pdftotext', '-bbox-layout', str(pdf), '-'])
        pages = word_pages(raw, bbox)
    else: pages, bbox = snapshot(pdf)
    text = ''.join(pages)
    source = soup.body or soup
    expected = compact(source.get_text())
    markers = {'•', '◦', '\uf0b7', '\uf0a1'}
    assert not markers.intersection(expected), 'Literal markers require a separate oracle'
    value = ''.join(c for c in text if c not in markers)
    counts = Counter(value)
    letter_bullets = 0
    if word:
        document = Document(pdf.parent.parent / 'renders' / (pdf.stem + '.docx'))
        letter_bullets = generated_o_markers(document, bbox)
        counts.subtract({'o': letter_bullets})
    repeated = None
    long = source.select('.resource-table .answer-detail')
    if any('BUDGET-PARA-60:' in node.get_text() for node in long):
        repeated = long_content(pages, source)
        table = source.select_one('.resource-table')
        heading = compact(table.thead.get_text())
        cells = table.select_one('tbody tr').find_all('td', recursive=False)
        identity = compact(cells[0].find('p', recursive=False).get_text() + cells[1].get_text() + cells[2].get_text())
        assert all(heading not in compact(n.get_text()) and identity not in compact(n.get_text()) for n in source.select('.answer-detail'))
        extra_heading = value.count(heading) - expected.count(heading)
        extra_identity = value.count(identity) - 1
        assert extra_heading >= 0 and extra_identity >= 0
        counts.subtract({c: n * extra_heading for c, n in Counter(heading).items()})
        counts.subtract({c: n * extra_identity for c, n in Counter(identity).items()})
    assert counts == Counter(expected), ('Rendered character inventory differs', dict(counts - Counter(expected)), dict(Counter(expected) - counts))
    paragraphs = Counter(compact(p.get_text()) for p in source.select('p') if compact(p.get_text()))
    assert all(value.count(p) >= n for p, n in paragraphs.items()), 'A paragraph is missing or interrupted'
    if word:
        # The Word cover can omit a footer; word_pages already validates every
        # footer that is removed. Here inspect all boxes against page bounds.
        from lxml import etree
        bounds = []
        for index, page in enumerate(etree.fromstring(bbox).findall('.//{*}page'), 1):
            for node in page.findall('.//{*}word'):
                if (float(node.get('xMin')) < 0 or float(node.get('yMin')) < 0 or
                    float(node.get('xMax')) > float(page.get('width')) or float(node.get('yMax')) > float(page.get('height'))):
                    bounds.append(dict(page=index, text=node.text, box=dict(node.attrib)))
    else: bounds = page_bounds(bbox)
    return dict(pages=len(pages), paragraph_counts_verified=sum(paragraphs.values()),
                character_inventory_verified=True, bounds_issues=bounds,
                verified_generated_o_markers=letter_bullets,
                line_box_overlaps=line_box_overlaps(bbox), long_row=repeated)


def pair(before, after, case, language, profile, partial_projection, compacted=False):
    old = before / 'renders' / (case + '-' + language)
    new = after / 'renders' / (case + '-' + profile + '-' + language)
    originals = [BeautifulSoup(p.with_suffix('.html').read_text(), 'html.parser') for p in [old, new]]
    review, candidate = originals
    for fmt in ['html', 'pdf', 'docx']:
        receipts = [json.loads(p.with_suffix('.' + fmt + '.fixture.json').read_text()) for p in [old, new]]
        for key in ['recipe_sha256', 'events_sha256', 'km_sha256']: assert receipts[0][key] == receipts[1][key]
        assert receipts[0]['output_profile'] == 'review' and receipts[1]['output_profile'] == profile
        for p, receipt in zip([before, after], receipts):
            manifest = json.loads((p / 'manifest.json').read_text()) if not compacted else json.loads((p.parent / 'provenance' / (p.name + '-manifest.json')).read_text())
            assert receipt['package_sha256'] == manifest['sha256'][language + '.zip']
    if profile == 'review':
        assert old.with_suffix('.html').read_bytes() == new.with_suffix('.html').read_bytes()
    else:
        expected_partial = partial_projection(review, language)
        compare(expected_partial, candidate, language)
    docs = [Document(p.with_suffix('.docx')) for p in [old, new]]
    source = candidate.body
    source_text = compact(source.get_text())
    from docx.text.paragraph import Paragraph
    question_nodes = body(docs[1])
    prefix = list(docs[1].element.body)[:-len(question_nodes)]
    prefix_texts = [Paragraph(p, docs[1]).text for node in prefix for p in node.iter(qn('w:p'))]
    word_text = compact(''.join(prefix_texts + paragraph_texts(docs[1], candidate)))
    word_counts = Counter(word_text)
    for header in {compact(t.get_text()) for t in source.select('.resource-table thead')}:
        assert all(header not in compact(n.get_text()) for n in source.select('.answer-detail'))
        tables = [t for t in docs[1].tables if compact(''.join(c.text for c in t.rows[0].cells)) == header]
        assert len(tables) == word_text.count(header)
        extra = len(tables) - source_text.count(header); assert extra >= 0
        word_counts.subtract({c: n * extra for c, n in Counter(header).items()})
    assert word_counts == Counter(source_text), 'DOCX body character inventory differs from HTML'
    paragraphs = Counter(compact(p.get_text()) for p in source.select('p') if compact(p.get_text()))
    assert all(word_text.count(text) >= count for text, count in paragraphs.items())
    assert [sorted(r.target_ref for r in d.part.rels.values() if r.is_external) for d in docs][0] == [sorted(r.target_ref for r in d.part.rels.values() if r.is_external) for d in docs][1]
    with zipfile.ZipFile(old.with_suffix('.docx')) as a, zipfile.ZipFile(new.with_suffix('.docx')) as b:
        for name in ['word/styles.xml', 'word/fontTable.xml']: assert a.read(name) == b.read(name)
    pdfs = [p.with_suffix('.pdf') for p in [old, new]]
    previews = [root / 'word-preview' / (p.name + '.pdf') for root, p in zip([before, after], [old, new])]
    if profile == 'review':
        assert [xml(n) for n in docs[0].element.body] == [xml(n) for n in docs[1].element.body]
        for paths in [pdfs, previews]:
            assert geometry(paths[0]) == geometry(paths[1]) and raster_digest(paths[0]) == raster_digest(paths[1])
    output = {}
    for kind, paths in [('native_pdf', pdfs), ('word_preview', previews)]:
        current = check_visible(paths[1], candidate, word=kind == 'word_preview')
        current['baseline_pages'] = len(geometry(paths[0]))
        assert current['pages'] <= current['baseline_pages'], (case, language, profile, kind, 'More pages')
        assert not current['bounds_issues'], 'Text outside page bounds'
        old_bbox = subprocess.check_output(['pdftotext', '-bbox-layout', str(paths[0]), '-'])
        current['baseline_line_box_overlaps'] = line_box_overlaps(old_bbox)
        # A count is only a screening signal; retain full coordinates for review.
        assert len(current['line_box_overlaps']) <= len(current['baseline_line_box_overlaps'])
        output[kind] = current
    paragraphs = verify_preview_paragraphs(docs[1], previews[1], candidate)
    if case == 'notice-mixed':
        assert candidate.select('.answer-detail .data-gap')
        for pdf in [pdfs[1], previews[1]]:
            assert 'AUTHORED-NOTICE:' in subprocess.check_output(['pdftotext', str(pdf), '-'], text=True)
    return dict(case=case, language=language, profile=profile, review_unchanged=profile == 'review',
                remaining_owned_notices=len(owned_gaps(candidate)), word_paragraphs_verified=paragraphs,
                rendered=output, artifacts={fmt: sha(new.with_suffix('.' + fmt)) for fmt in ['html', 'pdf', 'docx']},
                word_preview_sha256=sha(previews[1]))


def main():
    p = argparse.ArgumentParser(description=__doc__)
    for name in ['before', 'after', 'english', 'output']: p.add_argument('--' + name, type=Path, required=True)
    a = p.parse_args(); assert not a.output.exists()
    sys.path[:0] = [str(a.english.resolve() / 'scripts'), str(a.english.resolve() / 'tests')]
    from output_profile_contract import expected
    report = dict(selected_checks_passed=False, prototype_only=True, release_acceptance=False,
                  microsoft_word_acceptance=False, global_switch_complete=False,
                  checker_sha256=sha(Path(__file__)), rows=[])
    try:
        for case in ['empty', 'notice-mixed', 'budget-long-no-currency']:
            for language in ['english', 'chinese']:
                for profile in ['review', 'submission']:
                    row = pair(a.before, a.after, case, language, profile, expected)
                    report['rows'].append(row)
                    print(json.dumps(dict(case=case, language=language, profile=profile,
                                         pages={k: v['pages'] for k, v in row['rendered'].items()})), flush=True)
        report['selected_checks_passed'] = len(report['rows']) == 12
    except Exception as error: report['failure'] = repr(error); raise
    finally: a.output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n')


if __name__ == '__main__': main()
